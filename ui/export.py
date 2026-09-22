"""ui/export.py — THE table export (design-system decision #12: every data
table gets an Export action). Owner directive 2026-09-22.

One .xlsx per table, built from the RAW frame — never from display strings:

* numeric cells stay numbers and carry an Excel number format from the
  per-column ``formats`` spec (pct / usd / usd_k / x / int / …), so the
  sheet sorts and sums;
* absent values render as the literal ``n/a`` — never 0, never blank (the
  cardinal rule: no plausible-wrong number);
* header row bold + frozen, AutoFilter on, widths fitted;
* a ``Source`` sheet carries provenance (page, ticker, cert, source,
  as-of, units note, exported-at) so a file found on a desk a month later
  still says what it is;
* FDIC values stay in $thousands (owner choice 2026-09-22) — the ``usd_k``
  format is only accepted on a column whose header says ``($K)``, and the
  Source sheet spells the unit out.

Callers reach this through ``ui.chrome.table_export`` (re-exported) and
pass ``formats=`` / ``provenance=``. The workbook is built lazily — the
callable ``data`` runs only when the user clicks — so 40 exports on a
page cost nothing until one is wanted.
"""
from __future__ import annotations

import datetime as _dt
import html as _html
import io
import math
import re

import pandas as pd
import streamlit as st

# ── Format vocabulary ──────────────────────────────────────────────────
# Percent columns hold PERCENT UNITS throughout the app (11.24 means 11.24%,
# see utils.formatting.format_value), so the literal-% format is the honest
# one; Excel's native 0.00% would multiply by 100 and show 1124.00%.
FORMATS: dict[str, str] = {
    "pct":   '0.00"%"',
    "pct1":  '0.0"%"',
    "usd":   '$#,##0',        # whole US dollars
    "usd2":  '$#,##0.00',     # per-share / price
    "usd_k": '#,##0',         # FDIC $thousands — header must say ($K)
    "x":     '0.00"x"',
    "int":   '#,##0',
    "num":   '#,##0.00',
    "date":  'yyyy-mm-dd',
    "text":  '@',
}
NA = "n/a"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Provenance "Source" row shared by every Summary-of-Deposits export (branch
# details, market share, geography, proximity, merger HHI): the store is
# owner-resolved and post-survey mergers are re-attributed nightly.
SOD_SOURCE = ("FDIC Summary of Deposits (owner-resolved branches store; "
              "post-survey mergers re-attributed from FDIC merger history)")

_MISSING_STRINGS = {"", "—", "–", "-", "n/a", "N/A", "na", "nan", "NaN",
                    "None", "null"}
_HTML_TAG = re.compile(r"<[a-zA-Z/!][^>]*>")
_UNIT_SUFFIX = re.compile(r"\s*\(\$[BMK]\)\s*$")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_BAD_FILE_CHARS = re.compile(r"[^A-Za-z0-9._-]+")


# ── Cell coercion ──────────────────────────────────────────────────────

def _is_missing(v) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and math.isnan(v):
        return True
    if isinstance(v, str) and v.strip() in _MISSING_STRINGS:
        return True
    if v is pd.NA or v is pd.NaT:
        return True
    if hasattr(v, "dtype"):          # numpy / pandas scalar (never a frame)
        try:
            return bool(pd.isna(v))
        except (TypeError, ValueError):
            return False
    return False


def _strip_html(s: str) -> str:
    if "<" in s and ">" in s and _HTML_TAG.search(s):
        s = _HTML_TAG.sub("", s)
        return _html.unescape(s).strip()
    return s


def _coerce(v, fmt: str | None):
    """A cell value openpyxl can write: python scalars, tz-naive datetimes,
    ``NA`` for anything absent. Strings under a ``date`` format parse to a
    date when they are ISO-shaped; other strings lose any HTML markup."""
    if _is_missing(v):
        return NA
    if isinstance(v, bool):
        return v
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):   # numpy scalar
        try:
            v = v.item()
        except (TypeError, ValueError):
            pass
    if isinstance(v, pd.Timestamp):
        v = v.to_pydatetime()
    if isinstance(v, _dt.datetime):
        if v.tzinfo is not None:
            v = v.astimezone(_dt.timezone.utc).replace(tzinfo=None)
        return v
    if isinstance(v, _dt.date):
        return v
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isinf(v):
            return NA
        return v
    if isinstance(v, str):
        s = _strip_html(v)
        if fmt == "date":
            try:
                return _dt.date.fromisoformat(s[:10])
            except ValueError:
                return s
        return s
    if isinstance(v, (dict, list, tuple, set)):
        return str(v)
    return str(v)


# ── Naming ─────────────────────────────────────────────────────────────

def sheet_title(name: str, used: set[str] | None = None) -> str:
    """Excel-legal, ≤31 chars, unique within ``used`` (mutated)."""
    t = _BAD_SHEET_CHARS.sub(" ", str(name or "Data")).strip() or "Data"
    t = t[:31].rstrip()
    if used is not None:
        base, n = t, 2
        while t in used:
            suffix = f" ({n})"
            t = base[:31 - len(suffix)].rstrip() + suffix
            n += 1
        used.add(t)
    return t


def safe_filename(stem: str) -> str:
    s = _BAD_FILE_CHARS.sub("_", str(stem)).strip("_.")
    s = re.sub(r"_+", "_", s)
    return s or "export"


def export_header(label: str) -> str:
    """A screener metric label with a scaled-unit suffix — ``Mkt Cap ($B)``
    — sits over RAW-dollar cells in the export, so the suffix becomes
    ``($)``; the value is not rescaled (never guess units)."""
    return _UNIT_SUFFIX.sub(" ($)", label) if _UNIT_SUFFIX.search(label) else label


# ── Workbook build ─────────────────────────────────────────────────────

def _validate_formats(columns, formats: dict[str, str]) -> dict[str, str]:
    out = {}
    for col, f in (formats or {}).items():
        if col not in columns:
            continue          # tolerated: a spec written for an optional column
        nf = FORMATS.get(f, f)   # a vocabulary key or a raw Excel format string
        if f == "usd_k" and "($K)" not in str(col):
            raise ValueError(
                f"usd_k column {col!r} must carry '($K)' in its header — "
                "FDIC thousands are exported unscaled and the header is the "
                "only place the unit lives")
        out[col] = nf
    return out


def _units_note(columns, formats: dict[str, str]) -> str:
    parts = []
    k_cols = [c for c, f in formats.items() if f == "usd_k"]
    d_cols = [c for c, f in formats.items() if f in ("usd", "usd2")]
    p_cols = [c for c, f in formats.items() if f in ("pct", "pct1")]
    x_cols = [c for c, f in formats.items() if f == "x"]
    if d_cols:
        parts.append("Dollar columns are whole US dollars.")
    if k_cols:
        parts.append("Columns marked ($K) are FDIC-reported thousands of "
                     "dollars, unscaled.")
    if p_cols:
        parts.append("Percent columns hold percent units (11.24 = 11.24%).")
    if x_cols:
        parts.append("Multiples hold the raw ratio (10.2 = 10.2x).")
    return " ".join(parts) or "Values are exported as stored; see column headers."


def build_workbook(df: pd.DataFrame, *, sheet: str = "Data",
                   formats: dict[str, str] | None = None,
                   provenance: dict[str, object] | None = None,
                   freeze_cols: int = 0) -> bytes:
    """The .xlsx bytes for one table. ``formats`` maps column → FORMATS key
    (or a raw Excel number format). ``provenance`` rows land on the Source
    sheet after the always-present ``Exported`` stamp."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    formats = dict(formats or {})
    cols = [str(c) for c in df.columns]
    nf_by_col = _validate_formats(cols, formats)

    wb = Workbook()
    ws = wb.active
    used = set()
    ws.title = sheet_title(sheet, used)

    hdr_font = Font(bold=True)
    hdr_fill = PatternFill("solid", fgColor="E8EEF5")
    na_font = Font(italic=True, color="808080")
    na_align = Alignment(horizontal="right")

    ws.append(cols)
    for ci in range(1, len(cols) + 1):
        c = ws.cell(1, ci)
        c.font, c.fill = hdr_font, hdr_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)

    widths = [len(h) for h in cols]
    for rec in df.itertuples(index=False, name=None):
        row = []
        for ci, (col, v) in enumerate(zip(cols, rec)):
            cv = _coerce(v, formats.get(col))
            row.append(cv)
            widths[ci] = max(widths[ci], min(len(str(cv)), 40))
        ws.append(row)

    last_row = ws.max_row
    for ci, col in enumerate(cols, start=1):
        nf = nf_by_col.get(col)
        for rr in range(2, last_row + 1):
            c = ws.cell(rr, ci)
            if c.value == NA:
                c.font, c.alignment = na_font, na_align
            elif nf and (isinstance(c.value, (int, float, _dt.date))
                         and not isinstance(c.value, bool)):
                c.number_format = nf
            elif nf == FORMATS["text"] and isinstance(c.value, str):
                c.number_format = nf
        ws.column_dimensions[get_column_letter(ci)].width = max(8, min(42, widths[ci - 1] + 2))
    ws.freeze_panes = ws.cell(2, max(0, int(freeze_cols)) + 1).coordinate
    if cols:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(last_row, 1)}"

    # ── Source sheet ─────────────────────────────────────────────────
    src = wb.create_sheet(sheet_title("Source", used))
    rows: list[tuple[str, object]] = [
        ("Exported", _dt.datetime.now(_dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Table", ws.title),
        ("Rows", len(df)),
    ]
    for k, v in (provenance or {}).items():
        if v is None or (isinstance(v, str) and not v.strip()):
            continue
        rows.append((str(k), v if isinstance(v, (int, float)) else str(v)))
    rows.append(("Units", _units_note(cols, formats)))
    rows.append(("Missing values", f'"{NA}" = the source does not report the value or a '
                                   "precondition failed; nothing is estimated."))
    for k, v in rows:
        src.append([k, v])
    for rr in range(1, len(rows) + 1):
        src.cell(rr, 1).font = hdr_font
        src.cell(rr, 2).alignment = Alignment(wrap_text=True, vertical="top")
    src.column_dimensions["A"].width = 18
    src.column_dimensions["B"].width = 96

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Screener metric columns (config.METRICS) ───────────────────────────

def metric_format(m: dict) -> str:
    """FORMATS key for a config.METRICS entry. Scaled dollar formats
    (millions / billions / dollars_auto) export the raw dollars under a
    ``($)`` header — see export_header."""
    fmt, dec = m.get("format", "number"), int(m.get("decimals", 2) or 0)
    if fmt == "pct":
        return "pct1" if dec <= 1 else "pct"
    if fmt == "ratio":
        return "x"
    if fmt == "currency":
        return "usd2"
    if fmt in ("millions", "billions", "dollars_auto"):
        return "usd"
    if fmt == "number":
        return "int" if dec == 0 else "num"
    if fmt == "date":
        return "date"
    return "text"


def metric_columns(keys) -> tuple[dict[str, str], dict[str, str]]:
    """(rename, formats) for a list of screener metric keys: header labels
    with honest unit suffixes + the number format per label. Unknown keys
    pass through as themselves with no format."""
    from config import METRICS_BY_KEY
    rename, formats = {}, {}
    for k in keys:
        m = METRICS_BY_KEY.get(k)
        if not m:
            rename[k] = k
            continue
        label = export_header(m.get("label", k))
        rename[k] = label
        formats[label] = metric_format(m)
    return rename, formats


# ── The control ────────────────────────────────────────────────────────

def table_export(df, filename: str, key: str, *, sheet: str | None = None,
                 formats: dict[str, str] | None = None,
                 provenance: dict[str, object] | None = None,
                 freeze_cols: int = 0, label: str = "Export") -> None:
    """Small right-aligned Export action for a data table. Writes one
    .xlsx (see module doc); the workbook is built on click."""
    fname = safe_filename(filename)
    sheet_name = sheet or fname

    def _build() -> bytes:
        return build_workbook(df, sheet=sheet_name, formats=formats,
                              provenance=provenance, freeze_cols=freeze_cols)

    # The keyed container carries the compact right-aligned styling
    # (styles.py `st-key-tblexp_`); a bare download_button rendered as a
    # full-size button parked at the left under every table.
    with st.container(key=f"tblexp_{key}"):
        st.download_button(label, _build, file_name=f"{fname}.xlsx",
                           mime=XLSX_MIME, key=key)
