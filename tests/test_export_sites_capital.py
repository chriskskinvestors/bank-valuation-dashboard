"""Table exports on the shared .xlsx exporter (ui/export.py) for the three
Capital Adequacy tables ui/capital_dynamics.py renders as .ksk-grid markdown
(owner directive 2026-09-22: every data table gets an Export, built from the
RAW frame — never display strings):

  * holding-company capital table (_render_holdco_capital) — SEC inline XBRL;
  * holding-company regulatory-capital walk (_render_holdco_walk);
  * Capital Return Attribution quarterly detail
    (_render_capital_return_attribution) — SEC companyfacts.

Each test drives the real render function against fixture data, captures the
deferred workbook callable table_export hands st.download_button, loads the
bytes with openpyxl and asserts HAND-COMPUTED cells, the Excel number format
per row/column, "n/a" where the fixture had no value, and the Source sheet.

Units pinned here:
  * the SEC scraper stores capital RATIOS as FRACTIONS (0.1089) and the
    screen shows ×100 ("10.89%") → the sheet carries percent units under
    "(%)", and the test proves the exported number formats to the SAME
    string the screen rendered;
  * capital amounts / RWA / dividends / buybacks / net income are whole
    dollars → "($)";
  * the capital-return payout ratios are fractions (×100 on screen, 1 dp);
    share_change_pct is ALREADY percent units at source and is not rescaled.

Stub discipline: each render module's own ``st`` binding is patched with a
PRIVATE stub per test (mock.patch.object); nothing here touches sys.modules
or the shared package stub. ui.capital_dynamics is un-imported after the
class when this module was the first to import it, so test_render_smoke's
stub-swapping classes still see a fresh module under discovery (the same
guard tests/test_export_sites_highlights.py uses).

Run: PYTHONIOENCODING=utf-8 python -m unittest tests.test_export_sites_capital
"""
from __future__ import annotations

import io
import sys
import types
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).parent.parent))

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _noop(*a, **k):
    return None


class _Ctx:
    """Context-manager placeholder (columns / container / expander / st.empty
    slot): every attribute is a no-op."""

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def __getattr__(self, name):
        return _noop


class _StubSt:
    """Private streamlit stand-in: unknown widgets are no-ops; layout widgets
    return contexts; radio returns ``radio_value`` (the page's period
    toggle); markdown/caption text is recorded so a test can compare the
    on-screen string with the exported number."""

    session_state: dict = {}
    query_params: dict = {}

    def __init__(self, radio_value="Annual"):
        self.radio_value = radio_value
        self.text: list[str] = []

    def __getattr__(self, name):
        return _noop

    def markdown(self, s, *a, **k):
        self.text.append(str(s))

    caption = markdown

    def columns(self, spec, **k):
        n = spec if isinstance(spec, int) else len(spec)
        return [_Ctx() for _ in range(n)]

    def container(self, *a, **k):
        return _Ctx()

    expander = spinner = popover = empty = container

    def radio(self, label, options=None, **k):
        return self.radio_value


class _CapitalSite(unittest.TestCase):
    """Binds a private stub onto ui.capital_dynamics (+ the chrome/states
    modules its render paths reach through) and a download_button capture
    onto ui.export."""

    RADIO = "Annual"

    @classmethod
    def setUpClass(cls):
        import ui.chrome, ui.export, ui.states  # noqa: E401
        cls.X = ui.export
        cls._imported_here = "ui.capital_dynamics" not in sys.modules
        import ui.capital_dynamics
        cls.CD = ui.capital_dynamics
        cls._modules = (ui.chrome, ui.states, ui.capital_dynamics)

    @classmethod
    def tearDownClass(cls):
        if cls._imported_here:
            sys.modules.pop("ui.capital_dynamics", None)

    def setUp(self):
        self.calls = []          # [(label, data, kwargs)] from download_button
        self.stub = _StubSt(self.RADIO)
        fake_export_st = types.SimpleNamespace(
            download_button=lambda label, data, **k: self.calls.append((label, data, k)),
            container=lambda *a, **k: _Ctx())
        for mod, obj in [(m, self.stub) for m in self._modules] + [(self.X, fake_export_st)]:
            self._patch(mod, "st", obj)
        self._patch(self.CD, "get_cik", lambda t: 1281761)
        self._patch(self.CD, "get_fdic_cert", lambda t: 12368)
        self._patch(self.CD, "get_name", lambda t: "Regions Financial Corp")

    def _patch(self, obj, name, value):
        p = mock.patch.object(obj, name, value)
        p.start()
        self.addCleanup(p.stop)

    # ── helpers ──────────────────────────────────────────────────────
    def _book(self, i=0, expect=1):
        self.assertEqual(len(self.calls), expect,
                         f"expected {expect} Export control(s), got {len(self.calls)}")
        label, data, kw = self.calls[i]
        self.assertEqual(label, "Export")
        self.assertTrue(callable(data), "workbook must build lazily, on click")
        wb = load_workbook(io.BytesIO(data()))
        return wb, wb.worksheets[0], kw

    @staticmethod
    def _grid(ws):
        return [[c.value for c in r] for r in ws.iter_rows()]

    @staticmethod
    def _source(wb):
        return dict((r[0].value, r[1].value) for r in wb["Source"].iter_rows())

    def _assert_na(self, ws, coord):
        self.assertEqual(ws[coord].value, "n/a")
        self.assertTrue(ws[coord].font.italic)

    def _screen(self):
        return "\n".join(self.stub.text)


# ═══════════════════════════════════════════════════════════════════════
# Holding-company capital table + walk (SEC inline XBRL, scraper output)
# ═══════════════════════════════════════════════════════════════════════

META = {"form": "10-K", "date": "2026-02-24", "accession": "000128176126000012",
        "doc": "rf-20251231.htm"}

# FY2025: Regions-shaped (same values as test_render_smoke's holdco fixture):
# ratios are FRACTIONS as the scraper stores them; amounts whole dollars.
# The walk is flagged as reconciling with AOCI retained (opt-in); the UI
# trusts the scraper's _walk_reconciles flag, so the fixture sets it.
FY25 = {
    "cet1_ratio": 0.1089, "t1_ratio": 0.1199, "total_ratio": 0.1389,
    "lev_ratio": 0.0968, "cet1_cap": 13.49e9, "t1_cap": 14.859e9,
    "tier2_cap": 2.346e9, "total_cap": 17.205e9, "rwa": 123.9e9,
    "_anchored": True, "_walk_reconciles": True,
    "_walk": {"common_equity": 19.00e9, "goodwill": 5.733e9,
              "other_intangibles": 0.140e9, "aoci": -1.535e9,
              "subordinated_debt": 1.00e9, "intangibles": 5.873e9,
              "aoci_treatment": "included"},
}
# FY2024: leverage ratio, Tier 2 and RWA not tagged (→ n/a, never 0); the
# CET1 build does NOT reconcile, so every walk component step is n/a while
# the bridge totals (extracted amounts) still export.
FY24 = {
    "cet1_ratio": 0.1050, "t1_ratio": 0.1150, "total_ratio": 0.1350,
    "cet1_cap": 12.0e9, "t1_cap": 13.0e9, "total_cap": 15.0e9,
    "_anchored": True, "_walk_reconciles": False,
    "_walk": {"common_equity": 17.0e9, "goodwill": 5.7e9,
              "other_intangibles": None, "aoci": -2.0e9,
              "subordinated_debt": None, "intangibles": 5.7e9,
              "aoci_treatment": None},
}
# Q3 '25 (quarterly view): AOCI opt-OUT — the walk removes AOCI, so the
# "Less: AOCI removed" step is −AOCI = +1.2B (a loss added back).
Q3_25 = {
    "cet1_ratio": 0.1070, "t1_ratio": 0.1180, "total_ratio": 0.1370,
    "lev_ratio": 0.0950, "cet1_cap": 13.0e9, "t1_cap": 14.3e9,
    "tier2_cap": 2.2e9, "total_cap": 16.5e9, "rwa": 121.0e9,
    "_anchored": True, "_walk_reconciles": True,
    "_walk": {"common_equity": 20.0e9, "goodwill": 5.7e9,
              "other_intangibles": 0.1e9, "aoci": -1.2e9,
              "subordinated_debt": 1.1e9, "intangibles": 5.8e9,
              "aoci_treatment": "excluded"},
}


class _HoldcoSite(_CapitalSite):

    def _render(self, cap, *, quarterly=False):
        import data.ir_provider as irp
        import data.sec_filing_scraper as sfs
        res = {"meta": dict(META), "capital": cap}
        self._patch(irp, "fresh_capital", lambda cik: None)
        if quarterly:
            self._patch(sfs, "holdco_capital_quarterly_for", lambda cik, cert=None, n=8: res)
        else:
            self._patch(sfs, "holdco_capital_for", lambda cik, cert=None: res)
        self.CD._render_holdco_capital("RF")


class TestHoldcoCapitalAnnual(_HoldcoSite):
    RADIO = "Annual"

    def test_capital_table_fraction_to_percent_dollars_and_absent(self):
        self._render({"2025-12-31": dict(FY25), "2024-12-31": dict(FY24)})
        self.assertEqual(len(self.calls), 2, "capital table + walk exports")
        wb, ws, kw = self._book(0, expect=2)
        self.assertEqual(kw["file_name"], "capital_holdco_RF_Annual_2025-12-31.xlsx")
        self.assertEqual(kw["key"], "exp_hc_capital_RF_Annual")
        g = self._grid(ws)
        # Columns follow the screen: newest period first.
        self.assertEqual(g[0], ["Line item", "FY2025", "FY2024"])
        self.assertEqual(len(g), 1 + 9)
        rows = {r[0]: r for r in g[1:]}
        by_label = {r[0]: i + 1 for i, r in enumerate(g)}

        # ── ratios: stored fraction × 100 = percent units, "(%)" ──
        cet1 = rows["Common Equity Tier 1 ratio (%)"]
        self.assertAlmostEqual(cet1[1], 10.89, places=9)
        self.assertAlmostEqual(cet1[2], 10.50, places=9)
        self.assertAlmostEqual(rows["Tier 1 capital ratio (%)"][1], 11.99, places=9)
        self.assertAlmostEqual(rows["Total capital ratio (%)"][1], 13.89, places=9)
        self.assertAlmostEqual(rows["Tier 1 leverage ratio (%)"][1], 9.68, places=9)
        # The exported number formats to exactly the string the screen shows.
        self.assertEqual(f"{cet1[1]:.2f}%", "10.89%")
        self.assertIn("10.89%", self._screen())
        self.assertIn("9.68%", self._screen())
        r = by_label["Common Equity Tier 1 ratio (%)"]
        self.assertEqual(ws.cell(r, 2).number_format, '0.00"%"')
        self.assertEqual(ws.cell(r, 3).number_format, '0.00"%"')

        # ── amounts: whole dollars as tagged, "($)" ──
        self.assertEqual(rows["Common Equity Tier 1 capital ($)"][1:], [13_490_000_000, 12_000_000_000])
        self.assertEqual(rows["Tier 1 capital ($)"][1:], [14_859_000_000, 13_000_000_000])
        self.assertEqual(rows["Tier 2 capital ($)"][1], 2_346_000_000)
        self.assertEqual(rows["Total capital ($)"][1:], [17_205_000_000, 15_000_000_000])
        self.assertEqual(rows["Risk-weighted assets ($)"][1], 123_900_000_000)
        r = by_label["Common Equity Tier 1 capital ($)"]
        self.assertEqual(ws.cell(r, 2).number_format, "$#,##0")
        self.assertEqual(ws.cell(r, 3).number_format, "$#,##0")

        # ── FY2024 lines the filing did not tag → n/a, not 0 ──
        for lab in ("Tier 1 leverage ratio (%)", "Tier 2 capital ($)",
                    "Risk-weighted assets ($)"):
            self._assert_na(ws, f"C{by_label[lab]}")
        self.assertEqual(ws["A2"].number_format, "General")
        self.assertEqual(ws.freeze_panes, "B2")

        src = self._source(wb)
        self.assertIn("Capital Adequacy", src["Page"])
        self.assertIn("Regulatory Capital", src["Page"])
        self.assertIn("holding company capital", src["Page"])
        self.assertEqual(src["Ticker"], "RF")
        self.assertEqual(src["Company"], "Regions Financial Corp")
        self.assertEqual(src["SEC CIK"], 1281761)
        self.assertEqual(src["FDIC cert"], 12368)
        self.assertIn("SEC 10-K filed 2026-02-24", src["Source"])
        self.assertIn("inline XBRL", src["Source"])
        self.assertIn("holding-company consolidated", src["Source"])
        self.assertNotIn("bank-subsidiary basis", src["Source"])
        self.assertEqual(src["Filing"], "https://www.sec.gov/Archives/edgar/data/1281761/"
                                        "000128176126000012/rf-20251231.htm")
        self.assertEqual(src["View"], "Annual")
        self.assertEqual(src["Periods"], "FY2024 to FY2025")
        self.assertEqual(src["Report date"], "2025-12-31")
        self.assertIn("whole US dollars", src["Units"])
        self.assertIn("percent units", src["Units"])
        self.assertIn("capital ÷ ratio", src["Notes"])

    def test_walk_negated_deductions_bridges_and_non_reconciling_period(self):
        self._render({"2025-12-31": dict(FY25), "2024-12-31": dict(FY24)})
        wb, ws, kw = self._book(1, expect=2)
        self.assertEqual(kw["file_name"], "capital_walk_holdco_RF_Annual_2025-12-31.xlsx")
        self.assertEqual(kw["key"], "exp_hc_walk_RF_Annual")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Line item", "FY2025", "FY2024"])
        self.assertEqual(len(g), 1 + 11)
        rows = {r[0]: r for r in g[1:]}
        by_label = {r[0]: i + 1 for i, r in enumerate(g)}
        # Bold markers are screen markup — not in the row labels.
        self.assertIn("= Common Equity Tier 1 capital ($)", rows)
        self.assertFalse(any("*" in r[0] for r in g[1:]))

        # FY2025 (reconciles): components as tagged, deductions NEGATIVE.
        self.assertEqual(rows["Total common equity ($)"][1], 19_000_000_000)
        self.assertEqual(rows["Less: goodwill ($)"][1], -5_733_000_000)
        self.assertEqual(rows["Less: other intangibles ($)"][1], -140_000_000)
        # AOCI opt-in: no walk step — the screen's textual marker, verbatim.
        self.assertEqual(rows["Less: AOCI removed (opt-out) ($)"][1], "— (in CET1)")
        self.assertEqual(rows["= Common Equity Tier 1 capital ($)"][1], 13_490_000_000)
        # AT1 = Tier 1 − CET1 = 14.859B − 13.49B = 1.369B
        self.assertAlmostEqual(rows["Additional Tier 1 (qualifying preferred) ($)"][1],
                               1_369_000_000, delta=1)
        self.assertEqual(rows["= Tier 1 capital ($)"][1], 14_859_000_000)
        self.assertEqual(rows["Subordinated debt & qualifying Tier 2 ($)"][1], 1_000_000_000)
        # Other Tier 2 = Tier 2 − sub-debt = 2.346B − 1.00B = 1.346B
        self.assertAlmostEqual(rows["Other Tier 2 (allowance & adjustments) ($)"][1],
                               1_346_000_000, delta=1)
        self.assertEqual(rows["= Tier 2 capital ($)"][1], 2_346_000_000)
        self.assertEqual(rows["= Total capital ($)"][1], 17_205_000_000)
        # Screen shows the same values scaled to $B, from the same raw numbers.
        self.assertIn("19.00B", self._screen())
        self.assertIn("-5.73B", self._screen())
        self.assertIn("1.37B", self._screen())
        self.assertIn("— (in CET1)", self._screen())

        # FY2024 (does NOT reconcile): every component step n/a; bridge totals
        # (the FDIC-anchored extracted amounts) still export; AT1 derived.
        for lab in ("Total common equity ($)", "Less: goodwill ($)",
                    "Less: other intangibles ($)", "Less: AOCI removed (opt-out) ($)",
                    "Subordinated debt & qualifying Tier 2 ($)",
                    "Other Tier 2 (allowance & adjustments) ($)", "= Tier 2 capital ($)"):
            self._assert_na(ws, f"C{by_label[lab]}")
        self.assertEqual(rows["= Common Equity Tier 1 capital ($)"][2], 12_000_000_000)
        self.assertEqual(rows["Additional Tier 1 (qualifying preferred) ($)"][2], 1_000_000_000)
        self.assertEqual(rows["= Tier 1 capital ($)"][2], 13_000_000_000)
        self.assertEqual(rows["= Total capital ($)"][2], 15_000_000_000)

        r = by_label["Less: goodwill ($)"]
        self.assertEqual(ws.cell(r, 2).number_format, "$#,##0")
        r = by_label["= Total capital ($)"]
        self.assertEqual(ws.cell(r, 3).number_format, "$#,##0")
        # The textual AOCI cell carries no number format.
        r = by_label["Less: AOCI removed (opt-out) ($)"]
        self.assertEqual(ws.cell(r, 2).number_format, "General")
        self.assertEqual(ws.freeze_panes, "B2")

        src = self._source(wb)
        self.assertIn("regulatory capital walk (holding company)", src["Page"])
        self.assertEqual(src["Ticker"], "RF")
        self.assertEqual(src["SEC CIK"], 1281761)
        self.assertIn("SEC 10-K filed 2026-02-24", src["Source"])
        self.assertEqual(src["Periods"], "FY2024 to FY2025")
        self.assertEqual(src["Report date"], "2025-12-31")
        self.assertIn("NEGATIVE", src["Notes"])
        self.assertIn("Additional Tier 1 = Tier 1 − CET1", src["Notes"])
        self.assertIn("whole US dollars", src["Units"])

    def test_no_reconciling_period_exports_capital_table_only(self):
        # The walk section is gated off entirely (n/a note, no table) — so
        # there must be no walk export either: one Export control.
        self._render({"2024-12-31": dict(FY24)})
        wb, ws, kw = self._book(0, expect=1)
        self.assertEqual(kw["key"], "exp_hc_capital_RF_Annual")
        self.assertEqual(self._grid(ws)[0], ["Line item", "FY2024"])


class TestHoldcoCapitalQuarterly(_HoldcoSite):
    RADIO = "Quarterly"

    def test_quarter_labels_bank_basis_note_and_aoci_opt_out(self):
        q3 = {**Q3_25, "_basis": "bank"}
        self._render({"2025-09-30": q3, "2025-06-30": dict(FY25)}, quarterly=True)
        wb, ws, kw = self._book(0, expect=2)
        self.assertEqual(kw["file_name"], "capital_holdco_RF_Quarterly_2025-09-30.xlsx")
        self.assertEqual(kw["key"], "exp_hc_capital_RF_Quarterly")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Line item", "Q3 '25", "Q2 '25"])
        rows = {r[0]: r for r in g[1:]}
        self.assertAlmostEqual(rows["Common Equity Tier 1 ratio (%)"][1], 10.70, places=9)
        self.assertEqual(rows["Risk-weighted assets ($)"][1], 121_000_000_000)
        src = self._source(wb)
        self.assertEqual(src["View"], "Quarterly")
        self.assertEqual(src["Periods"], "Q2 '25 to Q3 '25")
        self.assertEqual(src["Report date"], "2025-09-30")
        self.assertIn("bank-subsidiary basis", src["Source"])

        wb, ws, kw = self._book(1, expect=2)
        self.assertEqual(kw["file_name"], "capital_walk_holdco_RF_Quarterly_2025-09-30.xlsx")
        self.assertEqual(kw["key"], "exp_hc_walk_RF_Quarterly")
        rows = {r[0]: r for r in self._grid(ws)[1:]}
        # Opt-out: the removed AOCI is −(−1.2B) = +1.2B (loss added back).
        self.assertEqual(rows["Less: AOCI removed (opt-out) ($)"][1], 1_200_000_000)
        self.assertEqual(rows["Less: AOCI removed (opt-out) ($)"][2], "— (in CET1)")
        self.assertEqual(rows["Less: goodwill ($)"][1], -5_700_000_000)
        # Other Tier 2 = 2.2B − 1.1B
        self.assertAlmostEqual(rows["Other Tier 2 (allowance & adjustments) ($)"][1],
                               1_100_000_000, delta=1)
        self.assertIn("1.20B", self._screen())


# ═══════════════════════════════════════════════════════════════════════
# Capital Return Attribution — quarterly detail (SEC companyfacts)
# ═══════════════════════════════════════════════════════════════════════

def _timeline():
    """Three quarters, hand-set the way analysis.capital_return produces
    them: whole dollars; ratios as FRACTIONS (payout 0.2 = 20%); a loss
    quarter (Q3) leaves every ratio NaN; buybacks untagged in Q2 (NaN, the
    dividend-only total still an observation); share_change_pct ALREADY in
    percent units (pct_change × 100)."""
    nan = np.nan
    return pd.DataFrame({
        "end": ["2025-06-30", "2025-09-30", "2025-12-31"],
        "year": [2025, 2025, 2025], "quarter": [2, 3, 4],
        "net_income_q": [500e6, -50e6, 600e6],
        "dividends_q": [100e6, 100e6, 120e6],
        "buybacks_q": [nan, 50e6, 60e6],
        "total_returned_q": [100e6, 150e6, 180e6],
        "payout_ratio_q": [0.2, nan, 0.2],
        "buyback_ratio_q": [nan, nan, 0.1],
        "total_return_ratio_q": [0.2, nan, 0.3],
        "shares_outstanding": [100e6, 99e6, 98e6],
        "share_change_pct": [nan, -1.0, (98e6 - 99e6) / 99e6 * 100],
        "date": pd.to_datetime(["2025-06-30", "2025-09-30", "2025-12-31"]),
    })


class TestCapitalReturnQuarterlyDetail(_CapitalSite):

    def _render(self, timeline, dividend_source="common-specific"):
        import analysis.capital_return as cr
        import data.cache as cache
        result = {"timeline": timeline, "ttm": {}, "growth": {}, "yield": {},
                  "dividend_source": dividend_source}
        self._patch(cr, "summarize_capital_return",
                    lambda cik, market_cap=None, lookback_quarters=20: result)
        self._patch(cache, "get", lambda key, *a, **k: None)
        self.CD._render_capital_return_attribution("RF")

    def test_flat_frame_dollars_fraction_ratios_and_loss_quarter(self):
        self._render(_timeline())
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "capital_return_quarterly_RF_2025Q4.xlsx")
        self.assertEqual(kw["key"], "exp_capret_q_RF")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Quarter", "Net Income ($)", "Dividends ($)", "Buybacks ($)",
                                "Total Returned ($)", "Payout (%)", "Buyback (%)",
                                "Total Ret (%)", "Share Chg (%)"])
        self.assertEqual(len(g), 1 + 3)
        q2, q3, q4 = g[1], g[2], g[3]
        self.assertEqual(q2[:3], ["2025Q2", 500_000_000, 100_000_000])
        self._assert_na(ws, "D2")                        # buybacks untagged → n/a
        self.assertEqual(q2[4], 100_000_000)
        self.assertAlmostEqual(q2[5], 20.0, places=9)    # 0.2 → 20.0 (%)
        self._assert_na(ws, "G2")                        # buyback ratio NaN
        self.assertAlmostEqual(q2[7], 20.0, places=9)
        self._assert_na(ws, "I2")                        # first quarter: no prior shares
        # Loss quarter: amounts export, every ratio n/a (never a negative payout).
        self.assertEqual(q3[:5], ["2025Q3", -50_000_000, 100_000_000, 50_000_000, 150_000_000])
        for col in "FGH":
            self._assert_na(ws, f"{col}3")
        self.assertEqual(q3[8], -1.0)
        self.assertEqual(q4[:5], ["2025Q4", 600_000_000, 120_000_000, 60_000_000, 180_000_000])
        self.assertAlmostEqual(q4[5], 20.0, places=9)
        self.assertAlmostEqual(q4[6], 10.0, places=9)
        self.assertAlmostEqual(q4[7], 30.0, places=9)
        # share_change_pct is not rescaled: (98 − 99) / 99 × 100 = −1.0101…
        self.assertAlmostEqual(q4[8], -1.0101010101, places=9)
        # The exported numbers format to the screen's strings.
        screen = self._screen()
        self.assertEqual(f"{q4[6]:.1f}%", "10.0%")
        self.assertIn("10.0%", screen)
        self.assertEqual(f"{q4[8]:+.2f}%", "-1.01%")
        self.assertIn("-1.01%", screen)
        self.assertEqual(ws["B2"].number_format, "$#,##0")
        self.assertEqual(ws["E4"].number_format, "$#,##0")
        self.assertEqual(ws["F2"].number_format, '0.0"%"')
        self.assertEqual(ws["H4"].number_format, '0.0"%"')
        self.assertEqual(ws["I4"].number_format, '0.00"%"')
        self.assertEqual(ws.freeze_panes, "B2")

        src = self._source(wb)
        self.assertIn("Capital Return Attribution", src["Page"])
        self.assertIn("quarterly detail", src["Page"])
        self.assertEqual(src["Ticker"], "RF")
        self.assertEqual(src["Company"], "Regions Financial Corp")
        self.assertEqual(src["SEC CIK"], 1281761)
        self.assertNotIn("FDIC cert", src)
        self.assertIn("SEC companyfacts (holding company)", src["Source"])
        self.assertIn("YTD differences", src["Source"])
        self.assertIn("common-specific", src["Dividend basis"])
        self.assertIn("excludes preferred", src["Dividend basis"])
        self.assertEqual(src["Periods"], "2025Q2 to 2025Q4")
        self.assertEqual(src["Report date"], "2025-12-31")
        self.assertIn("whole US dollars", src["Units"])
        self.assertIn("percent units", src["Units"])
        self.assertIn("not positive", src["Notes"])

    def test_dividend_basis_row_names_the_total_including_preferred_fallback(self):
        self._render(_timeline(), dividend_source="total (includes preferred)")
        wb, ws, kw = self._book()
        src = self._source(wb)
        self.assertTrue(src["Dividend basis"].startswith("total (includes preferred)"))
        self.assertIn("may overstate common", src["Dividend basis"])


if __name__ == "__main__":
    unittest.main()
