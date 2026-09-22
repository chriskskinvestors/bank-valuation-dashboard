"""ui/export.py — the shared .xlsx table export (owner directive 2026-09-22:
"audit and seriously improve our excel exporting").

Hermetic: builds workbooks in memory and reads them back with openpyxl,
asserting cell VALUES and number formats against hand-written expectations.
Each Phase-1 audit defect is pinned here:

  D1  scaled-unit header over raw dollars ("Mkt Cap ($B)" = 2,100,000,000)
  D3  display strings exported instead of numbers
  D4  FDIC $thousands with no unit anywhere in the file
  D5  no provenance (as-of / source / ticker) in the file
  plus: n/a never 0, HTML never lands, sheet names legal + unique.
"""
import datetime as dt
import io
import math
import types
import unittest

import pandas as pd
from openpyxl import load_workbook

import ui.export as ex


def _load(b: bytes):
    return load_workbook(io.BytesIO(b))


def _grid(ws, max_rows=None):
    rows = []
    for r in ws.iter_rows(min_row=1, max_row=max_rows or ws.max_row):
        rows.append([c.value for c in r])
    return rows


class TestCellsAndFormats(unittest.TestCase):

    def test_numbers_stay_numbers_with_excel_formats(self):
        df = pd.DataFrame({
            "Ticker": ["BANR", "EWBC"],
            "ROE (%)": [11.24, 15.8],
            "Total Assets ($)": [16_200_000_000, 74_500_000_000],
            "P/E (x)": [10.2, 9.1],
            "Branches": [120, 95],
            "Price": [61.5, 98.1],
        })
        wb = _load(ex.build_workbook(
            df, sheet="Peers",
            formats={"ROE (%)": "pct", "Total Assets ($)": "usd", "P/E (x)": "x",
                     "Branches": "int", "Price": "usd2"}))
        ws = wb["Peers"]
        self.assertEqual(_grid(ws, 2), [
            ["Ticker", "ROE (%)", "Total Assets ($)", "P/E (x)", "Branches", "Price"],
            ["BANR", 11.24, 16_200_000_000, 10.2, 120, 61.5]])
        self.assertIsInstance(ws["C2"].value, int)
        self.assertEqual(ws["B2"].number_format, '0.00"%"')
        self.assertEqual(ws["C2"].number_format, '$#,##0')
        self.assertEqual(ws["D2"].number_format, '0.00"x"')
        self.assertEqual(ws["E2"].number_format, '#,##0')
        self.assertEqual(ws["F2"].number_format, '$#,##0.00')
        self.assertEqual(ws["A2"].number_format, "General")

    def test_pct_format_is_percent_units_not_native_percent(self):
        # App-wide, 11.24 means 11.24% (utils.formatting.format_value). Excel's
        # native 0.00% would display 1124.00% — a plausible-wrong number.
        self.assertEqual(ex.FORMATS["pct"], '0.00"%"')
        self.assertNotIn("0.00%", ex.FORMATS["pct"].replace('"%"', ""))

    def test_missing_renders_na_never_zero_or_blank(self):
        df = pd.DataFrame({
            "Metric": ["a", "b", "c", "d", "e", "f"],
            "Value": [None, float("nan"), "—", "", "n/a", 0.0],
        })
        wb = _load(ex.build_workbook(df, formats={"Value": "pct"}))
        ws = wb.worksheets[0]
        vals = [ws.cell(r, 2).value for r in range(2, 8)]
        self.assertEqual(vals, ["n/a", "n/a", "n/a", "n/a", "n/a", 0.0])
        # a genuine zero keeps its number format; n/a cells carry none
        self.assertEqual(ws.cell(7, 2).number_format, '0.00"%"')
        self.assertEqual(ws.cell(2, 2).number_format, "General")
        self.assertTrue(ws.cell(2, 2).font.italic)

    def test_infinity_is_na(self):
        df = pd.DataFrame({"x": [math.inf, -math.inf, 1.5]})
        ws = _load(ex.build_workbook(df)).worksheets[0]
        self.assertEqual([ws.cell(r, 1).value for r in (2, 3, 4)], ["n/a", "n/a", 1.5])

    def test_html_never_lands(self):
        df = pd.DataFrame({
            "Ticker": ['<a href="?bank=JPM" target="_self">JPM</a>', "BAC"],
            "Name": ["Acme &amp; Sons <b>Bank</b>", "Plain"],
        })
        ws = _load(ex.build_workbook(df)).worksheets[0]
        self.assertEqual(_grid(ws), [["Ticker", "Name"],
                                     ["JPM", "Acme & Sons Bank"],
                                     ["BAC", "Plain"]])

    def test_date_strings_parse_under_date_format_only(self):
        df = pd.DataFrame({"Filed": ["2026-07-17", "not a date"],
                           "Label": ["2026-07-17", "x"]})
        ws = _load(ex.build_workbook(df, formats={"Filed": "date"})).worksheets[0]
        self.assertEqual(ws["A2"].value, dt.datetime(2026, 7, 17))
        self.assertEqual(ws["A2"].number_format, "yyyy-mm-dd")
        self.assertEqual(ws["A3"].value, "not a date")
        self.assertEqual(ws["B2"].value, "2026-07-17")   # no format → text

    def test_tz_aware_timestamps_become_naive_utc(self):
        ts = pd.Timestamp("2026-07-17 13:30", tz="US/Eastern")
        df = pd.DataFrame({"filed_at": [ts]})
        ws = _load(ex.build_workbook(df)).worksheets[0]
        self.assertEqual(ws["A2"].value, dt.datetime(2026, 7, 17, 17, 30))

    def test_bools_and_nested_values(self):
        df = pd.DataFrame({"flag": [True, False], "blob": [{"a": 1}, [1, 2]]})
        ws = _load(ex.build_workbook(df)).worksheets[0]
        self.assertIs(ws["A2"].value, True)
        self.assertEqual(ws["B2"].value, "{'a': 1}")
        self.assertEqual(ws["B3"].value, "[1, 2]")

    def test_numpy_scalars_become_python(self):
        df = pd.DataFrame({"n": pd.array([1, 2], dtype="int64"),
                           "f": pd.array([0.5, None], dtype="Float64")})
        ws = _load(ex.build_workbook(df, formats={"f": "num"})).worksheets[0]
        self.assertEqual([ws["A2"].value, ws["B2"].value, ws["B3"].value], [1, 0.5, "n/a"])


class TestUnitsContract(unittest.TestCase):
    """D4: FDIC values stay in $thousands (owner choice 2026-09-22) — the unit
    must live in the header and on the Source sheet."""

    def test_usd_k_requires_unit_in_header(self):
        df = pd.DataFrame({"Deposits": [1_812_344]})
        with self.assertRaises(ValueError):
            ex.build_workbook(df, formats={"Deposits": "usd_k"})

    def test_usd_k_labeled_column_exports_unscaled(self):
        df = pd.DataFrame({"Deposits ($K)": [1_812_344, None]})
        wb = _load(ex.build_workbook(df, formats={"Deposits ($K)": "usd_k"}))
        ws = wb.worksheets[0]
        self.assertEqual(ws["A2"].value, 1_812_344)      # NOT ×1000
        self.assertEqual(ws["A2"].number_format, "#,##0")
        self.assertEqual(ws["A3"].value, "n/a")
        src = dict(_grid(wb["Source"]))
        self.assertIn("($K)", src["Units"])
        self.assertIn("thousands", src["Units"])

    def test_export_header_rewrites_scaled_suffix_to_raw_dollars(self):
        # D1: the screener label says ($B) but the cell holds whole dollars.
        self.assertEqual(ex.export_header("Mkt Cap ($B)"), "Mkt Cap ($)")
        self.assertEqual(ex.export_header("Net Inc ($M)"), "Net Inc ($)")
        self.assertEqual(ex.export_header("Dep < $250K ($B)"), "Dep < $250K ($)")
        self.assertEqual(ex.export_header("P/E"), "P/E")
        self.assertEqual(ex.export_header("NIM (%)"), "NIM (%)")

    def test_metric_columns_from_config(self):
        rename, formats = ex.metric_columns(
            ["ticker", "market_cap", "pe_ratio", "change_pct", "price", "volume"])
        self.assertEqual(rename["ticker"], "ticker")          # not a metric → passthrough
        self.assertEqual(rename["market_cap"], "Mkt Cap ($)")
        self.assertEqual(formats["Mkt Cap ($)"], "usd")
        self.assertEqual(formats[rename["pe_ratio"]], "x")
        self.assertEqual(formats[rename["change_pct"]], "pct")
        self.assertEqual(formats[rename["price"]], "usd2")
        self.assertEqual(formats[rename["volume"]], "int")

    def test_screener_row_exports_raw_dollars_under_honest_header(self):
        rename, formats = ex.metric_columns(["market_cap"])
        df = pd.DataFrame({"market_cap": [2_100_000_000.0]}).rename(columns=rename)
        ws = _load(ex.build_workbook(df, formats=formats)).worksheets[0]
        self.assertEqual(ws["A1"].value, "Mkt Cap ($)")
        self.assertEqual(ws["A2"].value, 2_100_000_000)
        self.assertEqual(ws["A2"].number_format, "$#,##0")


class TestRowFormats(unittest.TestCase):
    """Statements carry the unit per ROW (line item), not per column."""

    def _stmt(self):
        df = pd.DataFrame({
            "Line item": ["Net income", "ROA", "Diluted EPS", "Deposits ($K)", "Header only"],
            "FY2025": [1_250_000, 1.12, 3.41, 812_000, None],
            "FY2026": [1_400_000, None, 3.85, 845_500, None],
        })
        return ex.build_workbook(df, row_formats={
            "Net income": "usd", "ROA": "pct", "Diluted EPS": "usd2",
            "Deposits ($K)": "usd_k"})

    def test_row_label_drives_the_format_across_period_columns(self):
        ws = _load(self._stmt()).worksheets[0]
        self.assertEqual((ws["B2"].value, ws["B2"].number_format), (1_250_000, "$#,##0"))
        self.assertEqual((ws["C2"].value, ws["C2"].number_format), (1_400_000, "$#,##0"))
        self.assertEqual((ws["B3"].value, ws["B3"].number_format), (1.12, '0.00"%"'))
        self.assertEqual(ws["C3"].value, "n/a")
        self.assertEqual((ws["B4"].value, ws["B4"].number_format), (3.41, "$#,##0.00"))
        self.assertEqual((ws["B5"].value, ws["B5"].number_format), (812_000, "#,##0"))
        self.assertEqual(ws["A2"].number_format, "General")   # the label column itself
        self.assertEqual(ws["B6"].value, "n/a")

    def test_column_format_wins_over_row_format(self):
        df = pd.DataFrame({"Line item": ["ROA"], "FY2025": [1.12], "Rank": [3]})
        ws = _load(ex.build_workbook(df, formats={"Rank": "int"},
                                     row_formats={"ROA": "pct"})).worksheets[0]
        self.assertEqual(ws["B2"].number_format, '0.00"%"')
        self.assertEqual(ws["C2"].number_format, "#,##0")

    def test_usd_k_row_label_must_say_k(self):
        df = pd.DataFrame({"Line item": ["Deposits"], "FY2025": [1]})
        with self.assertRaises(ValueError):
            ex.build_workbook(df, row_formats={"Deposits": "usd_k"})

    def test_units_note_covers_row_formats(self):
        d = dict(_grid(_load(self._stmt())["Source"]))
        self.assertIn("whole US dollars", d["Units"])
        self.assertIn("percent units", d["Units"])
        self.assertIn("($K)", d["Units"])


class TestSheetStructure(unittest.TestCase):

    def test_sheet_title_legal_truncated_unique(self):
        used = set()
        t1 = ex.sheet_title("Valuation & Market Multiples Overview: all/banks?", used)
        self.assertLessEqual(len(t1), 31)
        for ch in "[]:*?/\\":
            self.assertNotIn(ch, t1)
        t2 = ex.sheet_title("Valuation & Market Multiples Overview: all/banks?", used)
        self.assertNotEqual(t1, t2)
        self.assertLessEqual(len(t2), 31)
        self.assertTrue(t2.endswith("(2)"))
        self.assertEqual(ex.sheet_title("", set()), "Data")

    def test_data_sheet_named_source_does_not_collide(self):
        wb = _load(ex.build_workbook(pd.DataFrame({"a": [1]}), sheet="Source"))
        self.assertEqual(wb.sheetnames, ["Source", "Source (2)"])

    def test_freeze_autofilter_and_widths(self):
        df = pd.DataFrame({"Ticker": ["A", "B"], "Bank": ["x", "y"],
                           "V1": [1, 2], "V2": [3, 4]})
        ws = _load(ex.build_workbook(df)).worksheets[0]
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:D3")
        ws2 = _load(ex.build_workbook(df, freeze_cols=2)).worksheets[0]
        self.assertEqual(ws2.freeze_panes, "C2")
        self.assertGreaterEqual(ws.column_dimensions["A"].width, 8)
        self.assertTrue(ws["A1"].font.bold)

    def test_empty_frame_still_builds(self):
        wb = _load(ex.build_workbook(pd.DataFrame(columns=["a", "b"])))
        self.assertEqual(_grid(wb.worksheets[0]), [["a", "b"]])
        self.assertEqual(dict(_grid(wb["Source"]))["Rows"], 0)

    def test_format_for_absent_column_is_tolerated(self):
        ws = _load(ex.build_workbook(pd.DataFrame({"a": [1]}),
                                     formats={"missing": "pct"})).worksheets[0]
        self.assertEqual(ws["A2"].value, 1)


class TestProvenance(unittest.TestCase):
    """D5: a file found on a desk a month later must say what it is."""

    def test_source_sheet_rows(self):
        wb = _load(ex.build_workbook(
            pd.DataFrame({"a": [1]}), sheet="Peer rank",
            formats={"a": "pct"},
            provenance={"Page": "Company Analysis › Peers › Peer rank",
                        "Ticker": "BANR", "FDIC cert": 28489,
                        "Source": "FDIC Summary of Deposits, 2025 survey",
                        "Data as of": "2025-06-30", "Skipped": None, "Blank": ""}))
        rows = _grid(wb["Source"])
        keys = [r[0] for r in rows]
        self.assertEqual(keys[0], "Exported")
        self.assertRegex(str(rows[0][1]), r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2} UTC$")
        self.assertEqual(keys[1:3], ["Table", "Rows"])
        d = dict(rows)
        self.assertEqual(d["Table"], "Peer rank")
        self.assertEqual(d["Rows"], 1)
        self.assertEqual(d["Ticker"], "BANR")
        self.assertEqual(d["FDIC cert"], 28489)
        self.assertEqual(d["Data as of"], "2025-06-30")
        self.assertNotIn("Skipped", d)
        self.assertNotIn("Blank", d)
        self.assertIn("percent units", d["Units"])
        self.assertIn("n/a", d["Missing values"])
        self.assertEqual(keys[-2:], ["Units", "Missing values"])


class _Ctx:
    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


class TestControl(unittest.TestCase):
    """The Export control contract (design-system decision #12): label
    "Export", one .xlsx, deferred build (data is a callable), keyed
    container for the compact right-aligned style."""

    def _capture(self):
        calls, containers = [], []
        fake = types.SimpleNamespace(
            download_button=lambda label, data, **k: calls.append((label, data, k)),
            container=lambda *a, **k: containers.append(k.get("key")) or _Ctx())
        saved = ex.st
        ex.st = fake
        return calls, containers, saved

    def test_table_export_contract(self):
        calls, containers, saved = self._capture()
        try:
            df = pd.DataFrame({"Ticker": ["BANR"], "NPL Ratio (%)": [0.42]})
            ex.table_export(df, "Valuation & Market Multiples Overview_all banks",
                            key="exp_x", formats={"NPL Ratio (%)": "pct"},
                            provenance={"Ticker": "BANR"})
        finally:
            ex.st = saved
        self.assertEqual(containers, ["tblexp_exp_x"])
        label, data, kw = calls[0]
        self.assertEqual(label, "Export")
        self.assertEqual(kw["file_name"],
                         "Valuation_Market_Multiples_Overview_all_banks.xlsx")
        self.assertEqual(kw["mime"], ex.XLSX_MIME)
        self.assertEqual(kw["key"], "exp_x")
        self.assertTrue(callable(data), "workbook must build lazily, on click")
        wb = _load(data())
        ws = wb.worksheets[0]
        self.assertEqual(ws.title, "Valuation_Market_Multiples_Over")
        self.assertEqual(_grid(ws), [["Ticker", "NPL Ratio (%)"], ["BANR", 0.42]])
        self.assertEqual(ws["B2"].number_format, '0.00"%"')
        self.assertEqual(dict(_grid(wb["Source"]))["Ticker"], "BANR")

    def test_safe_filename(self):
        self.assertEqual(ex.safe_filename("branch_details_cert28489"), "branch_details_cert28489")
        self.assertEqual(ex.safe_filename("a b&c//d"), "a_b_c_d")
        self.assertEqual(ex.safe_filename("..."), "export")


if __name__ == "__main__":
    unittest.main()
