"""Templated (FDIC Call Report) statement export — ui/financials_statements.

The statement cells are formatted to display strings at the point of
computation; the export must carry the RAW values those strings came from
(the _V tagged-string mechanism), never a parsed display string. Pins:

  * dollar lines export the FDIC $thousands UNSCALED under a "($K)" label
    with #,##0 (owner decision 2026-09-22);
  * ratio lines are percent units with the literal-% format;
  * an absent input is "n/a", never 0 — and a ratio whose numerator is
    absent is "n/a" too;
  * section headers survive as label-only rows; the screen HTML is unchanged
    (compact "$100.0M" still renders while the sheet holds 100000).
"""
import io
import types
import unittest
from unittest import mock

import pandas as pd
from openpyxl import load_workbook

import ui.export as ex
import ui.financials_statements as FS


class _Ctx:
    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def _stub_st(period="Annual"):
    st = types.SimpleNamespace()
    st.radio = lambda *a, **k: period
    st.markdown = st.caption = st.warning = st.info = lambda *a, **k: None
    st.spinner = lambda *a, **k: _Ctx()
    st.columns = lambda spec, **k: [_Ctx() for _ in (spec if isinstance(spec, list) else range(spec))]
    st.container = lambda *a, **k: _Ctx()
    return st


HIST = pd.DataFrame([
    {"REPDTE": "2024-12-31", "NETINC": 100_000, "INTINC": 500_000, "EINTEXP": 200_000, "ROA": 1.25},
    {"REPDTE": "2025-12-31", "NETINC": None,    "INTINC": 520_000, "EINTEXP": 210_000, "ROA": None},
])

SPEC = [("Income", [
    ("Net income", "dollar", "NETINC"),
    ("Net interest income", "diff", "INTINC", "EINTEXP"),
    ("ROA", "pct", "ROA"),
    ("NI / interest income", "ratio", "NETINC", "INTINC"),
])]


class TestTemplatedStatementExport(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        calls, html = [], []
        fake_ex = types.SimpleNamespace(
            download_button=lambda label, data, **k: calls.append((label, data, k)),
            container=lambda *a, **k: _Ctx())
        fake_components = types.SimpleNamespace(html=lambda h, **k: html.append(h))
        import data.loaders as loaders
        with mock.patch.object(FS, "st", _stub_st()), \
             mock.patch.object(FS, "components", fake_components), \
             mock.patch.object(FS, "get_bank_info",
                               lambda t: {"name": "Test Bancorp", "fdic_cert": 4242, "cik": None}), \
             mock.patch.object(FS, "_render_statement_trends", lambda *a, **k: None), \
             mock.patch.object(loaders, "load_fdic_hist_df", lambda t, q: HIST.copy()), \
             mock.patch.object(ex, "st", fake_ex):
            FS.render_statement("TBNK", "inc", "Income Statement", SPEC, trends=[])
        cls.calls, cls.html = calls, html
        assert len(calls) == 1, calls
        label, data, kw = calls[0]
        cls.label, cls.kw = label, kw
        cls.wb = load_workbook(io.BytesIO(data()))
        cls.ws = cls.wb.worksheets[0]

    def _rows(self):
        return [[c.value for c in r] for r in self.ws.iter_rows()]

    def test_control_contract(self):
        self.assertEqual(self.label, "Export")
        self.assertEqual(self.kw["key"], "exp_stmt_inc_TBNK")
        self.assertEqual(self.kw["file_name"], "TBNK_inc_annual_2025-12-31.xlsx")

    def test_header_and_periods(self):
        rows = self._rows()
        self.assertEqual(rows[0], ["Line item", "FY2024", "FY2025"])
        self.assertEqual(rows[1], ["Income", "n/a", "n/a"])          # section header survives
        self.assertEqual(self.ws.freeze_panes, "B2")

    def test_dollar_lines_are_unscaled_thousands_with_k_label(self):
        rows = self._rows()
        self.assertEqual(rows[2], ["Net income ($K)", 100_000, "n/a"])
        self.assertEqual(self.ws["B3"].number_format, "#,##0")
        self.assertEqual(rows[3], ["Net interest income ($K)", 300_000, 310_000])   # 500,000−200,000 / 520,000−210,000
        self.assertEqual(self.ws["C4"].number_format, "#,##0")

    def test_ratio_lines_are_percent_units(self):
        rows = self._rows()
        self.assertEqual(rows[4], ["ROA", 1.25, "n/a"])
        self.assertEqual(self.ws["B5"].number_format, '0.00"%"')
        # 100,000 ÷ 500,000 × 100 = 20.00; FY2025 numerator absent → n/a, never 0
        self.assertEqual(rows[5][0], "NI / interest income")
        self.assertAlmostEqual(rows[5][1], 20.0)
        self.assertEqual(rows[5][2], "n/a")
        self.assertEqual(self.ws["B6"].number_format, '0.00"%"')

    def test_screen_html_unchanged(self):
        h = self.html[0]
        self.assertIn(">$100.0M<", h)        # compact display string still renders
        self.assertIn(">20.00%<", h)
        self.assertNotIn("100000", h)        # raw never leaks to the screen

    def test_source_sheet(self):
        d = dict([c.value for c in r] for r in self.wb["Source"].iter_rows())
        self.assertEqual(d["Page"], "Company Analysis › Financials › Templated › Income Statement")
        self.assertEqual(d["Period"], "Annual")
        self.assertEqual(d["Ticker"], "TBNK")
        self.assertEqual(d["FDIC cert"], 4242)
        self.assertEqual(d["Latest report date"], "2025-12-31")
        self.assertIn("FDIC Call Report", d["Source"])
        self.assertNotIn("SEC CIK", d)       # with_persh=False → no SEC row
        self.assertIn("($K)", d["Units"])


class TestTaggedStrings(unittest.TestCase):
    """_V behaves as str on screen and carries raw+unit for the export."""

    def test_usd_carries_thousands(self):
        v = FS._usd(1_234_567)
        self.assertEqual(str(v), "$1.23B")
        self.assertEqual((v.raw, v.unit), (1_234_567.0, "usd_k"))
        self.assertEqual(FS._usd(None), "—")
        self.assertIsNone(FS._usd(None).raw)

    def test_pct_and_inline_pct(self):
        self.assertEqual((FS._pct(3.4567).raw, FS._pct(3.4567).unit, str(FS._pct(3.4567))),
                         (3.4567, "pct", "3.46%"))
        self.assertEqual((str(FS._pctv(12.345, 1)), FS._pctv(12.345, 1).raw), ("12.3%", 12.345))

    def test_per_share(self):
        v = FS._psd(23.349)
        self.assertEqual((str(v), v.raw, v.unit), ("$23.35", 23.349, "usd2"))
        self.assertEqual(FS._psd(None), "—")

    def test_json_and_equality_behave_like_str(self):
        import json
        self.assertEqual(json.dumps({"v": FS._usd(1000)}), '{"v": "$1.0M"}')
        self.assertTrue(FS._usd(None) == "—")


if __name__ == "__main__":
    unittest.main()
