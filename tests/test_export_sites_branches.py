"""SOD-backed table exports on the shared .xlsx exporter (ui/export.py).

Five pages export FDIC Summary-of-Deposits frames whose deposit values are
$THOUSANDS at the source. Owner decision 2026-09-22: they export UNSCALED
with "($K)" in the header (never ×1000 — never guess units), shares in
percent units under the literal-% format, absent values as "n/a", and a
Source sheet naming the SOD survey year.

Render-driven (the real page helpers, streamlit stubbed, download_button
captured) for geo_view._ranking_table, deposit_market_share.
_render_market_table and deposit_lookup._render_market_share; the merger
and proximity sites are pinned by their column maps against the upstream
column lists (analysis.merger_hhi.PRO_FORMA_COLS / branches_store.
_COMPETITOR_PAIR_COLS) plus a hand-computed workbook build — so a new
upstream column can never export under a raw field code.

Run: python -m unittest tests.test_export_sites_branches
"""
from __future__ import annotations

import io
import sys
import unittest

import pandas as pd
from openpyxl import load_workbook

from tests._streamlit_stub import install

install()

import ui.export as ex  # noqa: E402

_MISSING = object()
_STUB_NAMES = ("markdown", "caption", "info", "warning", "empty",
               "container", "spinner", "download_button")

# Page modules these tests import lazily. Sibling suites that follow under
# discovery (tests/test_export_sites_earnings, _misc) REPLACE
# sys.modules["streamlit"] via test_render_smoke._install_streamlit_stub, and
# tests/test_merger_proximity_ui records widgets on the module current at ITS
# run — so a page module first imported HERE would stay bound to the earlier
# stub and its st.info/st.markdown calls would never reach those recorders
# (9 phantom failures, 2026-09-22). Leave the process as we found it: pop
# what we introduced, so later suites import fresh under their own stub.
_PAGE_MODULES = ("ui.geo_view", "ui.deposit_market_share", "ui.deposit_lookup",
                 "ui.merger_planning", "ui.branch_proximity")
_introduced: list[str] = []


def setUpModule():
    _introduced[:] = [m for m in _PAGE_MODULES if m not in sys.modules]


def tearDownModule():
    for m in _introduced:
        sys.modules.pop(m, None)


class _Ctx:
    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def _grid(ws):
    return [[c.value for c in r] for r in ws.iter_rows()]


class _ExportCase(unittest.TestCase):
    """Shared stub extension (the documented tests/_streamlit_stub pattern)
    + download_button capture on the st module ui.export is bound to."""

    def setUp(self):
        self.st = install()
        self._saved: list[tuple[object, str, object]] = []
        self._complete(self.st)
        self.calls: list = []
        self._saved_dl = ex.st.download_button
        ex.st.download_button = (
            lambda label, data, **k: self.calls.append((label, data, k)))

    def _complete(self, st):
        """No-op the widgets a render path touches on `st`, remembering
        what was there (or absent) for tearDown."""
        for n in _STUB_NAMES:
            self._saved.append((st, n, getattr(st, n, _MISSING)))
            setattr(st, n, (lambda *a, **k: _Ctx()) if n in ("container", "spinner")
                    else (lambda *a, **k: None))

    def _bind(self, *mods):
        """A module imported earlier in the process may be bound to a
        different streamlit stub object than install() returns; complete
        that one too (ui.components / tables / states / chrome render the
        headers and tables these sites draw)."""
        import ui.chrome, ui.components, ui.states, ui.tables  # noqa: E401
        seen = {id(self.st)}
        for m in mods + (ui.chrome, ui.components, ui.states, ui.tables):
            if id(m.st) not in seen:
                seen.add(id(m.st))
                self._complete(m.st)

    def tearDown(self):
        ex.st.download_button = self._saved_dl
        for st, n, v in reversed(self._saved):
            if v is _MISSING:
                if hasattr(st, n):
                    delattr(st, n)
            else:
                setattr(st, n, v)

    def _workbook(self):
        self.assertEqual(len(self.calls), 1, "expected exactly one Export")
        label, data, kw = self.calls[0]
        self.assertEqual(label, "Export")
        self.assertTrue(callable(data), "workbook must build lazily")
        wb = load_workbook(io.BytesIO(data()))
        return wb, wb.worksheets[0], dict(_grid(wb["Source"])), kw


class TestGeoRankingExport(_ExportCase):
    """geo_view._ranking_table: market ranking (state / MSA / county tabs).
    Fixture total = 43,400,000 + 6,600,000 = 50,000,000 $K →
    JPM 86.80%, private 13.20%; the bank with no deposits is n/a, never 0%."""

    def test_ranking_export_unscaled_k_and_percent_units(self):
        import ui.geo_view as gv
        self._bind(gv)
        banks = pd.DataFrame({
            "owner_key": ["JPM", "c999", "c555"],
            "ticker": ["JPM", None, None],
            "cert": [628, 999, 555],
            "bank_name": ["JPMorgan Chase Bank", "Private Bank", "No-Deposit Bank"],
            "n_branches": [120, 3, 1],
            "total_deposits": [43_400_000, 6_600_000, None],
        })
        gv._ranking_table(banks, "CA", "banks_by_state_CA",
                          "exp_banks_by_state_CA", tab="By State", year=2025)
        wb, ws, src, kw = self._workbook()
        self.assertEqual(kw["file_name"], "banks_by_state_CA_2025.xlsx")
        self.assertEqual(kw["key"], "exp_banks_by_state_CA")      # unchanged
        rows = _grid(ws)
        self.assertEqual(rows[0], ["Rank", "Ticker", "Bank", "FDIC cert",
                                   "Branches", "Deposits ($K)", "Market share (%)"])
        self.assertNotIn("owner_key", rows[0])
        self.assertEqual(rows[1][:6], [1, "JPM", "JPMorgan Chase Bank", 628, 120,
                                       43_400_000])
        self.assertNotEqual(ws["F2"].value, 43_400_000_000)        # never ×1000
        self.assertAlmostEqual(rows[1][6], 86.8)
        self.assertAlmostEqual(rows[2][6], 13.2)
        self.assertEqual(ws["F2"].number_format, "#,##0")
        self.assertEqual(ws["G2"].number_format, '0.00"%"')
        self.assertEqual(ws["D2"].number_format, "#,##0")
        self.assertEqual((ws["F4"].value, ws["G4"].value), ("n/a", "n/a"))
        self.assertEqual(ws["B3"].value, "n/a")                    # private: no ticker
        self.assertEqual(src["SOD survey year"], 2025)
        self.assertEqual(src["Market"], "CA")
        self.assertEqual(src["Page"], "Geographic › By State")
        self.assertIn("Summary of Deposits", src["Source"])
        self.assertIn("thousands", src["Units"])
        self.assertIn("($K)", src["Units"])
        self.assertIn("percent units", src["Units"])

    def test_year_absent_keeps_bare_stem(self):
        import ui.geo_view as gv
        self._bind(gv)
        banks = pd.DataFrame({"owner_key": ["A"], "ticker": ["A"], "cert": [1],
                              "bank_name": ["A"], "n_branches": [1],
                              "total_deposits": [10]})
        gv._ranking_table(banks, "X", "banks_by_msa_1", "exp_banks_by_msa_1")
        _wb, _ws, src, kw = self._workbook()
        self.assertEqual(kw["file_name"], "banks_by_msa_1.xlsx")
        self.assertNotIn("SOD survey year", src)      # never invented


class TestDepositMarketShareExport(_ExportCase):
    """deposit_market_share._render_market_table over _share_rows output.
    Market A: subject 600,000 / rival 400,000 $K → total 1,000,000, share
    60%, HHI 60² + 40² = 5,200, rank 1 of 2, top competitor Rival 40%.
    Market B: subject alone 50,000 → HHI 10,000, no competitor (n/a)."""

    def test_market_table_export(self):
        import data.bank_universe as bu
        import ui.deposit_market_share as dms
        self._bind(dms)
        df = pd.DataFrame([
            ("06001", "Alameda, CA", 111, "Subject Bank", "SUBJ", 3, 600_000),
            ("06001", "Alameda, CA", 222, "Rival Bank", "RVL", 5, 400_000),
            ("06013", "Contra Costa, CA", 111, "Subject Bank", "SUBJ", 1, 50_000),
        ], columns=["market_key", "market_label", "cert", "bank_name",
                    "ticker", "n_branches", "deposits"])
        rows = dms._share_rows(df, subject_cert=111)
        saved = bu.cert_ticker_map
        bu.cert_ticker_map = lambda: {222: "RVL"}
        try:
            dms._render_market_table(rows, "By County", "county_SUBJ",
                                     ticker="SUBJ", cert=111, year=2025)
        finally:
            bu.cert_ticker_map = saved
        wb, ws, src, kw = self._workbook()
        self.assertEqual(kw["file_name"], "deposit_share_county_SUBJ_2025.xlsx")
        self.assertEqual(kw["key"], "exp_depshare_county_SUBJ")     # unchanged
        rows_x = _grid(ws)
        self.assertEqual(rows_x[0], [
            "Market key", "Market", "Branches", "Deposits ($K)",
            "Market total ($K)", "Share (%)", "Rank", "Banks in market", "HHI",
            "Top competitor", "Top competitor share (%)",
            "Top competitor FDIC cert"])
        self.assertEqual(rows_x[1], ["06001", "Alameda, CA", 3, 600_000,
                                     1_000_000, 60.0, 1, 2, 5200.0,
                                     "Rival Bank", 40.0, 222])
        self.assertEqual(ws["D2"].number_format, "#,##0")     # usd_k, unscaled
        self.assertEqual(ws["E2"].number_format, "#,##0")
        self.assertEqual(ws["F2"].number_format, '0.00"%"')
        self.assertEqual(ws["I2"].number_format, "#,##0")     # HHI
        self.assertEqual(ws["K2"].number_format, '0.00"%"')
        self.assertEqual(rows_x[2][:4], ["06013", "Contra Costa, CA", 1, 50_000])
        self.assertEqual(rows_x[2][8], 10_000.0)
        self.assertEqual(rows_x[2][9:], ["n/a", "n/a", "n/a"])
        self.assertEqual(src["Ticker"], "SUBJ")
        self.assertEqual(src["FDIC cert"], 111)
        self.assertEqual(src["SOD survey year"], 2025)
        self.assertEqual(src["Page"], "Deposit Market Share › By County")
        self.assertIn("thousands", src["Units"])

    def test_every_format_header_is_a_mapped_label(self):
        import ui.deposit_market_share as dms
        self.assertTrue(set(dms._EXPORT_FORMATS) <= set(dms._EXPORT_COLS.values()))
        for h, f in dms._EXPORT_FORMATS.items():
            if f == "usd_k":
                self.assertIn("($K)", h)


class TestDepositLookupMarketShareExport(_ExportCase):
    """deposit_lookup._render_market_share (Market Share & Branches ›
    Deposit market share). market_share is ALREADY ×100 in _market_share:
    1,812,344 / 2,500,000 × 100 = 72.49376%."""

    def test_market_share_export(self):
        import ui.deposit_lookup as dl
        self._bind(dl)
        ms = pd.DataFrame({
            "owner_key": ["BANR", "c777"],
            "CERT": [28489, 777],
            "TICKER": ["BANR", None],
            "NAMEFULL": ["Banner Bank", "Local Savings"],
            "branches": [12, 2],
            "deposits": [1_812_344, 687_656],
            "market_share": [1_812_344 / 2_500_000 * 100, 687_656 / 2_500_000 * 100],
            "rank": [1, 2],
        })
        saved = dl._market_share
        dl._market_share = lambda kind, key, year: ms
        try:
            dl._render_market_share("county", "53063", "Spokane County, WA",
                                    2025, "BANR", "Banner Corp", lambda v: str(v))
        finally:
            dl._market_share = saved
        wb, ws, src, kw = self._workbook()
        self.assertEqual(kw["file_name"], "county_market_share_53063_2025.xlsx")
        self.assertEqual(kw["key"], "exp_county_market_share_53063")  # unchanged
        rows = _grid(ws)
        self.assertEqual(rows[0], ["Rank", "Ticker", "Bank", "FDIC cert",
                                   "Branches", "Deposits ($K)", "Market share (%)"])
        self.assertEqual(rows[1][:6], [1, "BANR", "Banner Bank", 28489, 12, 1_812_344])
        self.assertAlmostEqual(rows[1][6], 72.49376)
        self.assertAlmostEqual(rows[2][6], 27.50624)
        self.assertEqual(ws["F2"].number_format, "#,##0")
        self.assertEqual(ws["G2"].number_format, '0.00"%"')
        self.assertEqual(ws["B3"].value, "n/a")
        self.assertEqual(src["Company"], "Banner Corp")
        self.assertEqual(src["Market"], "Spokane County, WA")
        self.assertEqual(src["SOD survey year"], 2025)
        self.assertIn("thousands", src["Units"])


class TestMergerPlanningExportMap(unittest.TestCase):
    """merger_planning: every PRO_FORMA_COLS column has a readable header;
    hand-computed row (tests/test_merger_hhi fixture): A=400, B=300,
    C=200, D=100 → total 1,000; A 40%, B 30%, combined 70%; HHI pre
    40²+30²+20²+10² = 3,000, post 70²+20²+10² = 5,400, Δ 2,400, flagged."""

    def test_columns_cover_pro_forma_cols(self):
        from analysis.merger_hhi import PRO_FORMA_COLS
        import ui.merger_planning as mp
        self.assertEqual(set(mp._EXPORT_COLS), set(PRO_FORMA_COLS))
        self.assertTrue(set(mp._EXPORT_FORMATS) <= set(mp._EXPORT_COLS.values()))

    def test_hand_computed_row_builds(self):
        import ui.merger_planning as mp
        row = {"market_key": "11111", "market_label": "County 11111, PA",
               "n_banks": 4, "market_total": 1000,
               "branches_a": 1, "deposits_a": 400, "share_a_pct": 40.0,
               "branches_b": 1, "deposits_b": 300, "share_b_pct": 30.0,
               "combined_share_pct": 70.0, "hhi_pre": 3000.0,
               "hhi_post": 5400.0, "hhi_delta": 2400.0,
               "concentration_post": "highly concentrated",
               "screen_flag": True,
               "screen_reason": "post-merger HHI > 1,800 with ΔHHI > 200"}
        df = pd.DataFrame([row]).rename(columns=mp._EXPORT_COLS)
        wb = load_workbook(io.BytesIO(ex.build_workbook(
            df, formats=mp._EXPORT_FORMATS, provenance={"SOD survey year": 2024})))
        ws = wb.worksheets[0]
        hdr, vals = _grid(ws)
        d = dict(zip(hdr, vals))
        self.assertEqual(d["A deposits ($K)"], 400)          # unscaled thousands
        self.assertEqual(d["Market total ($K)"], 1000)
        self.assertEqual(d["A share (%)"], 40.0)
        self.assertEqual(d["Combined share (%)"], 70.0)
        self.assertEqual((d["HHI pre"], d["HHI post"], d["ΔHHI"]), (3000.0, 5400.0, 2400.0))
        self.assertIs(d["DOJ screen flagged"], True)
        col = {h: i + 1 for i, h in enumerate(hdr)}
        self.assertEqual(ws.cell(2, col["A deposits ($K)"]).number_format, "#,##0")
        self.assertEqual(ws.cell(2, col["A share (%)"]).number_format, '0.00"%"')
        self.assertEqual(ws.cell(2, col["HHI post"]).number_format, "#,##0")
        src = dict(_grid(wb["Source"]))
        self.assertEqual(src["SOD survey year"], 2024)
        self.assertIn("thousands", src["Units"])


class TestBranchProximityExportMap(unittest.TestCase):
    """branch_proximity: every competitor-pair column is labeled; both
    deposit columns are ($K) usd_k; distance is plain miles."""

    def test_columns_cover_pair_cols(self):
        import data.branches_store as bs
        import ui.branch_proximity as bp
        self.assertEqual(set(bp._EXPORT_COLS), set(bs._COMPETITOR_PAIR_COLS))
        self.assertTrue(set(bp._EXPORT_FORMATS) <= set(bp._EXPORT_COLS.values()))

    def test_pair_row_builds_unscaled(self):
        import data.branches_store as bs
        import ui.branch_proximity as bp
        pair = dict(zip(bs._COMPETITOR_PAIR_COLS, [
            1, "Main Office", "1 Main St", "Media", "PA", 40.0, -75.0, 500,
            2, 1, "CMP", "CompA", "CompA Br", "2 Elm St", "Media", "PA",
            "19063", 300, 40.1, -75.0, "12", 6.9093]))
        df = pd.DataFrame([pair]).rename(columns=bp._EXPORT_COLS)
        wb = load_workbook(io.BytesIO(ex.build_workbook(
            df, formats=bp._EXPORT_FORMATS,
            provenance={"Radius (mi)": 10, "SOD survey year": 2024})))
        ws = wb.worksheets[0]
        hdr, vals = _grid(ws)
        d = dict(zip(hdr, vals))
        col = {h: i + 1 for i, h in enumerate(hdr)}
        self.assertEqual(d["Subject deposits ($K)"], 500)
        self.assertEqual(d["Competitor deposits ($K)"], 300)
        self.assertEqual(ws.cell(2, col["Competitor deposits ($K)"]).number_format, "#,##0")
        self.assertAlmostEqual(d["Distance (mi)"], 6.9093)
        self.assertEqual(ws.cell(2, col["Distance (mi)"]).number_format, "#,##0.00")
        self.assertEqual(d["Competitor ZIP"], "19063")          # string kept
        self.assertEqual(d["Competitor latitude"], 40.1)
        self.assertEqual(ws.cell(2, col["Competitor latitude"]).number_format, "General")
        src = dict(_grid(wb["Source"]))
        self.assertEqual(src["Radius (mi)"], 10)
        self.assertEqual(src["SOD survey year"], 2024)


if __name__ == "__main__":
    unittest.main()
