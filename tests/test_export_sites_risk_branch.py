"""Table exports on ui/rate_sensitivity.py, ui/branch_analytics.py and
ui/hmda_view.py (shared exporter ui/export.py; owner directive 2026-09-22:
every data table gets an Export built from the RAW frame, never from display
strings).

Each test drives the real render function against small fixture data with the
data seams patched, captures the deferred workbook callable table_export hands
st.download_button, loads the bytes with openpyxl and asserts HAND-COMPUTED
values, exact Excel number formats, "n/a" for absent values and the Source
sheet rows. Units pinned here:

  * rate-sensitivity ΔNII is WHOLE dollars (earning_assets_usd = FDIC ERNAST
    $K ×1000 in analysis.rate_sensitivity.build_rate_sensitivity_inputs),
    ΔNIM is basis points, ΔEPS is $/share — never the $K/$M/$B display scale;
  * SOD deposits (branch list / competitors / demographics) are FDIC
    $thousands as stored and export UNSCALED under a "($K)" header;
  * Census dollars are whole dollars, unemployment is a percent;
  * HMDA volume_usd is raw dollars (data/hmda_client units contract).

Stub discipline: each module's own ``st`` binding is patched with a PRIVATE
stub via mock.patch.object — nothing here touches sys.modules or the shared
package stub, so the module composes with every other suite under discovery.

Run: PYTHONIOENCODING=utf-8 python -m unittest tests.test_export_sites_risk_branch
"""
from __future__ import annotations

import io
import sys
import types
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).parent.parent))

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _noop(*a, **k):
    return None


class _Ctx:
    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def __getattr__(self, name):
        return _noop


class _StubSt:
    """Private streamlit stand-in: unknown widgets are no-ops, layout widgets
    return contexts, input widgets return their default (first option /
    value=)."""

    session_state: dict = {}
    query_params: dict = {}

    def __getattr__(self, name):
        return _noop

    def columns(self, spec, **k):
        n = spec if isinstance(spec, int) else len(spec)
        return [_Ctx() for _ in range(n)]

    def container(self, *a, **k):
        return _Ctx()

    expander = spinner = popover = empty = container

    def selectbox(self, label, options=None, **k):
        return options[0] if options else None

    radio = selectbox

    def slider(self, *a, **k):
        return k.get("value", 0)

    def select_slider(self, label, options=None, value=None, **k):
        return value

    def checkbox(self, *a, **k):
        return bool(k.get("value", False))


class _ExportSite(unittest.TestCase):
    _st_modules: tuple = ()

    @classmethod
    def setUpClass(cls):
        import ui.chrome, ui.export, ui.states, ui.tables  # noqa: E401
        cls.X = ui.export
        cls._st_modules = (ui.chrome, ui.states, ui.tables) + tuple(
            __import__(m, fromlist=["_"]) for m in cls.PAGE_MODULES)

    def setUp(self):
        self.calls = []          # [(label, data, kwargs)] from download_button
        stub = _StubSt()
        fake_export_st = types.SimpleNamespace(
            download_button=lambda label, data, **k: self.calls.append((label, data, k)),
            container=lambda *a, **k: _Ctx())
        for mod, obj in [(m, stub) for m in self._st_modules] + [(self.X, fake_export_st)]:
            p = mock.patch.object(mod, "st", obj)
            p.start()
            self.addCleanup(p.stop)

    def _patch(self, target, **kw):
        p = mock.patch(target, **kw)
        p.start()
        self.addCleanup(p.stop)

    def _book(self, n_expected=1, idx=0):
        self.assertEqual(len(self.calls), n_expected,
                         f"expected {n_expected} Export control(s)")
        label, data, kw = self.calls[idx]
        self.assertEqual(label, "Export")
        self.assertTrue(callable(data), "workbook must build lazily, on click")
        wb = load_workbook(io.BytesIO(data()))
        return wb, wb.worksheets[0], self._source(wb), kw

    @staticmethod
    def _grid(ws):
        return [[c.value for c in r] for r in ws.iter_rows()]

    @staticmethod
    def _source(wb):
        return dict((r[0].value, r[1].value) for r in wb["Source"].iter_rows())

    def _assert_na(self, ws, coord):
        self.assertEqual(ws[coord].value, "n/a")
        self.assertTrue(ws[coord].font.italic)


# ── Rate sensitivity ─────────────────────────────────────────────────────

_LATEST = {"REPDTE": "20250630", "NIMY": 3.20, "ERNAST": 10_000_000}   # $K


class TestRateSensitivityExports(_ExportSite):
    PAGE_MODULES = ("ui.rate_sensitivity",)

    def setUp(self):
        super().setUp()
        import ui.rate_sensitivity as rs
        self.rs = rs
        self._patch("ui.rate_sensitivity.get_fdic_cert", return_value=12345)
        self._patch("ui.rate_sensitivity.get_name", return_value="Test Bancorp")

    def test_phased_table_exports_bps_dollars_and_eps(self):
        """EPS = ΔNII × (1 − 21%) / 50,000,000 shares:
        −1,250,000 × 0.79 / 5e7 = −0.01975. Second scenario has no share
        count → ΔEPS n/a, never 0."""
        result = {
            "inputs": {"current_nim_pct": 3.2, "earning_assets_usd": 1e10},
            "beta_used": 0.35, "beta_mode": "textbook", "tax_rate_used": 0.21,
            "shares_outstanding": 50_000_000, "ladder_source": "generic",
            "repricing_pace": {1: 0.40, 2: 0.70},
            "scenarios": [
                {"rate_change_bps": -100, "years": [
                    {"year": 1, "nim_delta_bps": -12.5, "nii_delta_usd": -1_250_000.0,
                     "eps_delta": -1_250_000.0 * 0.79 / 50_000_000},
                    {"year": 2, "nim_delta_bps": -20.0, "nii_delta_usd": -2_000_000.0,
                     "eps_delta": -2_000_000.0 * 0.79 / 50_000_000}]},
                {"rate_change_bps": 100, "years": [
                    {"year": 1, "nim_delta_bps": 15.0, "nii_delta_usd": 1_500_000.0,
                     "eps_delta": None},
                    {"year": 2, "nim_delta_bps": 24.0, "nii_delta_usd": 2_400_000.0,
                     "eps_delta": None}]},
            ],
        }
        self._patch("ui.rate_sensitivity.run_rate_sensitivity_phased", return_value=result)
        self._patch("ui.rate_sensitivity._render_phased_inputs")
        self._patch("ui.rate_sensitivity._render_assumptions_panel",
                    return_value=(None, None))
        self._patch("data.call_report_store.get_latest_ladder", return_value=None)
        self._patch("data.bank_mapping.get_cik", return_value=None)
        self._patch("data.nim_assumptions_store.get_assumptions", return_value=None)

        self.rs._render_phased_scenarios("TST", _LATEST, [_LATEST], "textbook", None)

        wb, ws, src, kw = self._book()
        self.assertEqual(kw["file_name"], "nim_phased_TST_2025-06-30.xlsx")
        self.assertEqual(kw["key"], "exp_nim_phased_TST")
        rows = self._grid(ws)
        self.assertEqual(rows[0], ["Rate shock (bp)",
                                   "Yr1 ΔNIM (bp)", "Yr1 ΔNII ($)", "Yr1 ΔEPS ($)",
                                   "Yr2 ΔNIM (bp)", "Yr2 ΔNII ($)", "Yr2 ΔEPS ($)"])
        self.assertEqual(rows[1][:3], [-100, -12.5, -1_250_000.0])
        self.assertAlmostEqual(rows[1][3], -0.01975)
        self.assertEqual(rows[1][4:6], [-20.0, -2_000_000.0])
        self.assertEqual(rows[2][:3], [100, 15.0, 1_500_000.0])
        self._assert_na(ws, "D3")
        self._assert_na(ws, "G3")
        self.assertEqual(ws["A2"].number_format, "#,##0")
        self.assertEqual(ws["B2"].number_format, "#,##0.00")
        self.assertEqual(ws["C2"].number_format, "$#,##0")      # whole dollars, unscaled
        self.assertEqual(ws["D2"].number_format, "$#,##0.00")
        self.assertEqual(src["Page"], "Interest Rate Risk › Multi-Year Impact (phased)")
        self.assertEqual(src["Ticker"], "TST")
        self.assertEqual(src["Company"], "Test Bancorp")
        self.assertEqual(src["FDIC cert"], 12345)
        self.assertEqual(src["Report date"], "2025-06-30")
        self.assertIn("model output", src["Source"])
        self.assertIn("NIMY", src["Source"])
        self.assertEqual(src["Rate shocks (bp)"], "-100, +100")
        self.assertEqual(src["Horizon (years)"], 1)              # selectbox default
        self.assertEqual(src["Floating-rate loan share"], "30%")  # generic default
        self.assertEqual(src["Repricing pace (cumulative)"], "Yr1=40%, Yr2=70%")
        self.assertEqual(src["Deposit beta mode"], "textbook")
        self.assertEqual(src["Effective tax rate"], "21%")
        self.assertEqual(src["Shares outstanding (SEC, TTM)"], 50_000_000)
        self.assertIn("generic industry average", src["Securities repricing source"])
        self.assertIn("whole US dollars", src["Units"])

    def test_named_scenarios_export(self):
        """ΔNII = $10B earning assets × ΔNIM(pp)/100: +25 bp → +25,000,000;
        −15 bp → −15,000,000."""
        result = {
            "inputs": {"current_nim_pct": 3.2, "earning_assets_usd": 1e10},
            "beta_used": 0.5, "beta_mode": "textbook", "asset_beta": 1.0,
            "scenarios": [
                {"name": "Parallel +100", "short_change_bps": 100, "long_change_bps": 100,
                 "nim_new_pct": 3.45, "nim_delta_bps": 25.0, "nii_delta_usd": 25_000_000.0,
                 "description": "All rates up 100bps (unchanged curve shape)."},
                {"name": "Bull Steepener", "short_change_bps": -100, "long_change_bps": -25,
                 "nim_new_pct": 3.05, "nim_delta_bps": -15.0, "nii_delta_usd": -15_000_000.0,
                 "description": "Short rates fall faster than long."},
            ],
        }
        self._patch("ui.rate_sensitivity.run_curve_sensitivity", return_value=result)

        self.rs._render_named_scenarios("TST", _LATEST, [_LATEST], "textbook", None, 1.0)

        wb, ws, src, kw = self._book()
        self.assertEqual(kw["file_name"], "nim_scenarios_TST_2025-06-30.xlsx")
        self.assertEqual(kw["key"], "exp_nim_scenarios_TST")
        rows = self._grid(ws)
        self.assertEqual(rows[0], ["Scenario", "Δ3M (bp)", "Δ5Y (bp)", "New NIM (%)",
                                   "ΔNIM (bp)", "ΔNII annual ($)", "Description"])
        self.assertEqual(rows[1], ["Parallel +100", 100, 100, 3.45, 25.0, 25_000_000.0,
                                   "All rates up 100bps (unchanged curve shape)."])
        self.assertEqual(rows[2][1:6], [-100, -25, 3.05, -15.0, -15_000_000.0])
        self.assertEqual(ws["B2"].number_format, "#,##0")
        self.assertEqual(ws["D2"].number_format, '0.00"%"')
        self.assertEqual(ws["E2"].number_format, "#,##0.00")
        self.assertEqual(ws["F2"].number_format, "$#,##0")
        self.assertEqual(src["Page"], "Interest Rate Risk › Named Curve Scenarios")
        self.assertEqual(src["Current NIM (%)"], 3.2)
        self.assertEqual(src["Earning assets ($)"], 1e10)
        self.assertEqual(src["Deposit beta (interest-bearing)"], 0.5)
        self.assertEqual(src["Asset beta (5Y pass-through to yields)"], 1.0)
        self.assertEqual(src["Report date"], "2025-06-30")
        self.assertIn("percent units", src["Units"])

    def test_curve_matrix_exports_two_grids(self):
        """Fixture rule: ΔNIM(bp) = (Δ5Y − Δ3M) × 0.25; ΔNII = ΔNIM(bp) × $1M
        ($10B × bp/10,000). Corner (Δ3M −100, Δ5Y +100) → +50 bp, +$50,000,000."""
        short = [-100, -50, 0, 50, 100]
        long_ = [-100, -50, 0, 50, 100]
        nim = [[(l - s) * 0.25 for l in long_] for s in short]
        nii = [[v * 1_000_000 for v in row] for row in nim]
        matrix = {"inputs": {"current_nim_pct": 3.2, "earning_assets_usd": 1e10},
                  "beta_used": 0.5, "beta_mode": "textbook", "asset_beta": 1.0,
                  "short_bps_range": short, "long_bps_range": long_,
                  "nim_delta_matrix_bps": nim, "nii_delta_matrix_usd": nii}
        self._patch("ui.rate_sensitivity.run_curve_matrix", return_value=matrix)

        self.rs._render_curve_matrix("TST", _LATEST, [_LATEST], "textbook", None, 1.0)

        wb, ws, src, kw = self._book(n_expected=2, idx=0)
        self.assertEqual(kw["file_name"], "nim_matrix_TST_2025-06-30.xlsx")
        self.assertEqual(kw["key"], "exp_nim_matrix_TST")
        rows = self._grid(ws)
        self.assertEqual(rows[0], ["Δ3M (bp)", "Δ5Y -100bp ΔNIM (bp)", "Δ5Y -50bp ΔNIM (bp)",
                                   "Δ5Y +0bp ΔNIM (bp)", "Δ5Y +50bp ΔNIM (bp)",
                                   "Δ5Y +100bp ΔNIM (bp)"])
        self.assertEqual(rows[1], [-100, 0.0, 12.5, 25.0, 37.5, 50.0])
        self.assertEqual(rows[5], [100, -50.0, -37.5, -25.0, -12.5, 0.0])
        self.assertEqual(ws["A2"].number_format, "#,##0")
        self.assertEqual(ws["F2"].number_format, "#,##0.00")
        self.assertEqual(src["Page"], "Interest Rate Risk › Curve Matrix (3M × 5Y) › ΔNIM")
        self.assertEqual(src["3M shocks (bp)"], "-100, -50, +0, +50, +100")
        self.assertEqual(src["5Y shocks (bp)"], "-100, -50, +0, +50, +100")

        wb2, ws2, src2, kw2 = self._book(n_expected=2, idx=1)
        self.assertEqual(kw2["file_name"], "nii_matrix_TST_2025-06-30.xlsx")
        self.assertEqual(kw2["key"], "exp_nii_matrix_TST")
        rows2 = self._grid(ws2)
        self.assertEqual(rows2[0][:2], ["Δ3M (bp)", "Δ5Y -100bp ΔNII annual ($)"])
        self.assertEqual(rows2[1][5], 50_000_000.0)          # whole dollars, not $M
        self.assertEqual(rows2[5][1], -50_000_000.0)
        self.assertEqual(ws2["F2"].number_format, "$#,##0")
        self.assertEqual(src2["Page"], "Interest Rate Risk › Curve Matrix (3M × 5Y) › ΔNII")
        self.assertIn("whole US dollars", src2["Units"])

    def test_missing_repdte_keeps_latest_stem_and_omits_report_date(self):
        result = {"inputs": {}, "beta_used": 0.5, "beta_mode": "textbook",
                  "asset_beta": 1.0, "scenarios": [
                      {"name": "Parallel +100", "short_change_bps": 100,
                       "long_change_bps": 100, "nim_new_pct": 3.45,
                       "nim_delta_bps": 25.0, "nii_delta_usd": 1.0, "description": ""}]}
        self._patch("ui.rate_sensitivity.run_curve_sensitivity", return_value=result)
        self.rs._render_named_scenarios("TST", {"NIMY": 3.2}, [], "textbook", None, 1.0)
        _wb, _ws, src, kw = self._book()
        self.assertEqual(kw["file_name"], "nim_scenarios_TST_latest.xlsx")
        self.assertNotIn("Report date", src)                 # never invented


# ── Branch analytics (FDIC SOD, owner-resolved store) ────────────────────

def _roster_fixture():
    """Three branches, two counties, SOD 2025. Deposits are $THOUSANDS as the
    store keeps them (DEPSUMBR). Walnut reports no deposits → n/a, never 0."""
    df = pd.DataFrame([
        ("Main Office", "1 Main St", "Oakland", "CA", "Alameda",
         "San Francisco-Oakland-Berkeley, CA", 300_000, "06001", 2025, 111, 37.80, -122.27),
        ("Elm Branch", "2 Elm St", "Berkeley", "CA", "Alameda",
         "San Francisco-Oakland-Berkeley, CA", 100_000, "06001", 2025, 111, 37.87, -122.27),
        ("Walnut Branch", "3 Walnut Blvd", "Walnut Creek", "CA", "Contra Costa",
         "San Francisco-Oakland-Berkeley, CA", None, "06013", 2025, 111, 37.90, -122.06),
    ], columns=["branch_name", "address", "city", "state", "county", "msa_name",
                "deposits", "stcntybr", "year", "cert", "lat", "lng"])
    return df, []


class TestBranchAnalyticsExports(_ExportSite):
    PAGE_MODULES = ("ui.branch_analytics",)

    def setUp(self):
        super().setUp()
        import ui.branch_analytics as ba
        self.ba = ba
        self._patch("ui.branch_analytics.get_fdic_cert", return_value=111)
        self._patch("ui.branch_analytics.get_bank_info",
                    return_value={"name": "Test Bancorp"})
        self._patch("ui.branch_analytics._roster",
                    side_effect=lambda cert: _roster_fixture())

    def test_branch_list_export_unscaled_k_and_share(self):
        """Total 300,000 + 100,000 = 400,000 $K → 75% / 25%; Walnut n/a."""
        self.ba.render_branch_list("TST")
        wb, ws, src, kw = self._book()
        self.assertEqual(kw["file_name"], "branches_TST_2025.xlsx")
        self.assertEqual(kw["key"], "exp_branches_TST")
        rows = self._grid(ws)
        self.assertEqual(rows[0], ["Branch", "Address", "City", "ST", "County", "MSA",
                                   "Deposits ($K)", "% of bank (%)"])
        self.assertEqual(rows[1], ["Main Office", "1 Main St", "Oakland", "CA", "Alameda",
                                   "San Francisco-Oakland-Berkeley, CA", 300_000, 75.0])
        self.assertEqual(rows[2][6:], [100_000, 25.0])
        self.assertNotEqual(ws["G2"].value, 300_000_000)         # never ×1000
        self._assert_na(ws, "G4")
        self._assert_na(ws, "H4")
        self.assertEqual(ws["G2"].number_format, "#,##0")
        self.assertEqual(ws["H2"].number_format, '0.00"%"')
        self.assertEqual(src["Page"], "Market Analysis › Branch List")
        self.assertEqual(src["Ticker"], "TST")
        self.assertEqual(src["Company"], "Test Bancorp")
        self.assertEqual(src["FDIC cert"], 111)
        self.assertEqual(src["SOD survey year"], 2025)
        self.assertEqual(src["Footprint"], "FDIC SOD 2025")
        self.assertIn("Summary of Deposits", src["Source"])
        self.assertIn("owner-resolved", src["Source"])
        self.assertIn("thousands", src["Units"])
        self.assertIn("($K)", src["Units"])

    def test_branch_competitors_export(self):
        """Footprint = Alameda (subject 400,000 + Rival 600,000) + Contra Costa
        (subject 50,000 + Rival 150,000) = 1,200,000 $K. Rival 750,000 =
        62.50%, 4 branches, 2 counties; subject 450,000 = 37.50%."""
        def county(fips, year):
            subj = (111, "Test Bank", "TST")
            rival = (222, "Rival Bank", None)
            data = {"06001": [(*subj, 2, 400_000), (*rival, 3, 600_000)],
                    "06013": [(*subj, 1, 50_000), (*rival, 1, 150_000)]}[fips]
            return pd.DataFrame(data, columns=["cert", "bank_name", "ticker",
                                               "n_branches", "total_deposits"])
        self._patch("ui.branch_analytics._county_banks", side_effect=county)

        self.ba.render_branch_competitors("TST")
        wb, ws, src, kw = self._book()
        self.assertEqual(kw["file_name"], "branch_competitors_TST_2025.xlsx")
        self.assertEqual(kw["key"], "exp_branch_competitors_TST")
        rows = self._grid(ws)
        self.assertEqual(rows[0], ["Bank", "Ticker", "Shared counties", "Branches",
                                   "Deposits ($K)", "% of footprint deposits (%)"])
        self.assertEqual(rows[1], ["Rival Bank", "n/a", 2, 4, 750_000.0, 62.5])
        self.assertEqual(rows[2], ["Test Bank", "TST", 2, 3, 450_000.0, 37.5])
        self.assertEqual(ws["C2"].number_format, "#,##0")
        self.assertEqual(ws["E2"].number_format, "#,##0")
        self.assertEqual(ws["F2"].number_format, '0.00"%"')
        self.assertEqual(src["Page"], "Market Analysis › Branch Competitors")
        self.assertEqual(src["Footprint counties"], 2)
        self.assertEqual(src["Footprint deposits, all banks ($K)"], 1_200_000.0)
        self.assertEqual(src["SOD survey year"], 2025)
        self.assertIn("thousands", src["Units"])

    def test_market_demographics_export(self):
        """Alameda: 300,000 + 100,000 = 400,000 $K bank deposits; Census whole
        dollars; unemployment 4.3%; home value absent → n/a. Contra Costa has
        no Census response → omitted and counted."""
        def demo(state_fips, county_fips):
            if (state_fips, county_fips) == ("06", "001"):
                return {"population": 1_600_000, "median_hh_income": 112_000,
                        "median_home_value": None, "unemployment_rate_pct": 4.3,
                        "vintage": "ACS5 2023"}
            return None
        self._patch("data.census_client.get_county_demographics", side_effect=demo)

        self.ba.render_market_demographics("TST")
        wb, ws, src, kw = self._book()
        self.assertEqual(kw["file_name"], "market_demographics_TST_2025.xlsx")
        self.assertEqual(kw["key"], "exp_market_demographics_TST")
        rows = self._grid(ws)
        self.assertEqual(rows[0], ["County FIPS", "County", "Bank deposits ($K)",
                                   "Population", "Median HH income ($)",
                                   "Median home value ($)", "Unemployment (%)", "Vintage"])
        self.assertEqual(len(rows), 2)                            # one county exported
        self.assertEqual(rows[1], ["06001", "Alameda, CA", 400_000, 1_600_000, 112_000,
                                   "n/a", 4.3, "ACS5 2023"])
        self.assertEqual(ws["A2"].number_format, "@")             # FIPS keeps its zero
        self.assertEqual(ws["C2"].number_format, "#,##0")
        self.assertEqual(ws["D2"].number_format, "#,##0")
        self.assertEqual(ws["E2"].number_format, "$#,##0")
        self.assertEqual(ws["G2"].number_format, '0.0"%"')
        self._assert_na(ws, "F2")
        self.assertEqual(src["Page"], "Market Analysis › Market Demographics")
        self.assertEqual(src["SOD survey year"], 2025)
        self.assertEqual(src["Counties omitted (no Census response)"], 1)
        self.assertIn("ACS 5-year", src["Census source"])
        self.assertIn("thousands", src["Units"])
        self.assertIn("whole US dollars", src["Units"])


# ── HMDA (CFPB, separate source) ─────────────────────────────────────────

class TestHmdaExports(_ExportSite):
    PAGE_MODULES = ("ui.hmda_view",)

    def test_by_year_and_by_state_exports(self):
        """2024 state split: WA 200,000,000 / 288,000,000 = 69.444…%,
        OR 88,000,000 / 288,000,000 = 30.555…%. Volumes raw dollars."""
        import ui.hmda_view as hv
        by_year = {2023: {"count": 1200, "volume_usd": 360_000_000.0},
                   2024: {"count": 900, "volume_usd": 288_000_000.0}}
        states = [{"state": "OR", "count": 300, "volume_usd": 88_000_000.0},
                  {"state": "WA", "count": 600, "volume_usd": 200_000_000.0}]
        self._patch("ui.hmda_view._lei_for", return_value="LEI123")
        self._patch("data.hmda_client.originations_by_year", return_value=by_year)
        self._patch("data.hmda_client.latest_breakdown", return_value=states)

        hv.render_hmda_mortgages("TST")

        wb, ws, src, kw = self._book(n_expected=2, idx=0)
        self.assertEqual(kw["file_name"], "hmda_originations_TST_2023_2024.xlsx")
        self.assertEqual(kw["key"], "exp_hmda_years_TST")
        rows = self._grid(ws)
        self.assertEqual(rows[0], ["Year", "Originations", "Volume ($)"])
        self.assertEqual(rows[1], [2024, 900, 288_000_000.0])   # newest first, as shown
        self.assertEqual(rows[2], [2023, 1200, 360_000_000.0])
        self.assertEqual(ws["B2"].number_format, "#,##0")
        self.assertEqual(ws["C2"].number_format, "$#,##0")
        self.assertEqual(src["Page"], "Company › HMDA Mortgages")
        self.assertEqual(src["Ticker"], "TST")
        self.assertEqual(src["LEI"], "LEI123")
        self.assertEqual(src["Source"], "CFPB HMDA LAR 2023–2024 (public loan-level data)")
        self.assertEqual(src["HMDA years"], "2023–2024")
        self.assertIn("$10,000-bucket midpoint", src["Note"])
        self.assertNotIn("FDIC cert", src)                        # not an FDIC source

        wb2, ws2, src2, kw2 = self._book(n_expected=2, idx=1)
        self.assertEqual(kw2["file_name"], "hmda_by_state_TST_2024.xlsx")
        self.assertEqual(kw2["key"], "exp_hmda_states_TST")
        rows2 = self._grid(ws2)
        self.assertEqual(rows2[0], ["State", "Originations", "Volume ($)", "% of volume (%)"])
        self.assertEqual(rows2[1][:3], ["WA", 600, 200_000_000.0])   # sorted by volume
        self.assertAlmostEqual(rows2[1][3], 200 / 288 * 100)
        self.assertEqual(rows2[2][:3], ["OR", 300, 88_000_000.0])
        self.assertAlmostEqual(rows2[2][3], 88 / 288 * 100)
        self.assertEqual(ws2["C2"].number_format, "$#,##0")
        self.assertEqual(ws2["D2"].number_format, '0.0"%"')
        self.assertEqual(src2["Source"], "CFPB HMDA LAR 2024 (public loan-level data)")
        self.assertEqual(src2["HMDA year"], 2024)
        self.assertIn("top 2 of 2", src2["States"])

    def test_no_lei_renders_no_export(self):
        import ui.hmda_view as hv
        self._patch("ui.hmda_view._lei_for", return_value=None)
        hv.render_hmda_mortgages("TST")
        self.assertEqual(self.calls, [])


if __name__ == "__main__":
    unittest.main()
