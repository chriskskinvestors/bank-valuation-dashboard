"""Fabricated-zero fixes (owner 2026-09-24, "fix the pre existing").

Each site here used to turn ABSENT data into a displayed 0 / 0% / partial
total (memory: fabricated-aggregate-defect-class). Pins:

  * utils.aggregate.strict_sum — a total is known only when every part is;
  * branches_store — a blank DEPSUMBR is stored NULL (never 0) and a bank's
    SQL total is NULL when any of its branches is NULL (never a partial);
  * deposit_lookup._market_share — unknown deposits → NaN share for the whole
    market, never 0% / never a share of a partial total; screen shows "—";
  * branch_analytics demographics — a county of unreported branches exports
    n/a, not $0;
  * financial_highlights TCE/TA — absent equity/assets → n/a, not a ratio
    computed from 0;
  * data_quality — an unreported FDIC field shows "n/a" (not "$nan"); a
    missing securities duration shows "— (not reported)" (not "0.00 yrs").
"""
import io
import math
import types
import unittest
from unittest import mock

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.pool import StaticPool

from utils.aggregate import strict_sum
from utils.formatting import fmt_dollars_from_thousands


class TestStrictSum(unittest.TestCase):

    def test_known_values_sum(self):
        self.assertEqual(strict_sum([1, 2, 3.5]), 6.5)
        self.assertEqual(strict_sum(["12", 3]), 15.0)          # numeric strings count

    def test_any_unknown_makes_total_unknown(self):
        self.assertIsNone(strict_sum([1, None, 3]))
        self.assertIsNone(strict_sum([1, float("nan")]))
        self.assertIsNone(strict_sum([1, "n/a"]))
        self.assertIsNone(strict_sum([]))                       # nothing to total

    def test_na_sentinel_for_pandas(self):
        v = strict_sum([1, None], na=float("nan"))
        self.assertTrue(math.isnan(v))
        s = pd.Series([1.0, float("nan")])
        self.assertTrue(math.isnan(strict_sum(s, na=float("nan"))))
        self.assertEqual(strict_sum(pd.Series([1.0, 2.0])), 3.0)


class TestStoreNeverFabricatesZero(unittest.TestCase):
    """SQLite store: blank DEPSUMBR → NULL; owner total NULL when any branch
    NULL; ORDER BY puts unknown totals last."""

    def setUp(self):
        import data.db as db
        import data.branches_store as bs
        self.bs = bs
        self._eng = create_engine("sqlite://", connect_args={"check_same_thread": False},
                                  poolclass=StaticPool)
        self._orig = db.get_engine
        db.get_engine = lambda: self._eng
        bs._engine = None
        bs.init_branches_schema()

    def tearDown(self):
        import data.db as db
        db.get_engine = self._orig
        self.bs._engine = None

    def _sod_df(self, rows):
        base = {"YEAR": 2025, "NAMEBR": "b", "ADDRESBR": "a", "CITYBR": "c",
                "STALPBR": "WA", "ZIPBR": "99362", "CNTYNAMB": "Walla Walla",
                "STCNTYBR": "53071", "MSABR": "", "MSANAMB": "",
                "SIMS_LATITUDE": None, "SIMS_LONGITUDE": None, "BRSERTYP": "11"}
        return pd.DataFrame([{**base, **r} for r in rows])

    def test_blank_depsumbr_stored_null_and_total_unknown(self):
        # Bank A: one branch reported (500), one blank → total UNKNOWN.
        self.bs.upsert_branches("AAAA", 1, self._sod_df([
            {"BRNUM": 1, "DEPSUMBR": 500}, {"BRNUM": 2, "DEPSUMBR": ""}]))
        # Bank B: both reported → 300 + 0 (a real LPO zero) = 300.
        self.bs.upsert_branches("BBBB", 2, self._sod_df([
            {"BRNUM": 1, "DEPSUMBR": 300}, {"BRNUM": 2, "DEPSUMBR": 0}]))
        with self._eng.connect() as c:
            stored = c.execute(text(
                "SELECT cert, brnum, deposits FROM branches ORDER BY cert, brnum")).fetchall()
        self.assertEqual([tuple(r) for r in stored],
                         [(1, 1, 500), (1, 2, None), (2, 1, 300), (2, 2, 0)])
        df = self.bs.get_banks_by_county("53071", year=2025)
        by = {int(r.cert): r for r in df.itertuples(index=False)}
        self.assertEqual(by[2].total_deposits, 300)
        self.assertTrue(pd.isna(by[1].total_deposits), "partial total must be unknown")
        self.assertEqual(int(by[1].n_branches), 2)
        # Unknown sorts last, not first (NULLS LAST).
        self.assertEqual([int(c) for c in df["cert"]], [2, 1])


class _Ctx:
    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


class TestMarketShareUnknownPropagates(unittest.TestCase):

    def _share(self, deposits):
        import ui.deposit_lookup as dl
        import data.branches_store as bs
        frame = pd.DataFrame({
            "owner_key": [f"o{i}" for i in range(len(deposits))],
            "cert": list(range(1, len(deposits) + 1)),
            "ticker": ["A", "B", None][:len(deposits)],
            "bank_name": ["Alpha", "Beta", "Gamma"][:len(deposits)],
            "n_branches": [3, 2, 1][:len(deposits)],
            "total_deposits": deposits,
        })
        with mock.patch.object(bs, "get_banks_by_county", lambda k, year=None: frame):
            return dl._market_share("county", "53071", 2025)

    def test_all_known_shares_sum_to_100(self):
        df = self._share([600.0, 400.0])
        self.assertEqual(list(df["market_share"].round(6)), [60.0, 40.0])
        self.assertEqual(list(df["rank"]), [1, 2])

    def test_one_unknown_makes_every_share_unknown_never_zero(self):
        df = self._share([600.0, None, 400.0])
        self.assertTrue(df["market_share"].isna().all())
        self.assertFalse((df["deposits"] == 0).any())          # never fabricated 0
        # Unknown deposits sort LAST; ranks still sequential.
        self.assertEqual(list(df["NAMEFULL"]), ["Alpha", "Gamma", "Beta"])
        self.assertTrue(pd.isna(df["deposits"].iloc[-1]))

    def test_render_shows_dash_and_na_not_nan(self):
        import ui.deposit_lookup as dl
        captured, caps = [], []
        fake_st = types.SimpleNamespace(
            caption=lambda t, *a, **k: caps.append(t), warning=lambda *a, **k: None,
            markdown=lambda *a, **k: None, container=lambda *a, **k: _Ctx())
        ms = pd.DataFrame({"owner_key": ["o0", "o1"], "CERT": [1, 2], "TICKER": ["A", None],
                           "NAMEFULL": ["Alpha", "Beta"], "branches": [3, 2],
                           "deposits": [600.0, float("nan")],
                           "market_share": [float("nan"), float("nan")], "rank": [1, 2]})
        fake_export = types.SimpleNamespace(download_button=lambda *a, **k: None,
                                            container=lambda *a, **k: _Ctx())
        import ui.export as ex
        with mock.patch.object(dl, "st", fake_st), \
             mock.patch.object(dl, "_market_share", lambda *a, **k: ms), \
             mock.patch.object(dl, "_skeleton", lambda: _Ctx()), \
             mock.patch.object(dl, "ksk_table", lambda df, **k: captured.append(df)), \
             mock.patch.object(dl, "_linked_tickers", lambda s: list(s)), \
             mock.patch.object(ex, "st", fake_export):
            dl._render_market_share("county", "53071", "Walla Walla County, WA",
                                    2025, "o0", "Alpha", lambda v: fmt_dollars_from_thousands(v, 2))
        self.assertEqual(list(captured[0]["Share"]), ["—", "—"])
        self.assertEqual(list(captured[0]["Deposits"])[1], "—")
        self.assertIn("**n/a** share", caps[0])
        self.assertNotIn("nan", caps[0])


class TestDemographicsCountyUnknown(unittest.TestCase):

    def test_county_of_unreported_branches_exports_na_not_zero(self):
        import ui.branch_analytics as ba
        import ui.export as ex
        import data.census_client as cc
        roster = pd.DataFrame({
            "stcntybr": ["53071", "53071", "53005"], "county": ["Walla Walla", "Walla Walla", "Benton"],
            "state": ["WA", "WA", "WA"], "deposits": [float("nan"), float("nan"), 250_000.0],
            "year": [2025, 2025, 2025],
        })
        demo = {"population": 60_000, "median_hh_income": 62_000, "median_home_value": 310_000,
                "unemployment_rate_pct": 4.1, "vintage": "2023 ACS 5-yr"}
        calls, shown = [], []
        fake_st = types.SimpleNamespace(caption=lambda *a, **k: None, markdown=lambda *a, **k: None,
                                        dataframe=lambda df, **k: shown.append(df),
                                        container=lambda *a, **k: _Ctx())
        fake_export = types.SimpleNamespace(
            download_button=lambda label, data, **k: calls.append((data, k)),
            container=lambda *a, **k: _Ctx())
        with mock.patch.object(ba, "st", fake_st), \
             mock.patch.object(ba, "_roster", lambda cert: (roster, [])), \
             mock.patch.object(ba, "get_fdic_cert", lambda t: 28489, create=True), \
             mock.patch.object(ba, "get_bank_info", lambda t: {"name": "Banner"}, create=True), \
             mock.patch.object(cc, "get_county_demographics", lambda s, c: demo), \
             mock.patch.object(ex, "st", fake_export):
            ba.render_market_demographics("BANR")
        self.assertEqual(len(calls), 1)
        ws = load_workbook(io.BytesIO(calls[0][0]())).worksheets[0]
        hdr = [c.value for c in ws[1]]
        dep_col = hdr.index("Bank deposits ($K)") + 1
        by_county = {ws.cell(r, hdr.index("County") + 1).value: ws.cell(r, dep_col).value
                     for r in range(2, ws.max_row + 1)}
        self.assertEqual(by_county["Benton, WA"], 250_000)
        self.assertEqual(by_county["Walla Walla, WA"], "n/a")     # was 0
        # Known county sorts first; on screen the unknown county shows "—".
        self.assertEqual(ws.cell(2, hdr.index("County") + 1).value, "Benton, WA")
        disp = shown[0]
        self.assertEqual(list(disp["Bank deposits"]), ["$250.0M", "—"])


class TestTceTaNeverFromFabricatedZero(unittest.TestCase):

    def _b(self, rec):
        import ui.financial_highlights as fh
        payloads = []

        def P(v, *a, **k):
            payloads.append((v, k.get("raw")))
            return v
        b = fh._tce_ta_builder({0: rec}, {0: "Dec 31, 2025"}, "http://fdic", P)
        b(0)
        return payloads[0]

    def test_all_present_hand_computed(self):
        # (1,000 − 100) ÷ (12,400 − 100) × 100 = 7.317073…
        v, raw = self._b({"EQTOT": 1000, "INTAN": 100, "ASSET": 12_400})
        self.assertEqual(v, "7.32%")
        self.assertAlmostEqual(raw, 900 / 12_300 * 100, places=9)

    def test_absent_intangibles_means_none_reported(self):
        v, raw = self._b({"EQTOT": 1000, "ASSET": 12_400})
        self.assertAlmostEqual(raw, 1000 / 12_400 * 100, places=9)

    def test_absent_equity_or_assets_is_na(self):
        for rec in ({"INTAN": 100, "ASSET": 12_400}, {"EQTOT": 1000, "INTAN": 100}):
            v, raw = self._b(rec)
            self.assertEqual((v, raw), ("—", None), rec)   # was −0.81% / n/a-by-luck


class TestDataQualityAbsentIsNa(unittest.TestCase):

    def test_ladder_missing_duration_says_not_reported(self):
        import ui.data_quality as dq
        import data.call_report_store as crs
        import data.ffiec_client as fc
        import ui.export as ex
        shown = []
        fake_st = types.SimpleNamespace(
            dataframe=lambda df, **k: shown.append(df), caption=lambda *a, **k: None,
            markdown=lambda *a, **k: None, warning=lambda *a, **k: None,
            info=lambda *a, **k: None, success=lambda *a, **k: None,
            container=lambda *a, **k: _Ctx(), subheader=lambda *a, **k: None)
        fake_export = types.SimpleNamespace(download_button=lambda *a, **k: None,
                                            container=lambda *a, **k: _Ctx())
        with mock.patch.object(dq, "st", fake_st), \
             mock.patch.object(fc, "is_configured", lambda: True), \
             mock.patch.object(fc, "health_check", lambda: {"ok": True, "days_until_expiry": 120}), \
             mock.patch.object(crs, "get_latest_ladder", lambda cert: {
                 "reporting_period": "12/31/2025", "weighted_avg_duration_years": None,
                 "floating_loan_share": None, "source": "ffiec"}), \
             mock.patch.object(ex, "st", fake_export):
            dq._render_ffiec_status(28489, "BANR")
        rows = {r["Field"]: r["Value"] for r in shown[0].to_dict("records")}
        self.assertEqual(rows["Securities duration (wtd-avg)"], "— (not reported)")   # was "0.00 yrs"
        self.assertEqual(rows["Floating-loan share (RC-C Memo 2)"], "— (not reported)")


if __name__ == "__main__":
    unittest.main()
