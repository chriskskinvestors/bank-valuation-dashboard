"""Export call sites migrated onto ui.export.table_export (owner directive
2026-09-22): Corporate Structure, Ownership (Detailed + Crossholdings) and
Insider Activity (transactions + per-insider summary).

Each test drives the real render function with a small fixture, captures the
deferred workbook builder that table_export hands st.download_button, builds
the .xlsx and reads it back with openpyxl — asserting header labels (units in
the header), that numeric cells are NUMBERS carrying the exact Excel number
format, "n/a" for absent values, and the Source-sheet provenance rows. All
expectations are hand-computed from the fixtures.

Run: PYTHONIOENCODING=utf-8 python -m unittest tests.test_export_sites_ownership
"""
from __future__ import annotations

import datetime as dt
import io
import types
import unittest
from unittest.mock import patch

from tests._streamlit_stub import install

install()

import ui.export as ex  # noqa: E402

_MISSING = object()


class _Ctx:
    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def _noop(*a, **k):
    return None


def _load(data):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(data() if callable(data) else data))


def _grid(ws, max_rows=None):
    return [[c.value for c in r]
            for r in ws.iter_rows(min_row=1, max_row=max_rows or ws.max_row)]


def _source(wb) -> dict:
    return dict(_grid(wb["Source"]))


class _ExportCapture(unittest.TestCase):
    """Shared-stub extension (the documented tests/_streamlit_stub pattern:
    install() then setattr the extras these pages reach; restore after) +
    download_button capture on the st module ui.export is bound to. Never
    replaces the streamlit module or reloads ui modules — under discovery
    every ui module is bound to this ONE module object, and swapping it
    strands modules imported earlier on a different object."""

    _WIDGETS = {
        "markdown": _noop, "caption": _noop, "subheader": _noop, "info": _noop,
        "dataframe": _noop, "plotly_chart": _noop, "download_button": _noop,
        "spinner": lambda *a, **k: _Ctx(),
        "container": lambda *a, **k: _Ctx(),
        "expander": lambda *a, **k: _Ctx(),
        "columns": lambda spec, **k: [_Ctx() for _ in range(
            spec if isinstance(spec, int) else len(spec))],
        "selectbox": lambda label, options=None, **k: (options[0] if options else None),
        "radio": lambda label, options=None, **k: (options[0] if options else None),
        "empty": lambda: types.SimpleNamespace(markdown=_noop, empty=_noop),
        "column_config": types.SimpleNamespace(LinkColumn=_noop, TextColumn=_noop),
    }

    def setUp(self):
        self.st = install()
        saved = {n: getattr(self.st, n, _MISSING) for n in self._WIDGETS}
        for n, v in self._WIDGETS.items():
            setattr(self.st, n, v)

        def _restore():
            for n, v in saved.items():
                if v is _MISSING:
                    if hasattr(self.st, n):
                        delattr(self.st, n)
                else:
                    setattr(self.st, n, v)
        self.addCleanup(_restore)

    def _capture(self):
        calls = []
        saved = ex.st.download_button
        ex.st.download_button = (
            lambda label, data, **k: calls.append((label, data, k)))
        self.addCleanup(setattr, ex.st, "download_button", saved)
        return calls


# ── Corporate Structure ─────────────────────────────────────────────────

def _tree():
    return {
        "entity": {"name": "BANNER CORPORATION", "rssd": 2126977,
                   "type": "Bank Holding Company", "location": "Walla Walla, WA"},
        "children": [
            {"entity": {"name": "BANNER BANK", "rssd": 352772,
                        "type": "State Non-member Bank", "location": "Walla Walla, WA"},
             "ownership_pct": 100.0, "relationship": "Controlled",
             "children": [
                 {"entity": {"name": "COMMUNITY FINANCIAL CORPORATION",
                             "rssd": 111, "type": "Domestic Entity (Other)",
                             "location": "Lake Oswego, OR"},
                  "ownership_pct": 51.5, "relationship": "Controlled",
                  "children": []},
             ]},
            {"entity": {"name": "BANNER CAPITAL TRUST V", "rssd": 222,
                        "type": "Domestic Entity (Other)", "location": None},
             "ownership_pct": None, "relationship": "Controlled", "children": []},
        ],
        "as_of": "2026-09-01",
    }


class TestCorporateStructureExport(_ExportCapture):

    def test_depth_column_numeric_and_entity_unindented(self):
        import ui.corporate_structure as cs
        calls = self._capture()
        with patch.object(cs, "get_name", lambda t: "Banner Corporation"), \
             patch.object(cs, "get_fdic_cert", lambda t: 28489), \
             patch("data.fdic_client.get_rssd_for_cert", lambda c: 352772), \
             patch("data.nic_client.get_parent",
                   side_effect=[{"rssd": 2126977, "name": "BANNER CORPORATION"}, None]), \
             patch("data.nic_client.get_org_hierarchy", lambda r: _tree()):
            cs.render_corporate_structure("BANR")

        self.assertEqual(len(calls), 1)
        _label, data, kw = calls[0]
        self.assertEqual(kw["file_name"], "BANR_corporate_structure_2026-09-01.xlsx")
        self.assertEqual(kw["key"], "exp_struct_BANR")
        wb = _load(data)
        ws = wb["Corporate structure"]
        rows = _grid(ws)
        self.assertEqual(rows[0], ["Entity", "Depth", "Type", "Location",
                                   "Ownership (%)", "Control"])
        # pre-order: root, bank, bank's sub, then the trust
        self.assertEqual([r[0] for r in rows[1:]],
                         ["BANNER CORPORATION", "BANNER BANK",
                          "COMMUNITY FINANCIAL CORPORATION", "BANNER CAPITAL TRUST V"])
        for r in rows[1:]:
            self.assertEqual(r[0], r[0].lstrip(), "Entity must not carry indent spaces")
        self.assertEqual([r[1] for r in rows[1:]], [0, 1, 2, 1])
        for rr in range(2, 6):
            self.assertIsInstance(ws.cell(rr, 2).value, int)
            self.assertEqual(ws.cell(rr, 2).number_format, "#,##0")
        # root has no ownership/control → n/a, never 0 / blank
        self.assertEqual(rows[1][4:], ["n/a", "n/a"])
        self.assertEqual(ws["E3"].value, 100.0)
        self.assertEqual(ws["E3"].number_format, '0.00"%"')
        self.assertEqual(ws["E4"].value, 51.5)
        self.assertEqual(rows[4][3:5], ["n/a", "n/a"])   # trust: no location, no pct
        self.assertEqual(ws.freeze_panes, "B2")

        src = _source(wb)
        self.assertEqual(src["Page"], "Company Analysis › Overview › Corporate Structure")
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["FDIC cert"], 28489)
        self.assertEqual(src["Bank RSSD"], 352772)
        self.assertEqual(src["Top holder RSSD"], 2126977)
        self.assertIn("NIC", src["Source"])
        self.assertIn("PCT_EQUITY", src["Source"])
        self.assertEqual(src["As of"], "2026-09-01")
        self.assertIn("percent units", src["Units"])


# ── Ownership ───────────────────────────────────────────────────────────

_HOLDERS = [
    {"filer_name": "VANGUARD GROUP INC", "filer_cik": "0000102909",
     "accession": "0000102909-26-000001", "shares": 3_500_000.0,
     "value_usd": 210_000_000.0, "date_filed": "2026-08-14"},
    {"filer_name": "NEW HOLDER LLC", "filer_cik": "0000000002",
     "accession": "0000000002-26-000001", "shares": 100_000.0,
     "value_usd": None, "date_filed": "2026-08-10"},
]

_HIST = {
    "VANGUARD GROUP INC": {"2026Q2": {"shares": 3_500_000.0, "value_usd": 210_000_000.0},
                           "2026Q1": {"shares": 3_000_000.0, "value_usd": 171_000_000.0}},
    "NEW HOLDER LLC": {"2026Q2": {"shares": 100_000.0, "value_usd": None}},
}


class TestOwnershipDetailedExport(_ExportCapture):

    def test_raw_numerics_with_units_and_qoq_provenance(self):
        import ui.ownership as own
        calls = self._capture()
        with patch.object(own, "get_name", lambda t: "Banner Corporation"), \
             patch.object(own, "fetch_institutional_holdings",
                          lambda t, n, max_filers=30: [dict(h) for h in _HOLDERS]), \
             patch("data.form13f_client.get_holder_history",
                   lambda t, quarters=20: _HIST), \
             patch("data.fmp_client.get_quote", lambda t: {"price": 60.0}):
            own.render_ownership_detailed("BANR", {"shares_outstanding": 35_000_000})

        self.assertEqual(len(calls), 1)
        _label, data, kw = calls[0]
        self.assertEqual(kw["file_name"], "ownership_detailed_BANR.xlsx")
        self.assertEqual(kw["key"], "exp_owndet_BANR")
        wb = _load(data)
        ws = wb["Ownership detailed"]
        rows = _grid(ws)
        self.assertEqual(rows[0], [
            "Holder", "Filer CIK", "Accession", "Shares", "Δ Shares (QoQ)",
            "Δ Shares (QoQ) (%)", "New Position", "% CSO (%)", "Mkt Value ($)",
            "Reported Value ($)", "Date Filed"])
        # Vanguard: Δ = 3.5M − 3.0M = 500,000 (+16.67%); %CSO = 3.5M/35M = 10%;
        # Mkt = 3.5M × $60 = $210,000,000.
        v = rows[1]
        self.assertEqual(v[:5], ["VANGUARD GROUP INC", "0000102909",
                                 "0000102909-26-000001", 3_500_000, 500_000])
        self.assertAlmostEqual(v[5], 16.666666, places=4)
        self.assertIs(v[6], False)
        self.assertEqual(v[7], 10.0)
        self.assertEqual(v[8], 210_000_000)
        self.assertEqual(v[9], 210_000_000)
        self.assertEqual(v[10], dt.datetime(2026, 8, 14))
        # New holder: absent from prior snapshot → New, deltas n/a; %CSO =
        # 100,000/35M = 0.2857%; Mkt = 100,000 × $60 = $6,000,000; no
        # reported value → n/a (never 0).
        n = rows[2]
        self.assertEqual(n[0], "NEW HOLDER LLC")
        self.assertEqual(n[4:7], ["n/a", "n/a", True])
        self.assertAlmostEqual(n[7], 0.285714, places=5)
        self.assertEqual(n[8], 6_000_000)
        self.assertEqual(n[9], "n/a")
        # formats: whole-number shares, percent units, whole dollars, ISO date
        self.assertEqual(ws["D2"].number_format, "#,##0")
        self.assertEqual(ws["E2"].number_format, "#,##0")
        self.assertEqual(ws["F2"].number_format, '0.00"%"')
        self.assertEqual(ws["H2"].number_format, '0.00"%"')
        self.assertEqual(ws["I2"].number_format, "$#,##0")
        self.assertEqual(ws["J2"].number_format, "$#,##0")
        self.assertEqual(ws["K2"].number_format, "yyyy-mm-dd")
        self.assertEqual(ws["B2"].number_format, "@")   # CIK keeps leading zeros
        self.assertEqual(ws["J3"].number_format, "General")   # n/a carries no format
        self.assertTrue(ws["J3"].font.italic)

        src = _source(wb)
        self.assertEqual(src["Page"], "Company Analysis › Ownership › Ownership Detailed")
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["Company"], "Banner Corporation")
        self.assertIn("13F-HR", src["Source"])
        self.assertEqual(src["Prior quarter (QoQ)"], "2026Q1")
        self.assertEqual(src["Shares outstanding (% CSO)"], 35_000_000)
        self.assertEqual(src["Price ($) (Mkt Value)"], 60.0)
        self.assertIn("whole US dollars", src["Units"])
        self.assertIn("percent units", src["Units"])


class TestCrossholdingsExport(_ExportCapture):

    def test_quarter_in_filename_and_provenance(self):
        import ui.ownership as own
        calls = self._capture()
        x = {"quarter": "2026Q2", "coverage": 3, "rows": [
            {"holder": "VANGUARD GROUP INC", "subject_value_usd": 210_000_000.0,
             "others": [{"ticker": "EWBC", "shares": 4_000_000.0, "value_usd": 392_000_000.0},
                        {"ticker": "CATY", "shares": 900_000.0, "value_usd": None}]},
        ]}
        with patch.object(own, "get_name", lambda t: "Banner Corporation"), \
             patch("data.form13f_client.get_crossholdings", lambda t: x):
            own.render_crossholdings("BANR")

        self.assertEqual(len(calls), 1)
        _label, data, kw = calls[0]
        self.assertEqual(kw["file_name"], "crossholdings_BANR_2026Q2.xlsx")
        self.assertEqual(kw["key"], "exp_crossholdings_BANR")
        wb = _load(data)
        ws = wb["Crossholdings"]
        self.assertEqual(_grid(ws), [
            ["Institution", "Position in BANR ($)", "Other Bank", "Shares", "Reported Value ($)"],
            ["VANGUARD GROUP INC", 210_000_000, "EWBC", 4_000_000, 392_000_000],
            ["VANGUARD GROUP INC", 210_000_000, "CATY", 900_000, "n/a"],
        ])
        self.assertEqual(ws["B2"].number_format, "$#,##0")
        self.assertEqual(ws["D2"].number_format, "#,##0")
        self.assertEqual(ws["E2"].number_format, "$#,##0")
        src = _source(wb)
        self.assertEqual(src["Quarter"], "2026Q2")
        self.assertIn("3 other universe banks", src["Coverage"])
        self.assertEqual(src["Page"], "Company Analysis › Ownership › Crossholdings")


# ── Insider Activity ────────────────────────────────────────────────────

_TXS = [
    {"date": "2026-09-10", "insider": "SMITH JOHN", "role": "Director",
     "type": "Open market purchase", "code": "P", "shares": 1000.0, "price": 61.25,
     "value_usd": 61_250.0, "shares_after": 11_000.0, "direction": "Buy",
     "form_type": "non-derivative", "accession": "0000946673-26-000010",
     "filed_at": "2026-09-11T20:05:00+00:00"},
    {"date": "2026-08-01", "insider": "SMITH JOHN", "role": "Director",
     "type": "Open market sale", "code": "S", "shares": 400.0, "price": 60.0,
     "value_usd": 24_000.0, "shares_after": 10_000.0, "direction": "Sell",
     "form_type": "non-derivative", "accession": "0000946673-26-000009",
     "filed_at": "2026-08-02T21:00:00+00:00"},
    {"date": "2026-07-15", "insider": "DOE JANE", "role": "CFO",
     "type": "Option exercise", "code": "M", "shares": 500.0, "price": 30.0,
     "strike_price": 30.0, "value_usd": None, "shares_after": None,
     "direction": "Exercise", "form_type": "derivative",
     "accession": "0000946673-26-000008", "filed_at": None},
]


class TestInsiderActivityExports(_ExportCapture):

    def _render(self):
        import ui.insider_activity as ia
        calls = self._capture()
        with patch.object(ia, "get_cik", lambda t: 946673), \
             patch.object(ia, "get_name", lambda t: "Banner Corporation"), \
             patch.object(ia, "fetch_insider_trades",
                          lambda cik, months_back=12: [dict(t) for t in _TXS]), \
             patch("data.fmp_client.get_history", lambda t, rng: None):
            ia.render_insider_activity("BANR")
        self.assertEqual(len(calls), 2, "transactions + per-insider exports")
        return calls

    def test_transactions_export_raw_numerics(self):
        _label, data, kw = self._render()[0]
        self.assertEqual(kw["file_name"], "insider_transactions_BANR.xlsx")
        self.assertEqual(kw["key"], "exp_insider_transactions_BANR")
        wb = _load(data)
        ws = wb["Insider transactions"]
        rows = _grid(ws)
        self.assertEqual(rows[0], [
            "Date", "Insider", "Role", "Type", "Code", "Direction", "Security Type",
            "Shares", "Price ($)", "Value ($)", "Shares After", "Accession",
            "Filed At (UTC)", "Filing URL"])
        buy = rows[1]
        self.assertEqual(buy[0], dt.datetime(2026, 9, 10))
        self.assertEqual(buy[1:7], ["SMITH JOHN", "Director", "Open market purchase",
                                    "P", "Buy", "non-derivative"])
        self.assertEqual(buy[7:11], [1000, 61.25, 61_250, 11_000])
        self.assertEqual(buy[13], "https://www.sec.gov/Archives/edgar/data/946673/"
                                  "000094667326000010/0000946673-26-000010-index.htm")
        # option exercise: strike price kept as Price, Value/Shares After/Filed At n/a
        ex = rows[3]
        self.assertEqual(ex[1], "DOE JANE")
        self.assertEqual(ex[8], 30.0)
        self.assertEqual(ex[9:11], ["n/a", "n/a"])
        self.assertEqual(ex[12], "n/a")
        self.assertEqual(ws["A2"].number_format, "yyyy-mm-dd")
        self.assertEqual(ws["H2"].number_format, "#,##0")
        self.assertEqual(ws["I2"].number_format, "$#,##0.00")
        self.assertEqual(ws["J2"].number_format, "$#,##0")
        self.assertEqual(ws["K2"].number_format, "#,##0")
        self.assertEqual(ws.freeze_panes, "C2")
        src = _source(wb)
        self.assertEqual(src["Page"], "Company Analysis › Insiders › Insider Activity")
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["CIK"], 946673)
        self.assertIn("Form 4", src["Source"])
        self.assertEqual(src["Filter"], "All transactions · All · limit 20")

    def test_per_insider_summary_export(self):
        _label, data, kw = self._render()[1]
        self.assertEqual(kw["file_name"], "insider_summary_BANR.xlsx")
        self.assertEqual(kw["key"], "exp_insider_summary_BANR")
        wb = _load(data)
        ws = wb["Activity by insider"]
        # Only P/S non-derivative trades count: SMITH bought $61,250 and sold
        # $24,000 → net +$37,250 over 2 txns; DOE's exercise is excluded.
        self.assertEqual(_grid(ws), [
            ["Insider", "Role", "Buys ($)", "Sells ($)", "Net ($)", "Txns"],
            ["SMITH JOHN", "Director", 61_250, 24_000, 37_250, 2],
        ])
        for col in "CDE":
            self.assertEqual(ws[f"{col}2"].number_format, "$#,##0")
        self.assertEqual(ws["F2"].number_format, "#,##0")
        self.assertIsInstance(ws["F2"].value, int)
        src = _source(wb)
        self.assertEqual(src["CIK"], 946673)
        self.assertIn("open-market P/S", src["Scope"])
        self.assertIn("Activity by insider", src["Page"])


if __name__ == "__main__":
    unittest.main()
