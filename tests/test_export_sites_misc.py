"""The six table-export sites migrated onto ui.export.table_export
(owner directive 2026-09-22): historicals, macro (FDIC national rates +
print board), peer_rank (scorecard + leaderboard), filings,
recent_documents, transactions deal comps.

Each test drives the real renderer against synthetic data with the
streamlit stub installed, captures the deferred download, opens the
workbook and asserts cell VALUES and Excel number formats against
hand-written expectations — units especially:

  * FDIC $thousands export UNSCALED under a "($K)" header (owner choice);
  * percent values are percent units (3.42 = 3.42%) under a (%) header;
  * deal-comps dollars are whole dollars (ma_history / deal_comps contract);
  * P/Assets and core-deposit premium are fractions at source, exported
    ×100 under (%) — exactly what the on-screen table shows;
  * absent values are the literal "n/a", never 0.

Run: PYTHONIOENCODING=utf-8 python -m unittest tests.test_export_sites_misc
"""
from __future__ import annotations

import datetime as dt
import importlib
import io
import sys
import types
import unittest
from contextlib import contextmanager
from datetime import date, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from tests.test_render_smoke import _install_streamlit_stub  # noqa: E402

# Reload order matters: each module must be re-bound to the fresh stub AFTER
# the modules it imports from (chrome after export, recent_documents after
# filings, transactions after components).
_UI_MODULES = ("ui.export", "ui.chrome", "ui.components", "ui.states",
               "ui.historicals", "ui.macro", "ui.peer_rank", "ui.filings",
               "ui.recent_documents", "ui.transactions")


def _fresh_ui() -> dict:
    """Install the rich stub and (re)bind every module under test to it.
    Under discovery these modules may already be imported against the thin
    package stub (tests/__init__), whose `st` lacks the widgets the renderers
    call — reloading re-executes them against the rich one."""
    _install_streamlit_stub()
    mods = {}
    for name in _UI_MODULES:
        mods[name] = (importlib.reload(sys.modules[name]) if name in sys.modules
                      else importlib.import_module(name))
    return mods


@contextmanager
def _patched(obj, **attrs):
    saved = {k: getattr(obj, k) for k in attrs}
    for k, v in attrs.items():
        setattr(obj, k, v)
    try:
        yield
    finally:
        for k, v in saved.items():
            setattr(obj, k, v)


@contextmanager
def _fake_module(name: str, **attrs):
    """Serve a fake module for a renderer's function-local `from X import Y`
    (the import system returns the sys.modules entry, so the real data
    module is never imported and nothing reaches a network or DB)."""
    saved = sys.modules.get(name)
    mod = types.ModuleType(name)
    for k, v in attrs.items():
        setattr(mod, k, v)
    sys.modules[name] = mod
    try:
        yield mod
    finally:
        if saved is None:
            del sys.modules[name]
        else:
            sys.modules[name] = saved


def _grid(ws, max_rows=None):
    return [[c.value for c in r]
            for r in ws.iter_rows(min_row=1, max_row=max_rows or ws.max_row)]


class _SiteCase(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.m = _fresh_ui()
        cls.export = cls.m["ui.export"]

    def setUp(self):
        self.calls = []
        self._saved_dl = self.export.st.download_button
        self.export.st.download_button = (
            lambda label, data, **k: self.calls.append((label, data, k)))

    def tearDown(self):
        self.export.st.download_button = self._saved_dl

    def _open(self, i=0):
        """(workbook, data sheet, header→column index, download kwargs)."""
        self.assertGreater(len(self.calls), i, "export button was not rendered")
        label, data, kw = self.calls[i]
        self.assertEqual(label, "Export")
        self.assertTrue(callable(data), "workbook must build lazily")
        wb = load_workbook(io.BytesIO(data()))
        ws = wb.worksheets[0]
        col = {c.value: c.column for c in ws[1]}
        return wb, ws, col, kw

    @staticmethod
    def _source(wb) -> dict:
        return dict(_grid(wb["Source"]))


# ═══════════════════════════════════════════════════════════════════════
# ui/historicals.py — FDIC field codes → labeled ($K) / (%) columns
# ═══════════════════════════════════════════════════════════════════════

def _hist_row(repdte, **vals):
    row = {"CERT": 28489, "REPDTE": repdte,
           "Period": f"{str(repdte)[:4]}Q{(int(str(repdte)[4:6]) - 1) // 3 + 1}"}
    for f in ("ASSET", "DEP", "LNLSNET", "EQTOT", "NETINC", "NIMY", "ROA", "ROE",
              "EEFFR", "NCLNLSR", "IDT1CER", "INTINCY", "INTEXPY", "INTINC",
              "EINTEXP", "NONII", "NONIX", "NTLNLS", "SC", "COREDEP", "BRO",
              "LNLSGR"):
        row[f] = vals.get(f)
    return row


class TestHistoricalsExport(_SiteCase):
    # FDIC financials values as the API returns them: $THOUSANDS for dollar
    # fields (fetch_historical stores them verbatim; the charts do v*1000),
    # percent units for ratios.
    Q2 = _hist_row(20260630, NETINC=12_345, ASSET=2_500_000, NIMY=3.42, ROA=1.23,
                   EEFFR=58.7, NCLNLSR=None, IDT1CER=11.8)
    Q1 = _hist_row(20260331, NETINC=6_100, ASSET=2_450_000, NIMY=3.38, ROA=1.19,
                   EEFFR=59.4, NCLNLSR=0.44, IDT1CER=11.6)
    Q4_25 = _hist_row(20251231, NETINC=24_800, ASSET=2_400_000, NIMY=3.30, ROA=1.10,
                      EEFFR=60.1, NCLNLSR=0.40, IDT1CER=11.5)

    def _render(self, tab_index: int, rows):
        h = self.m["ui.historicals"]
        df = pd.DataFrame(rows)
        for col in df.columns:
            if col not in ("CERT", "REPDTE", "Period"):
                df[col] = pd.to_numeric(df[col], errors="coerce")
        with _patched(h, fetch_historical=lambda cert, quarters=12: df,
                      get_fdic_cert=lambda t: 28489,
                      get_name=lambda t: "Banner Corp",
                      lazy_tabs=lambda labels, key, default=0: labels[tab_index]):
            h.render_historicals("BANR")

    def test_quarterly_detail_labels_and_units(self):
        self._render(1, [self.Q2, self.Q1])           # "Quarterly Detail"
        wb, ws, col, kw = self._open()
        self.assertEqual(kw["file_name"], "historicals_quarterly_BANR.xlsx")
        self.assertEqual(kw["key"], "exp_historicals_quarterly_BANR")
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[0], "Period")
        # No raw FDIC field code survives as a header.
        for code in ("NETINC", "ROA", "NIMY", "EEFFR", "ASSET", "CERT", "REPDTE"):
            self.assertNotIn(code, headers)
        self.assertEqual(ws.cell(2, col["Period"]).value, "2026Q2")
        # $K: UNSCALED thousands under a ($K) header, #,##0 (never ×1000, never $)
        c = ws.cell(2, col["Net Income ($K)"])
        self.assertEqual(c.value, 12_345)
        self.assertEqual(c.number_format, "#,##0")
        self.assertEqual(ws.cell(2, col["Total Assets ($K)"]).value, 2_500_000)
        # %: percent units under a (%) header with the literal-% format
        c = ws.cell(2, col["NIM (%)"])
        self.assertEqual(c.value, 3.42)
        self.assertEqual(c.number_format, '0.00"%"')
        self.assertEqual(ws.cell(2, col["Efficiency Ratio (%)"]).value, 58.7)
        self.assertEqual(ws.cell(2, col["Efficiency Ratio (%)"]).number_format, '0.00"%"')
        # absent ratio → n/a, never 0
        self.assertEqual(ws.cell(2, col["NPL Ratio (%)"]).value, "n/a")
        self.assertEqual(ws.cell(3, col["NPL Ratio (%)"]).value, 0.44)
        self.assertEqual(ws.freeze_panes, "B2")
        src = self._source(wb)
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["Company"], "Banner Corp")
        self.assertEqual(src["FDIC cert"], 28489)
        self.assertEqual(src["Latest quarter"], "2026Q2")
        self.assertEqual(src["Source"], "FDIC/FFIEC bank-subsidiary financials")
        self.assertIn("($K)", src["Units"])
        self.assertIn("thousands", src["Units"])
        self.assertIn("percent units", src["Units"])

    def test_annual_summary_ytd_label_and_units(self):
        self._render(2, [self.Q2, self.Q1, self.Q4_25])   # "Annual Summary"
        wb, ws, col, kw = self._open()
        self.assertEqual(kw["file_name"], "historicals_annual_BANR.xlsx")
        self.assertEqual(kw["key"], "exp_historicals_annual_BANR")
        periods = [ws.cell(r, col["Period"]).value for r in (2, 3)]
        self.assertEqual(periods, ["2026 YTD", "2025"])
        # Flow = the year's latest YTD-cumulative print: 2026 YTD → Q2's 12,345;
        # 2025 → Q4's 24,800. Ratio = the year's last reported quarter.
        ni = col["Net Income ($K)"]
        self.assertEqual([ws.cell(2, ni).value, ws.cell(3, ni).value], [12_345, 24_800])
        self.assertEqual(ws.cell(2, ni).number_format, "#,##0")
        self.assertEqual(ws.cell(2, col["ROAA (%)"]).value, 1.23)
        self.assertEqual(ws.cell(3, col["ROAA (%)"]).value, 1.10)
        self.assertEqual(ws.cell(2, col["ROAA (%)"]).number_format, '0.00"%"')
        src = self._source(wb)
        self.assertIn("YTD", src["Note"])
        self.assertEqual(src["Latest quarter"], "2026Q2")

    def test_export_table_format_by_label_suffix(self):
        h = self.m["ui.historicals"]
        df = pd.DataFrame([self.Q2])
        out, fmt = h._export_table(df)
        self.assertEqual(list(out.columns)[0], "Period")
        self.assertEqual(fmt["Net Income ($K)"], "usd_k")
        self.assertEqual(fmt["NIM (%)"], "pct")
        self.assertTrue(all(v in ("usd_k", "pct") for v in fmt.values()),
                        "every HIST_METRICS label carries ($K) or (%)")


# ═══════════════════════════════════════════════════════════════════════
# ui/filings.py + ui/recent_documents.py — EDGAR filing records
# ═══════════════════════════════════════════════════════════════════════

def _filing(form, filed, report_date="", items="", is_earnings=False, size=45_123,
            acc="0001193125-26-077437"):
    return {"form": form, "date": filed, "report_date": report_date,
            "description": f"{form} primary doc", "items": items,
            "accession": acc,
            "url": f"https://www.sec.gov/Archives/edgar/data/946673/{acc.replace('-', '')}/doc.htm",
            "index_url": f"https://www.sec.gov/Archives/edgar/data/946673/{acc.replace('-', '')}/index.html",
            "is_earnings": is_earnings, "size": size}


_EXPECTED_FILING_HEADERS = ["Filed", "Form", "Description", "Report date", "Items",
                            "Accession", "URL", "Index URL", "Earnings-related",
                            "Size (bytes)"]


class TestFilingsExport(_SiteCase):
    FILINGS = [
        _filing("8-K", "2026-07-17", report_date="2026-07-16", items="2.02,9.01",
                is_earnings=True, size=45_123),
        _filing("10-Q", "2026-08-05", report_date="2026-06-30", size=3_210_987,
                acc="0000946673-26-000042"),
        _filing("8-K", "2026-06-02", report_date="", items="", size=0,
                acc="0001193125-26-060001"),
    ]

    def test_filings_table_export(self):
        fil = self.m["ui.filings"]
        with _patched(fil, _event_summaries=lambda t: {}):
            fil._render_filings_table(self.FILINGS, key_prefix="all",
                                      ticker="BANR", cik=946673)
        wb, ws, col, kw = self._open()
        self.assertEqual(kw["file_name"], "filings_all_BANR.xlsx")
        self.assertEqual(kw["key"], "exp_filings_all_BANR")
        self.assertEqual([c.value for c in ws[1]], _EXPECTED_FILING_HEADERS)
        # Dates parse to real date cells under yyyy-mm-dd
        c = ws.cell(2, col["Filed"])
        self.assertEqual(c.value, dt.datetime(2026, 7, 17))
        self.assertEqual(c.number_format, "yyyy-mm-dd")
        self.assertEqual(ws.cell(2, col["Report date"]).value, dt.datetime(2026, 7, 16))
        self.assertEqual(ws.cell(3, col["Report date"]).value, dt.datetime(2026, 6, 30))
        # Empty report date / items → n/a, not "" and not 1900-01-00
        self.assertEqual(ws.cell(4, col["Report date"]).value, "n/a")
        self.assertEqual(ws.cell(4, col["Items"]).value, "n/a")
        self.assertEqual(ws.cell(2, col["Items"]).value, "2.02,9.01")
        self.assertIs(ws.cell(2, col["Earnings-related"]).value, True)
        self.assertIs(ws.cell(3, col["Earnings-related"]).value, False)
        c = ws.cell(3, col["Size (bytes)"])
        self.assertEqual(c.value, 3_210_987)
        self.assertEqual(c.number_format, "#,##0")
        self.assertEqual(ws.cell(2, col["Accession"]).value, "0001193125-26-077437")
        self.assertTrue(str(ws.cell(2, col["URL"]).value).startswith("https://www.sec.gov/"))
        src = self._source(wb)
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["CIK"], 946673)
        self.assertEqual(src["Source"], "SEC EDGAR")
        self.assertEqual(src["Rows"], 3)

    def test_filings_export_frame_empty_keeps_headers(self):
        fil = self.m["ui.filings"]
        df, fmt = fil.filings_export_frame([])
        self.assertEqual(list(df.columns), _EXPECTED_FILING_HEADERS)
        self.assertEqual(fmt, {"Filed": "date", "Report date": "date",
                               "Size (bytes)": "int"})

    def test_recent_documents_export(self):
        rd = self.m["ui.recent_documents"]
        # Dated inside the stub's default "1 Year" range (selectbox → options[0]).
        filed = (date.today() - timedelta(days=40)).isoformat()
        filings = [_filing("10-K", filed, report_date="2025-12-31", size=9_876_543)]
        info = {"cik": 946673, "recent_filings": filings}
        with _patched(rd, get_filing_info=lambda cik, max_filings=1000: info,
                      fetch_key_exhibits=lambda cik: [],
                      get_fdic_cert=lambda t: None,
                      _handle_pdf_request=lambda *a, **k: None,
                      _skeleton=contextmanager(lambda: iter([None]))), \
                _fake_module("data.fmp_transcripts", get_transcript_dates=lambda t: []):
            rd._render_body("BANR", 946673)
        wb, ws, col, kw = self._open()
        self.assertEqual(kw["file_name"], "recent_documents_BANR.xlsx")
        self.assertEqual(kw["key"], "exp_rd_BANR")
        self.assertEqual([c.value for c in ws[1]], _EXPECTED_FILING_HEADERS)
        y, mo, d = (int(p) for p in filed.split("-"))
        self.assertEqual(ws.cell(2, col["Filed"]).value, dt.datetime(y, mo, d))
        self.assertEqual(ws.cell(2, col["Filed"]).number_format, "yyyy-mm-dd")
        self.assertEqual(ws.cell(2, col["Report date"]).value, dt.datetime(2025, 12, 31))
        self.assertEqual(ws.cell(2, col["Size (bytes)"]).value, 9_876_543)
        src = self._source(wb)
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["CIK"], 946673)
        self.assertEqual(src["Date range"], "1 Year")
        self.assertEqual(src["Source"], "SEC EDGAR")


# ═══════════════════════════════════════════════════════════════════════
# ui/transactions.py — universe deal comps
# ═══════════════════════════════════════════════════════════════════════

# Hand-computed from whole-dollar inputs (the snapshot stores what
# compute_multiples produced):
#   value 191,100,000 ÷ TBV 80,000,000           = 2.38875x
#   value ÷ comp assets 900,000,000              = 0.212333… → 21.2333…%
#   (value − TBV) ÷ core deposits 600,000,000    = 0.185166… → 18.5166…%
_PRICED = {
    "buyer_ticker": "BANR", "buyer_name": "Banner Corp", "buyer_cert": 28489,
    "target_name": "Skagit Bank", "target_cert": 17874, "status": "completed",
    "announce_date": "2026-03-02", "completion_date": "2026-07-01",
    "termination_date": None,
    "value_usd": 191_100_000, "value_basis": "stated", "value_note": None,
    "announce_url": "https://www.sec.gov/Archives/edgar/data/946673/000094667326000010/ex99.htm",
    "target_assets": 920_000_000, "target_assets_repdte": "2026-06-30",
    "tbv_usd": 80_000_000, "tbv_basis": "bank-sub", "tbv_asof": "2025-12-31",
    "p_tbv": 191_100_000 / 80_000_000,
    "price_assets": 191_100_000 / 900_000_000,
    "core_dep_premium": (191_100_000 - 80_000_000) / 600_000_000,
    "comp_assets": 900_000_000, "flagged": None,
}
_UNPRICED = {
    "buyer_ticker": "EWBC", "buyer_name": "East West Bancorp", "buyer_cert": 31628,
    "target_name": "Private Bank NA", "target_cert": 55555, "status": "completed",
    "announce_date": None, "completion_date": "2025-11-14", "termination_date": None,
    "value_usd": None, "value_basis": None, "value_note": None, "announce_url": None,
    "target_assets": 450_000_000, "target_assets_repdte": "2025-09-30",
    "tbv_usd": None, "tbv_basis": None, "tbv_asof": None,
    "p_tbv": None, "price_assets": None, "core_dep_premium": None,
    "comp_assets": None, "flagged": None,
}
_SNAPSHOT = {"built_at": "2026-09-21T06:12:03", "banks_covered": 2,
             "deals_total": 2, "deals_priced": 1, "deals": [_PRICED, _UNPRICED]}


class TestDealCompsExport(_SiteCase):

    def test_deal_comps_units_and_na(self):
        tx = self.m["ui.transactions"]
        with _patched(tx, _cert_ticker_map=lambda: {}), \
                _fake_module("data.deal_comps",
                             get_comps_snapshot=lambda: dict(_SNAPSHOT)):
            tx._render_comps()
        wb, ws, col, kw = self._open()
        # stem: filters at their defaults ("All years" / "All") + snapshot date
        self.assertEqual(kw["file_name"], "deal_comps_all_2026-09-21.xlsx")
        self.assertEqual(kw["key"], "comps_export")
        # No raw snapshot key survives as a header
        for raw in ("value_usd", "p_tbv", "price_assets", "core_dep_premium",
                    "comp_assets", "tbv_usd"):
            self.assertNotIn(raw, col)
        # ── priced deal (row 2) ──
        c = ws.cell(2, col["Deal value ($)"])
        self.assertEqual(c.value, 191_100_000)            # whole dollars, unscaled
        self.assertEqual(c.number_format, "$#,##0")
        self.assertEqual(ws.cell(2, col["Target assets ($)"]).value, 920_000_000)
        self.assertEqual(ws.cell(2, col["Target assets ($)"]).number_format, "$#,##0")
        self.assertEqual(ws.cell(2, col["Comp assets ($)"]).value, 900_000_000)
        self.assertEqual(ws.cell(2, col["TBV ($)"]).value, 80_000_000)
        self.assertEqual(ws.cell(2, col["TBV basis"]).value, "bank-sub")
        c = ws.cell(2, col["P/TBV (x)"])
        self.assertAlmostEqual(c.value, 2.38875, places=6)
        self.assertEqual(c.number_format, '0.00"x"')
        c = ws.cell(2, col["P/Assets (%)"])
        self.assertAlmostEqual(c.value, 21.2333333, places=5)   # fraction ×100, as on screen
        self.assertEqual(c.number_format, '0.00"%"')
        c = ws.cell(2, col["Core deposit premium (%)"])
        self.assertAlmostEqual(c.value, 18.5166667, places=5)
        self.assertEqual(c.number_format, '0.00"%"')
        c = ws.cell(2, col["Announced"])
        self.assertEqual(c.value, dt.datetime(2026, 3, 2))
        self.assertEqual(c.number_format, "yyyy-mm-dd")
        self.assertEqual(ws.cell(2, col["Completed"]).value, dt.datetime(2026, 7, 1))
        self.assertEqual(ws.cell(2, col["TBV as of"]).value, dt.datetime(2025, 12, 31))
        self.assertEqual(ws.cell(2, col["Terminated"]).value, "n/a")
        c = ws.cell(2, col["Target FDIC cert"])
        self.assertEqual(c.value, 17874)
        self.assertEqual(c.number_format, "#,##0")
        self.assertEqual(ws.cell(2, col["Buyer"]).value, "BANR")
        self.assertEqual(ws.cell(2, col["Status"]).value, "completed")
        # ── unpriced deal (row 3): value and every multiple are n/a, never 0 ──
        for h in ("Deal value ($)", "P/TBV (x)", "P/Assets (%)",
                  "Core deposit premium (%)", "TBV ($)", "Comp assets ($)",
                  "Announced", "Announcement URL"):
            self.assertEqual(ws.cell(3, col[h]).value, "n/a", h)
        self.assertEqual(ws.cell(3, col["Target assets ($)"]).value, 450_000_000)
        self.assertEqual(ws.cell(3, col["Completed"]).value, dt.datetime(2025, 11, 14))
        src = self._source(wb)
        self.assertEqual(src["Snapshot built"], "2026-09-21")
        self.assertEqual(src["Announced since"], "All years")
        self.assertEqual(src["Status filter"], "All")
        self.assertIn("EDGAR", src["Source"])
        self.assertIn("whole US dollars", src["Units"])
        self.assertIn("percent units", src["Units"])
        self.assertIn("raw ratio", src["Units"])
        self.assertEqual(src["Rows"], 2)


# ═══════════════════════════════════════════════════════════════════════
# ui/peer_rank.py — scorecard + leaderboard
# ═══════════════════════════════════════════════════════════════════════

def _ctx():
    """Fresh each call — render_peer_rank pops _meta from it."""
    return {
        "_meta": {"tier": "$10B–$50B", "cohort_size": 18, "mode": "size"},
        "roaa": {"value": 1.23, "percentile": 72.2, "raw": 72.2, "median": 1.05,
                 "n": 18, "rank": 6, "out_of": 18, "higher_better": True, "label": "ROAA"},
        "npl_ratio": {"value": 0.31, "percentile": 88.9, "raw": 11.1, "median": 0.55,
                      "n": 18, "rank": 3, "out_of": 18, "higher_better": False,
                      "label": "NPL Ratio"},
    }


class TestPeerRankExport(_SiteCase):
    COHORT = [{"ticker": "BANR", "roaa": 1.23, "npl_ratio": 0.31},
              {"ticker": "EWBC", "roaa": 1.61, "npl_ratio": 0.55}]

    def _render(self):
        pr = self.m["ui.peer_rank"]
        with _patched(pr, metric_percentile_context=lambda *a, **k: _ctx(),
                      get_peer_group_for_bank=lambda t, m, mode="size": m,
                      get_name=lambda t: f"{t} Bancorp"):
            pr.render_peer_rank("BANR", self.COHORT)

    def test_every_ranked_metric_is_a_percent(self):
        # The scorecard headers say (%) and _metric_row prints "x.xx%" for
        # every row — both are honest only while every RANK_GROUPS metric is a
        # config.METRICS pct. Adding a non-percent metric must fail here.
        pr = self.m["ui.peer_rank"]
        from config import METRICS_BY_KEY
        for k in pr._ALL_KEYS:
            self.assertEqual(METRICS_BY_KEY[k]["format"], "pct", k)

    def test_scorecard_export(self):
        self._render()
        wb, ws, col, kw = self._open(0)
        self.assertEqual(kw["file_name"], "peer_rank_BANR.xlsx")
        self.assertEqual(kw["key"], "exp_peer_rank_BANR")
        self.assertEqual(_grid(ws), [
            ["Metric", "Value (%)", "Percentile", "Rank", "Out of", "Peer median (%)"],
            ["ROAA", 1.23, 72.2, 6, 18, 1.05],
            ["NPL ratio", 0.31, 88.9, 3, 18, 0.55],
        ])
        self.assertEqual(ws.cell(2, col["Value (%)"]).number_format, '0.00"%"')
        self.assertEqual(ws.cell(2, col["Peer median (%)"]).number_format, '0.00"%"')
        self.assertEqual(ws.cell(2, col["Percentile"]).number_format, "#,##0.00")
        self.assertEqual(ws.cell(2, col["Rank"]).number_format, "#,##0")
        self.assertEqual(ws.cell(2, col["Out of"]).number_format, "#,##0")
        src = self._source(wb)
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["Company"], "BANR Bancorp")
        self.assertIn("18 tracked", src["Peer set"])
        self.assertIn("Asset size", src["Peer set"])
        self.assertIn("goodness", src["Note"])

    def test_leaderboard_export_keeps_contract_and_adds_format(self):
        self._render()
        # selectbox → first ranked key ("roaa") → leaderboard is the 2nd export
        wb, ws, col, kw = self._open(1)
        self.assertEqual(kw["file_name"], "peer_leaderboard_BANR_roaa.xlsx")
        self.assertEqual(kw["key"], "exp_peer_leaderboard_BANR_roaa")
        rows = _grid(ws)
        self.assertEqual(rows[0][:4], ["Rank", "Ticker", "Bank", "Value"])
        self.assertEqual(rows[1][:4], [1, "EWBC", "EWBC Bancorp", 1.61])   # higher is better
        self.assertEqual(rows[2][:4], [2, "BANR", "BANR Bancorp", 1.23])
        # roaa is config format pct / 2 decimals → percent-unit format
        self.assertEqual(ws.cell(2, col["Value"]).number_format, '0.00"%"')
        self.assertEqual(ws.cell(2, col["Rank"]).number_format, "#,##0")
        src = self._source(wb)
        self.assertIn("ROAA", src["Metric"])
        self.assertIn("higher is better", src["Metric"])
        self.assertIn("2 banks", src["Peer set"])
        self.assertEqual(src["Ticker"], "BANR")


# ═══════════════════════════════════════════════════════════════════════
# ui/macro.py — FDIC national rates + the FRED print board
# ═══════════════════════════════════════════════════════════════════════

class TestMacroExports(_SiteCase):

    def test_fdic_national_rates_export(self):
        macro = self.m["ui.macro"]
        rates = {"asof": "2026-09-15",
                 "savings": {"rate_pct": 0.42, "cap_pct": 1.17},
                 "cd_12mo": {"rate_pct": 1.86, "cap_pct": 2.61}}
        hist = [{"asof": "2026-08-18", "savings": {"rate_pct": 0.41},
                 "cd_12mo": {"rate_pct": 1.90}},
                {"asof": "2026-09-15", "savings": {"rate_pct": 0.42},
                 "cd_12mo": {"rate_pct": 1.86}}]
        with _patched(macro, latest_value=lambda sid: 4.33,
                      fetch_series=lambda sid, years=5: pd.DataFrame()), \
                _fake_module("data.national_rates",
                             get_national_rates=lambda: rates,
                             get_national_rate_history=lambda weeks=104: list(hist)):
            macro._render_funding_deposits()
        wb, ws, col, kw = self._open()
        self.assertEqual(kw["file_name"], "fdic_national_rates_2026-09-15.xlsx")
        self.assertEqual(kw["key"], "fdic_rates_export")
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers[0], "As of")
        self.assertEqual(headers[1:], [f"{lbl} (%)" for _f, lbl in macro._DEPOSIT_PRODUCTS])
        for code in ("asof", "savings", "cd_12mo", "mmda"):
            self.assertNotIn(code, headers)
        c = ws.cell(2, col["As of"])
        self.assertEqual(c.value, dt.datetime(2026, 8, 18))
        self.assertEqual(c.number_format, "yyyy-mm-dd")
        c = ws.cell(2, col["Savings (%)"])
        self.assertEqual(c.value, 0.41)                    # percent units, as shown
        self.assertEqual(c.number_format, '0.00"%"')
        self.assertEqual(ws.cell(3, col["12-Month CD (%)"]).value, 1.86)
        self.assertEqual(ws.cell(2, col["Money Market (%)"]).value, "n/a")   # unpublished → n/a
        src = self._source(wb)
        self.assertEqual(src["Data as of"], "2026-09-15")
        self.assertIn("FDIC national rates", src["Source"])
        self.assertIn("percent units", src["Units"])

    def test_print_board_export_mixed_units_stay_plain(self):
        macro = self.m["ui.macro"]
        rows = [
            {"key": "cpi", "label": "CPI", "series_id": "CPIAUCSL", "basis": "yoy_pct",
             "freq": "M", "theme": "Inflation", "favorable": "down",
             "latest": 2.9, "prior": 3.1, "delta": -0.2, "zscore": 0.4,
             "as_of": pd.Timestamp("2026-08-31"), "spark": [3.1, 2.9]},
            {"key": "payrolls", "label": "Nonfarm payrolls", "series_id": "PAYEMS",
             "basis": "mom_chg_k", "freq": "M", "theme": "Labor", "favorable": "up",
             "latest": 142.0, "prior": 98.0, "delta": 44.0, "zscore": -0.3,
             "as_of": pd.Timestamp("2026-09-05"), "spark": [98.0, 142.0]},
            {"key": "claims", "label": "Initial claims", "series_id": "ICSA",
             "basis": "level_k", "freq": "W", "theme": "Labor", "favorable": "down",
             "latest": None, "prior": None, "delta": None, "zscore": None,
             "as_of": None, "spark": []},
        ]
        with _patched(macro, _cached_print_board=lambda: [dict(r) for r in rows],
                      _render_macro_grid=lambda: None,
                      _render_indicator_explorer=lambda: None), \
                _fake_module("data.econ_calendar",
                             get_recent_releases=lambda **k: [],
                             get_upcoming_releases=lambda **k: []):
            macro._render_economy_calendar()
        wb, ws, col, kw = self._open()
        self.assertEqual(kw["file_name"], "macro_print_board_2026-09-05.xlsx")
        self.assertEqual(kw["key"], "macro_print_board_export")
        self.assertEqual([c.value for c in ws[1]],
                         ["Theme", "Indicator", "Basis", "Latest", "Prior", "Δ",
                          "Z-score", "As of", "FRED series"])
        # A YoY-% row and a thousands row share the column: values raw, the
        # format PLAIN for both (a % format here would be a plausible-wrong 142%).
        self.assertEqual(ws.cell(2, col["Latest"]).value, 2.9)
        self.assertEqual(ws.cell(3, col["Latest"]).value, 142.0)
        for r in (2, 3):
            for h in ("Latest", "Prior", "Δ", "Z-score"):
                self.assertEqual(ws.cell(r, col[h]).number_format, "#,##0.00", (r, h))
        self.assertEqual(ws.cell(2, col["Basis"]).value, "yoy_pct")
        self.assertEqual(ws.cell(3, col["Basis"]).value, "mom_chg_k")
        self.assertEqual(ws.cell(3, col["Δ"]).value, 44.0)
        c = ws.cell(2, col["As of"])
        self.assertEqual(c.value, dt.datetime(2026, 8, 31))
        self.assertEqual(c.number_format, "yyyy-mm-dd")
        # failed series → n/a everywhere numeric, never 0
        for h in ("Latest", "Prior", "Δ", "Z-score", "As of"):
            self.assertEqual(ws.cell(4, col[h]).value, "n/a", h)
        self.assertEqual(ws.cell(4, col["FRED series"]).value, "ICSA")
        src = self._source(wb)
        self.assertEqual(src["Data as of"], "2026-09-05")
        self.assertEqual(src["Source"], "FRED")
        self.assertIn("Basis", src["Value units"])
        self.assertIn("thousands", src["Value units"])
        self.assertNotIn("percent units", src["Units"])   # no pct format was applied


if __name__ == "__main__":
    unittest.main()
