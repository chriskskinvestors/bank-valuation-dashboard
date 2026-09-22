"""ui/earnings.py table exports on the shared .xlsx exporter (ui/export.py;
owner directive 2026-09-22: exports are Excel from the RAW frame, never
display strings).

Each test drives the real render function with small fixture data, captures
the deferred workbook callable table_export hands st.download_button, loads
the bytes with openpyxl and asserts against HAND-COMPUTED values: header
labels carry units, numeric cells are numbers (not "$1.21" / "3.42%" /
"3/4"), the Excel number format per column, "n/a" where the fixture had no
value, and the Source sheet's provenance rows.

Former display-string sites pinned here: consistency metrics, per-firm
estimate matrix, beat/miss summary, surprise rankings (Consensus/Actual),
sector metric breakdown. Raw sites: per-bank surprise history, consensus vs
actual, earnings calendar, reported-results board.

Stub discipline: the render modules' own ``st`` bindings are patched with a
PRIVATE no-op stub per test (mock.patch.object) — nothing here touches
sys.modules or the shared package stub, so this module composes with every
other suite under discovery (a sys.modules swap here broke
test_merger_proximity_ui's recorders, 2026-09-22).

Run: python -m unittest tests.test_export_sites_earnings
"""
from __future__ import annotations

import datetime as dt
import io
import sys
import types
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).parent.parent))

from openpyxl import load_workbook  # noqa: E402


def _noop(*a, **k):
    return None


class _Ctx:
    """Context-manager placeholder (columns / container / st.empty slot):
    every attribute is a no-op."""

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def __getattr__(self, name):
        return _noop


class _StubSt:
    """Private streamlit stand-in for the earnings render paths: unknown
    widgets are no-ops; layout widgets return contexts; selectbox returns
    the first option (the page's default filter)."""

    session_state: dict = {}
    query_params: dict = {}
    column_config = types.SimpleNamespace(LinkColumn=lambda *a, **k: None)

    def __getattr__(self, name):
        return _noop

    def columns(self, spec, **k):
        n = spec if isinstance(spec, int) else len(spec)
        return [_Ctx() for _ in range(n)]

    def container(self, *a, **k):
        return _Ctx()

    expander = spinner = popover = empty = container

    def selectbox(self, label, options=None, **k):
        return options[0] if options else None


def _fake_name(t):
    return f"{t} Bancorp"


class _ExportSite(unittest.TestCase):
    """Binds the private stub onto every module the render paths reach
    through, and a download_button capture onto ui.export."""

    @classmethod
    def setUpClass(cls):
        import ui.chrome, ui.export, ui.states, ui.tables  # noqa: E401
        import ui.earnings
        cls.E, cls.X = ui.earnings, ui.export
        cls._st_modules = (ui.earnings, ui.chrome, ui.states, ui.tables)

    def setUp(self):
        self.calls = []          # [(label, data, kwargs)] from download_button
        stub = _StubSt()
        fake_export_st = types.SimpleNamespace(
            download_button=lambda label, data, **k: self.calls.append((label, data, k)),
            container=lambda *a, **k: _Ctx())
        for mod, obj in [(m, stub) for m in self._st_modules] + [(self.X, fake_export_st)]:
            p = mock.patch.object(mod, "st", obj)
            p.start()
            self.addCleanup(p.stop)

    # ── helpers ──────────────────────────────────────────────────────
    def _book(self):
        self.assertEqual(len(self.calls), 1, "expected exactly one Export control")
        _label, data, kw = self.calls[0]
        self.assertTrue(callable(data), "workbook must build lazily, on click")
        wb = load_workbook(io.BytesIO(data()))
        return wb, wb.worksheets[0], kw

    @staticmethod
    def _grid(ws):
        return [[c.value for c in r] for r in ws.iter_rows()]

    @staticmethod
    def _source(wb):
        return dict((r[0].value, r[1].value) for r in wb["Source"].iter_rows())

    def _assert_na(self, ws, coord):
        self.assertEqual(ws[coord].value, "n/a")
        self.assertTrue(ws[coord].font.italic)

    # ── fixtures shared by the consensus-driven sites ────────────────
    ALL_CONSENSUS = {
        "AAA": [{"period": "2026Q1", "n_firms": 1, "firms": ["F"], "metric_count": 2}],
        "BBB": [{"period": "2026Q1", "n_firms": 1, "firms": ["F"], "metric_count": 1}],
    }
    CONSENSUS = {
        "AAA": {"metrics": [
            {"key": "eps", "name": "EPS", "value": 1.00, "unit": "$"},
            {"key": "nim", "name": "NIM", "value": 3.40, "unit": "%"}]},
        "BBB": {"metrics": [
            {"key": "eps", "name": "EPS", "value": 2.00, "unit": "$"}]},
    }

    def _consensus_patches(self, actuals: dict):
        E = self.E
        return (
            mock.patch.object(E, "compile_consensus",
                              lambda t, p: self.CONSENSUS.get(t)),
            mock.patch.object(E, "period_actuals",
                              lambda t, p: actuals.get(t)),
            mock.patch.object(E, "get_name", _fake_name),
        )


class TestSurpriseHistory(_ExportSite):
    """Site: _render_surprise_history_grid (raw yfinance history)."""

    def test_numeric_eps_and_surprise_with_dates(self):
        past = [
            {"date": "2026-07-15", "eps_estimate": 1.20, "eps_actual": 1.26,
             "surprise_pct": 5.0},
            {"date": "2026-10-15", "eps_estimate": 1.30, "eps_actual": None,
             "surprise_pct": None},
        ]
        self.E._render_surprise_history_grid("BANR", past)
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "earnings_surprises_BANR.xlsx")
        self.assertEqual(kw["key"], "exp_earnings_surprises_BANR")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Date", "EPS Est ($)", "EPS Act ($)", "Surprise (%)"])
        self.assertEqual(g[1], [dt.datetime(2026, 7, 15), 1.2, 1.26, 5.0])
        self.assertEqual(g[2][:2], [dt.datetime(2026, 10, 15), 1.3])
        self._assert_na(ws, "C3")
        self._assert_na(ws, "D3")
        self.assertEqual(ws["A2"].number_format, "yyyy-mm-dd")
        self.assertEqual(ws["B2"].number_format, "$#,##0.00")
        self.assertEqual(ws["C2"].number_format, "$#,##0.00")
        self.assertEqual(ws["D2"].number_format, '0.00"%"')
        src = self._source(wb)
        self.assertEqual(src["Ticker"], "BANR")
        self.assertIn("Surprise History", src["Page"])
        self.assertIn("Yahoo Finance", src["Source"])


class TestConsensusVsActual(_ExportSite):
    """Site: _render_comparison_table — the REAL compare_consensus_to_actual
    output (contract pinned: Δ = actual − consensus in the canonical unit)."""

    def test_raw_values_unit_column_and_named_file(self):
        from data.consensus import compare_consensus_to_actual
        consensus = {"metrics": [
            {"key": "eps", "name": "EPS", "value": 1.20, "unit": "$",
             "low": 1.15, "high": 1.25, "n_firms": 2},
            {"key": "nii", "name": "NII", "value": 150.0, "unit": "$M",
             "n_firms": 1},
            {"key": "revenue", "name": "Revenue", "value": 200.0, "unit": "$M"},
        ]}
        # nii actual is RAW dollars in the metrics dict → $M in the comparison.
        comparison = compare_consensus_to_actual(
            consensus, {"eps": 1.26, "net_interest_income": 147_000_000.0})
        self.E._render_comparison_table(comparison, ticker="BANR", period="2026Q1")
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "consensus_vs_actual_BANR_2026Q1.xlsx")
        self.assertEqual(kw["key"], "exp_consensus_vs_actual")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Metric", "Unit", "Consensus", "Low", "High",
                                "Firms", "Actual", "Δ", "Δ (%)", "Result"])
        # eps: 1.26 − 1.20 = +0.06 → +5.0% → beat
        self.assertEqual(g[1][:7], ["Earnings Per Share", "$", 1.2, 1.15, 1.25, 2, 1.26])
        self.assertAlmostEqual(g[1][7], 0.06, places=9)
        self.assertAlmostEqual(g[1][8], 5.0, places=9)
        self.assertEqual(g[1][9], "beat")
        # nii: 147 − 150 = −3 ($M) → −2.0% → miss; single firm → no range
        self.assertEqual(g[2][:3], ["Net Interest Income", "$M", 150.0])
        self._assert_na(ws, "D3")
        self._assert_na(ws, "E3")
        self.assertEqual(g[2][5:], [1, 147.0, -3.0, -2.0, "miss"])
        # revenue has no actual counterpart → everything after Consensus n/a
        self.assertEqual(g[3][:3], ["Revenue", "$M", 200.0])
        for col in "DEFGHIJ":
            self._assert_na(ws, f"{col}4")
        for coord in ("C2", "D2", "E2", "G2", "H2"):
            self.assertIsInstance(ws[coord].value, float)
            self.assertEqual(ws[coord].number_format, "#,##0.00")
        self.assertEqual(ws["F2"].number_format, "#,##0")
        self.assertEqual(ws["I2"].number_format, '0.00"%"')
        src = self._source(wb)
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["Period"], "2026Q1")
        self.assertIn("Consensus vs Actual", src["Page"])
        self.assertIn("SEC companyfacts", src["Source"])
        self.assertIn("$M = millions", src["Value units"])


class TestConsistencyMetrics(_ExportSite):
    """Site: heat-map Consistency Metrics — formerly exported "1/3 (33%)",
    "+2.7%", "6.4pp" strings."""

    ESTIMATES = {
        "AAA": {"earnings_history": [
            {"date": "2026-07-15", "eps_estimate": 1.00, "eps_actual": 1.10, "surprise_pct": 10.0},
            {"date": "2026-04-15", "eps_estimate": 1.00, "eps_actual": 0.98, "surprise_pct": -2.0},
            {"date": "2026-01-15", "eps_estimate": 1.00, "eps_actual": 1.00, "surprise_pct": 0.0},
        ]},
        "BBB": {"earnings_history": [
            {"date": "2026-07-16", "eps_estimate": 2.00, "eps_actual": 2.10, "surprise_pct": 5.0},
        ]},
    }

    def test_counts_rates_and_volatility_are_numbers(self):
        import data.bank_mapping as bm
        import data.estimates as de
        with mock.patch.object(de, "fetch_all_estimates", lambda tks: self.ESTIMATES), \
                mock.patch.object(bm, "get_name", _fake_name), \
                mock.patch.object(self.E, "get_name", _fake_name):
            self.E._render_surprise_heatmap(["AAA", "BBB"])
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "earnings_consistency_metrics.xlsx")
        self.assertEqual(kw["key"], "exp_earnings_consistency_metrics")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Ticker", "Bank", "Quarters", "Beats", "Misses",
                                "Inline", "Beat Rate (%)", "Avg Surprise (%)",
                                "Volatility (pp)", "Last Qtr Surprise (%)"])
        # AAA (sorted first: latest surprise 10 > 5): surprises [10, −2, 0]
        #   beats 1 (>1), misses 1 (<−1), inline 1; rate 1/3 = 33.33%;
        #   mean 8/3 = 2.6667; sample sd = sqrt(((10−8/3)² + (−2−8/3)² +
        #   (0−8/3)²)/2) = sqrt(82.6667/2) = 6.4291
        self.assertEqual(g[1][:6], ["AAA", "AAA Bancorp", 3, 1, 1, 1])
        self.assertAlmostEqual(g[1][6], 100 / 3, places=9)
        self.assertAlmostEqual(g[1][7], 8 / 3, places=9)
        self.assertAlmostEqual(g[1][8], 6.429100507, places=6)
        self.assertEqual(g[1][9], 10.0)
        # BBB: one quarter → 1/1 beat, no volatility (needs ≥2)
        self.assertEqual(g[2][:8], ["BBB", "BBB Bancorp", 1, 1, 0, 0, 100.0, 5.0])
        self._assert_na(ws, "I3")
        self.assertEqual(g[2][9], 5.0)
        for coord in ("C2", "D2", "E2", "F2"):
            self.assertIsInstance(ws[coord].value, int)
            self.assertEqual(ws[coord].number_format, "#,##0")
        for coord in ("G2", "H2", "J2"):
            self.assertEqual(ws[coord].number_format, '0.0"%"')
        self.assertEqual(ws["I2"].number_format, "#,##0.00")
        src = self._source(wb)
        self.assertIn("Consistency Metrics", src["Page"])
        self.assertEqual(src["Banks"], 2)
        self.assertIn("Yahoo Finance", src["Source"])


class TestFirmMatrix(_ExportSite):
    """Site: _render_firm_matrix — formerly exported "$1.21" / "3.40%–3.50%"."""

    def test_per_firm_raw_values_with_unit_column(self):
        detail = {
            "ticker": "BANR", "period": "2026Q1", "firms": ["Firm A", "Firm B"],
            "metrics": [
                {"key": "nim", "name": "Net Interest Margin", "unit": "%",
                 "by_firm": {"Firm A": 3.40, "Firm B": 3.50},
                 "mean": 3.45, "low": 3.40, "high": 3.50, "n": 2},
                {"key": "eps", "name": "Earnings Per Share", "unit": "$",
                 "by_firm": {"Firm A": 1.21},
                 "mean": 1.21, "low": 1.21, "high": 1.21, "n": 1},
            ]}
        self.E._render_firm_matrix(detail, "BANR_2026Q1")
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "estimates_BANR_2026Q1.xlsx")
        self.assertEqual(kw["key"], "exp_estimates_BANR_2026Q1")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Metric", "Unit", "Firm A", "Firm B", "Mean",
                                "Low", "High", "Firms"])
        # sorted by metric name → EPS first; Firm B gave no EPS → n/a
        self.assertEqual(g[1][:3], ["Earnings Per Share", "$", 1.21])
        self._assert_na(ws, "D2")
        self.assertEqual(g[1][4:], [1.21, 1.21, 1.21, 1])
        self.assertEqual(g[2], ["Net Interest Margin", "%", 3.4, 3.5, 3.45,
                                3.4, 3.5, 2])
        for coord in ("C3", "D3", "E3", "F3", "G3"):
            self.assertIsInstance(ws[coord].value, float)
            self.assertEqual(ws[coord].number_format, "#,##0.00")
        self.assertIsInstance(ws["H3"].value, int)
        self.assertEqual(ws["H3"].number_format, "#,##0")
        src = self._source(wb)
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["Period"], "2026Q1")
        self.assertEqual(src["Firms"], "Firm A, Firm B")
        self.assertIn("Uploaded consensus", src["Source"])
        self.assertIn("Unit column", src["Value units"])


class TestBeatMissSummary(_ExportSite):
    """Site: _render_beat_miss_summary — formerly exported "+$0.05",
    "-0.10%", "1/2" strings."""

    def test_counts_and_numeric_deltas(self):
        # AAA reported: eps 1.05 vs 1.00 → +0.05 (+5%) beat;
        #               nim 3.30 vs 3.40 → −0.10 (−2.94%) miss.
        # BBB not reported (period_actuals None) → 0 comparable metrics.
        actuals = {"AAA": {"eps": 1.05, "nim": 3.30}}
        p1, p2, p3 = self._consensus_patches(actuals)
        with p1, p2, p3:
            self.E._render_beat_miss_summary(self.ALL_CONSENSUS)
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "beat_miss_summary.xlsx")
        self.assertEqual(kw["key"], "exp_beat_miss_summary")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Ticker", "Bank", "Period", "Metrics", "Beats",
                                "Misses", "Inline", "EPS Δ ($)", "NIM Δ (%)"])
        self.assertEqual(g[1][:7], ["AAA", "AAA Bancorp", "2026Q1", 2, 1, 1, 0])
        self.assertAlmostEqual(g[1][7], 0.05, places=9)
        self.assertAlmostEqual(g[1][8], -0.10, places=9)
        self.assertEqual(g[2][:7], ["BBB", "BBB Bancorp", "2026Q1", 0, 0, 0, 0])
        self._assert_na(ws, "H3")
        self._assert_na(ws, "I3")
        self.assertIsInstance(ws["D2"].value, int)
        self.assertEqual(ws["D2"].number_format, "#,##0")
        self.assertEqual(ws["H2"].number_format, "$#,##0.00")
        self.assertEqual(ws["I2"].number_format, '0.00"%"')
        src = self._source(wb)
        self.assertIn("Beat / Miss Summary", src["Page"])
        self.assertIn("latest uploaded period", src["Period"])
        self.assertIn("SEC companyfacts", src["Source"])


class TestSurpriseRankings(_ExportSite):
    """Site: _render_surprise_rankings — Consensus/Actual were "$1.00"
    strings; Unit column carries the per-row unit."""

    def test_uploaded_and_yfinance_rows_are_numeric(self):
        actuals = {"AAA": {"eps": 1.05, "nim": 3.30}}
        yf = {"AAA": {"earnings_history": [
            {"date": "2026-07-15", "eps_estimate": 1.00, "eps_actual": 1.10,
             "surprise_pct": 10.0}]}}
        p1, p2, p3 = self._consensus_patches(actuals)
        with p1, p2, p3, \
                mock.patch.object(self.E, "fetch_all_estimates", lambda tks: yf):
            self.E._render_surprise_rankings(
                {"AAA": self.ALL_CONSENSUS["AAA"]}, ["AAA"])
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "surprise_rankings.xlsx")
        self.assertEqual(kw["key"], "exp_surprise_rankings")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Ticker", "Bank", "Metric", "Unit", "Period",
                                "Consensus", "Actual", "Surprise (%)", "Result",
                                "Source"])
        # sorted by |surprise| desc: yfinance 10% > eps 5% > nim −2.94%
        self.assertEqual(g[1], ["AAA", "AAA Bancorp", "EPS", "$", "2026-07-15",
                                1.0, 1.1, 10.0, "beat", "Yahoo Finance"])
        self.assertEqual(g[2][:7], ["AAA", "AAA Bancorp", "Earnings Per Share",
                                    "$", "2026Q1", 1.0, 1.05])
        self.assertAlmostEqual(g[2][7], 5.0, places=9)
        self.assertEqual(g[2][8:], ["beat", "Uploaded"])
        self.assertEqual(g[3][:7], ["AAA", "AAA Bancorp", "Net Interest Margin",
                                    "%", "2026Q1", 3.4, 3.3])
        self.assertAlmostEqual(g[3][7], -0.10 / 3.40 * 100, places=9)
        self.assertEqual(g[3][8:], ["miss", "Uploaded"])
        # (1.00 round-trips through the sheet XML as the integer 1 — still a
        # number, never the "$1.00" string the display frame carries.)
        for coord in ("F2", "G2"):
            self.assertIsInstance(ws[coord].value, (int, float))
            self.assertEqual(ws[coord].number_format, "#,##0.00")
        self.assertEqual(ws["H2"].number_format, '0.00"%"')
        src = self._source(wb)
        self.assertIn("Surprise Magnitude Rankings", src["Page"])
        self.assertEqual(src["Filter"], "All · All Metrics")
        self.assertIn("Yahoo Finance", src["Source"])
        self.assertIn("Unit column", src["Value units"])


class TestSectorBreakdown(_ExportSite):
    """Site: sector Per-Metric Breakdown — formerly exported "50%" / "+2.8%"."""

    def test_beat_rate_and_avg_surprise_numeric(self):
        # AAA: eps +5.0% beat, nim −2.94% miss. BBB: eps 2.01 vs 2.00 →
        # +0.5% → inline. eps: 2 banks, 1 beat, 1 inline, rate 50%,
        # avg (5.0 + 0.5)/2 = 2.75. nim: 1 bank, 1 miss, rate 0%, avg −2.94.
        actuals = {"AAA": {"eps": 1.05, "nim": 3.30}, "BBB": {"eps": 2.01}}
        p1, p2, p3 = self._consensus_patches(actuals)
        with p1, p2, p3, \
                mock.patch.object(self.E, "fetch_all_estimates", lambda tks: {}):
            self.E._render_sector_aggregates(self.ALL_CONSENSUS, ["AAA", "BBB"])
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "sector_metric_breakdown.xlsx")
        self.assertEqual(kw["key"], "exp_sector_metric_breakdown")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Metric", "Banks", "Beat", "Miss", "Inline",
                                "Beat Rate (%)", "Avg Surprise (%)"])
        self.assertEqual(g[1][:6], ["Earnings Per Share", 2, 1, 0, 1, 50.0])
        self.assertAlmostEqual(g[1][6], 2.75, places=9)
        self.assertEqual(g[2][:6], ["Net Interest Margin", 1, 0, 1, 0, 0.0])
        self.assertAlmostEqual(g[2][6], -0.10 / 3.40 * 100, places=9)
        for coord in ("B2", "C2", "D2", "E2"):
            self.assertIsInstance(ws[coord].value, int)
            self.assertEqual(ws[coord].number_format, "#,##0")
        self.assertEqual(ws["F2"].number_format, '0.0"%"')
        self.assertEqual(ws["G2"].number_format, '0.0"%"')
        src = self._source(wb)
        self.assertIn("Per-Metric Breakdown", src["Page"])
        self.assertIn("SEC companyfacts", src["Source"])


class TestEarningsCalendar(_ExportSite):
    """Site: _render_earnings_calendar (raw agenda rows via the real
    build_calls_agenda). Revenue estimate is FMP's RAW dollars."""

    def test_calendar_rows_dates_and_raw_revenue(self):
        import data.bank_universe as bu
        import data.earnings_call as ec
        today = dt.date.today()
        d = (today + dt.timedelta(days=3)).isoformat()
        yf_cal = [{"ticker": "AAA", "next_earnings_date": d, "eps_estimate": 1.21}]
        fmp_cal = [{"symbol": "AAA", "date": d, "time": "bmo", "confirmed": True,
                    "epsEstimated": 1.20, "revenueEstimated": 150_000_000,
                    "periodEnding": "2026-06-30"}]
        with mock.patch.object(bu, "get_universe",
                               lambda: {"AAA": {"share_class": "common"}}), \
                mock.patch.object(self.E, "fetch_earnings_calendar", lambda u: yf_cal), \
                mock.patch.object(self.E, "_fmp_earnings_window", lambda a, b: fmp_cal), \
                mock.patch.object(ec, "merged_call_info", lambda: {}), \
                mock.patch.object(self.E, "get_name", _fake_name):
            self.E._render_earnings_calendar(["AAA"])
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], f"earnings_calendar_{today.isoformat()}.xlsx")
        self.assertEqual(kw["key"], "exp_earnings_calendar")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Ticker", "Bank", "Release Date", "Days Until",
                                "When", "Confirmed", "Period Ending", "Call Date",
                                "Call Time", "Webcast URL", "Dial-in",
                                "EPS Est ($)", "Rev Est ($)"])
        row = dict(zip(g[0], g[1]))
        self.assertEqual(row["Ticker"], "AAA")
        self.assertEqual(row["Bank"], "AAA Bancorp")
        self.assertEqual(row["Release Date"], dt.datetime.fromisoformat(d))
        self.assertEqual(row["Days Until"], 3)
        self.assertIs(row["Confirmed"], True)
        self.assertEqual(row["Period Ending"], dt.datetime(2026, 6, 30))
        for col in ("Call Date", "Call Time", "Webcast URL", "Dial-in"):
            self.assertEqual(row[col], "n/a")
        self.assertEqual(row["EPS Est ($)"], 1.21)          # yfinance preferred
        self.assertEqual(row["Rev Est ($)"], 150_000_000)   # unscaled
        self.assertEqual(ws["C2"].number_format, "yyyy-mm-dd")
        self.assertEqual(ws["D2"].number_format, "#,##0")
        self.assertEqual(ws["G2"].number_format, "yyyy-mm-dd")
        self.assertEqual(ws["L2"].number_format, "$#,##0.00")
        self.assertEqual(ws["M2"].number_format, "$#,##0")
        src = self._source(wb)
        self.assertIn("Earnings Calendar", src["Page"])
        self.assertEqual(src["Data as of"], today.isoformat())
        self.assertEqual(src["Horizon"], "75 days")
        self.assertIn("FMP", src["Source"])


class TestResultsBoard(_ExportSite):
    """Site: _render_results_board (raw board rows). Revenue columns —
    FMP actual/estimate and the release's total_revenue — are RAW dollars
    under ($), never rescaled."""

    def test_board_row_sources_and_release_columns(self):
        import data.earnings_results as er
        today = dt.date.today()
        rows = [{
            "ticker": "AAA", "date": "2026-07-15", "when": "Before open",
            "period_ending": "2026-06-30", "eps_act": 1.26, "eps_est": 1.20,
            "eps_surprise": 5.0, "eps_act_src": "release, adj.",
            "rev_act": 150_000_000.0, "rev_est": 148_000_000.0,
            "rev_surprise": 1.35, "reaction_session": "2026-07-15",
            "pr_headline": "AAA reports Q2", "pr_url": "https://ir.example/q2",
            "pending": False, "awaiting": False, "px_react": 2.5,
            "px_react_live": False, "rel": None,
        }, {
            "ticker": "BBB", "date": "2026-07-15", "when": None,
            "period_ending": None, "eps_act": None, "eps_est": 0.80,
            "eps_surprise": None, "rev_act": None, "rev_est": None,
            "rev_surprise": None, "reaction_session": "2026-07-15",
            "pr_headline": None, "pr_url": None, "pending": False,
            "awaiting": True, "px_react": None, "px_react_live": False,
            "rel": None,
        }]
        with mock.patch.object(er, "results_board", lambda *a, **k: rows), \
                mock.patch.object(self.E, "get_name", _fake_name):
            self.E._render_results_board()
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], f"earnings_results_{today.isoformat()}.xlsx")
        self.assertEqual(kw["key"], "exp_earnings_results")
        g = self._grid(ws)
        hdr = g[0]
        self.assertEqual(hdr[:20], [
            "Ticker", "Bank", "Reported", "When", "Period Ending", "Status",
            "EPS Act ($)", "EPS Act Source", "EPS Est ($)", "EPS Surprise (%)",
            "Rev Act ($)", "Rev Act Source", "Rev Est ($)", "Rev Surprise (%)",
            "Px React (%)", "Px React Live", "Reaction Session", "Release URL",
            "Release Headline", "Release Period End"])
        # One column per _REL_METRICS entry, unit in the header ($M → $).
        self.assertEqual(hdr[20:23], ["Release EPS adj ($)", "Release EPS GAAP ($)",
                                      "Release Revenue ($)"])
        self.assertIn("Release NIM (%)", hdr)
        self.assertIn("Release TBV/sh ($)", hdr)
        self.assertEqual(len(hdr), 20 + len(self.E._REL_METRICS))
        a = dict(zip(hdr, g[1]))
        self.assertEqual(a["Ticker"], "AAA")
        self.assertEqual(a["Reported"], dt.datetime(2026, 7, 15))
        self.assertEqual(a["Period Ending"], dt.datetime(2026, 6, 30))
        self.assertEqual(a["Status"], "reported")
        self.assertEqual(a["EPS Act ($)"], 1.26)
        self.assertEqual(a["EPS Act Source"], "release, adj.")
        self.assertEqual(a["EPS Surprise (%)"], 5.0)
        self.assertEqual(a["Rev Act ($)"], 150_000_000.0)   # unscaled
        self.assertEqual(a["Rev Act Source"], "FMP")
        self.assertEqual(a["Rev Est ($)"], 148_000_000.0)
        self.assertEqual(a["Px React (%)"], 2.5)
        self.assertIs(a["Px React Live"], False)
        self.assertEqual(a["Release URL"], "https://ir.example/q2")
        self.assertEqual(a["Release Revenue ($)"], "n/a")     # no release attached
        b = dict(zip(hdr, g[2]))
        self.assertEqual(b["Status"], "awaiting")
        for col in ("When", "Period Ending", "EPS Act ($)", "EPS Act Source",
                    "Rev Act ($)", "Rev Act Source", "Px React (%)", "Release URL"):
            self.assertEqual(b[col], "n/a", col)
        self.assertEqual(b["EPS Est ($)"], 0.8)
        col = {h: i + 1 for i, h in enumerate(hdr)}
        self.assertEqual(ws.cell(2, col["Reported"]).number_format, "yyyy-mm-dd")
        self.assertEqual(ws.cell(2, col["EPS Act ($)"]).number_format, "$#,##0.00")
        self.assertEqual(ws.cell(2, col["EPS Surprise (%)"]).number_format, '0.00"%"')
        self.assertEqual(ws.cell(2, col["Rev Act ($)"]).number_format, "$#,##0")
        self.assertEqual(ws.cell(2, col["Rev Est ($)"]).number_format, "$#,##0")
        self.assertEqual(ws.cell(2, col["Px React (%)"]).number_format, '0.00"%"')
        src = self._source(wb)
        self.assertIn("Reported Results", src["Page"])
        self.assertEqual(src["Data as of"], today.isoformat())
        self.assertIn("8-K", src["Source"])


if __name__ == "__main__":
    unittest.main()
