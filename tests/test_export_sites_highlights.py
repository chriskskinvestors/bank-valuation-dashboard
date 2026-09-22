"""Table exports on the shared .xlsx exporter (ui/export.py) for the
Financial Highlights, Capital Adequacy (RC-R walk) and Data Quality tables
(owner directive 2026-09-22: every data table gets an Export, built from the
RAW frame — never display strings).

Each test drives the real render function against small fixture data,
captures the deferred workbook callable table_export hands st.download_button,
loads the bytes with openpyxl and asserts HAND-COMPUTED values: row labels
carry the unit, cells are numbers (never "$1.2B" / "12.00%"), the Excel
number format per row/column, "n/a" where the fixture had no value, and the
Source sheet's provenance rows.

Units pinned here:
  * FDIC Call Report and RC-R values are $thousands as filed → exported
    UNSCALED under a "($K)" row label (owner choice 2026-09-22);
  * ratios are percent units (12.0 = 12.0%) under "(%)";
  * SEC per-share values are dollars per share; SEC USD amounts whole dollars.

Stub discipline: each render module's own ``st`` binding is patched with a
PRIVATE no-op stub per test (mock.patch.object); nothing here touches
sys.modules or the shared package stub, so this module composes with every
other suite under discovery.

Run: PYTHONIOENCODING=utf-8 python -m unittest tests.test_export_sites_highlights
"""
from __future__ import annotations

import datetime as dt
import io
import sys
import types
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).parent.parent))

import pandas as pd  # noqa: E402
from openpyxl import load_workbook  # noqa: E402


def _noop(*a, **k):
    return None


class _Ctx:
    """Context-manager placeholder (columns / container / expander / st.empty
    slot): every attribute is a no-op."""

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def __getattr__(self, name):
        return _noop


class _StubSt:
    """Private streamlit stand-in: unknown widgets are no-ops; layout widgets
    return contexts; radio returns ``radio_value`` (the page's period toggle)."""

    session_state: dict = {}
    query_params: dict = {}

    def __init__(self, radio_value="Annual"):
        self.radio_value = radio_value

    def __getattr__(self, name):
        return _noop

    def columns(self, spec, **k):
        n = spec if isinstance(spec, int) else len(spec)
        return [_Ctx() for _ in range(n)]

    def container(self, *a, **k):
        return _Ctx()

    expander = spinner = popover = empty = container

    def radio(self, label, options=None, **k):
        return self.radio_value


class _ExportSite(unittest.TestCase):
    """Binds a private stub onto every module the render paths reach through,
    and a download_button capture onto ui.export."""

    ST_MODULES: tuple = ()
    RADIO = "Annual"

    @classmethod
    def setUpClass(cls):
        import ui.chrome, ui.export, ui.states  # noqa: E401
        cls.X = ui.export
        cls._base_modules = (ui.chrome, ui.states)

    def setUp(self):
        self.calls = []          # [(label, data, kwargs)] from download_button
        self.stub = _StubSt(self.RADIO)
        fake_export_st = types.SimpleNamespace(
            download_button=lambda label, data, **k: self.calls.append((label, data, k)),
            container=lambda *a, **k: _Ctx())
        mods = list(self._base_modules) + list(self.ST_MODULES)
        for mod, obj in [(m, self.stub) for m in mods] + [(self.X, fake_export_st)]:
            p = mock.patch.object(mod, "st", obj)
            p.start()
            self.addCleanup(p.stop)

    def _patch(self, obj, name, value):
        p = mock.patch.object(obj, name, value)
        p.start()
        self.addCleanup(p.stop)

    # ── helpers ──────────────────────────────────────────────────────
    def _book(self, i=0, expect=1):
        self.assertEqual(len(self.calls), expect,
                         f"expected {expect} Export control(s), got {len(self.calls)}")
        label, data, kw = self.calls[i]
        self.assertEqual(label, "Export")
        self.assertTrue(callable(data), "workbook must build lazily, on click")
        wb = load_workbook(io.BytesIO(data()))
        return wb, wb.worksheets[0], kw

    @staticmethod
    def _grid(ws):
        return [[c.value for c in r] for r in ws.iter_rows()]

    @staticmethod
    def _source(wb):
        return dict((r[0].value, r[1].value) for r in wb["Source"].iter_rows())

    def _assert_na(self, ws, coord):
        self.assertEqual(ws[coord].value, "n/a")
        self.assertTrue(ws[coord].font.italic)


# ═══════════════════════════════════════════════════════════════════════
# ui/financial_highlights.py — metrics in rows × periods (row_formats)
# ═══════════════════════════════════════════════════════════════════════

def _fdic_row(repdte, **v):
    return {"REPDTE": repdte, **v}


# FY2024 / FY2025 year-end Call Reports ($thousands) + one Q3 '25 row that the
# Annual view must NOT show as a column (it is not a December quarter).
FDIC_HIST = [
    _fdic_row("2024-12-31", ASSET=900_000, LNLSNET=504_000, DEP=720_000, EQTOT=81_000,
              SC=162_000, NETINC=9_900, ROA=1.10, NIMY=3.30, EEFFR=62.0, NCLNLSR=0.60,
              NTLNLSR=0.12, LNATRESR=1.30, IDT1CER=10.5, RBCRWAJ=12.5, RBCT1JR=8.5,
              INTAN=9_000),
    _fdic_row("2025-09-30", ASSET=980_000, LNLSNET=590_000, DEP=790_000, EQTOT=95_000,
              SC=195_000, NETINC=9_000, ROA=1.15, NIMY=3.35, EEFFR=61.0, NCLNLSR=0.55,
              NTLNLSR=0.11, LNATRESR=1.28, IDT1CER=10.8, RBCRWAJ=12.8, RBCT1JR=8.8,
              INTAN=9_500),
    _fdic_row("2025-12-31", ASSET=1_000_000, LNLSNET=600_000, DEP=800_000, EQTOT=100_000,
              SC=200_000, NETINC=12_000, ROA=1.20, NIMY=3.40, EEFFR=60.0, NCLNLSR=0.50,
              NTLNLSR=0.10, LNATRESR=1.25, IDT1CER=11.0, RBCRWAJ=13.0, RBCT1JR=9.0,
              INTAN=10_000),
]

# Holding-company companyfacts: FY2025 only, so every FY2024 per-share cell
# is genuinely absent (→ n/a). Equity 1,500,000 − goodwill 300,000 over
# 100,000 shares: BVPS 15.00, TBVPS 12.00.
_FY25 = {"form": "10-K", "start": "2025-01-01", "end": "2025-12-31",
         "accn": "0000946673-26-000010", "filed": "2026-02-20"}
SEC_FACTS = {"facts": {
    "us-gaap": {
        "EarningsPerShareDiluted": {"units": {"USD/shares": [{**_FY25, "val": 2.49}]}},
        "CommonStockDividendsPerShareDeclared": {"units": {"USD/shares": [{**_FY25, "val": 0.65}]}},
        "StockholdersEquity": {"units": {"USD": [
            {"form": "10-K", "end": "2025-12-31", "val": 1_500_000,
             "accn": _FY25["accn"], "filed": _FY25["filed"]}]}},
        "Goodwill": {"units": {"USD": [{"form": "10-K", "end": "2025-12-31", "val": 300_000}]}},
    },
    "dei": {"EntityCommonStockSharesOutstanding": {"units": {"shares": [
        {"form": "10-K", "end": "2025-12-31", "val": 100_000,
         "accn": _FY25["accn"], "filed": _FY25["filed"]}]}}},
}}

BANK_INFO = {"name": "Banner Corp", "fdic_cert": 28489, "cik": "946673"}


class _HighlightsSite(_ExportSite):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        import ui.financial_highlights
        cls.FH = ui.financial_highlights
        cls.ST_MODULES = (ui.financial_highlights,)

    def _render(self, facts):
        import data.cache as cache
        import data.loaders as dl
        FH = self.FH
        self._patch(FH, "components", types.SimpleNamespace(html=_noop))
        self._patch(FH, "get_bank_info", lambda t: dict(BANK_INFO))
        self._patch(FH, "_render_fh_trends", _noop)
        self._patch(FH.sec_client, "fetch_company_facts", lambda cik: facts)
        self._patch(dl, "load_fdic_hist_df",
                    lambda ticker, quarters=36: pd.DataFrame(FDIC_HIST))
        self._patch(cache, "get_age", lambda key: None)
        FH.render_financial_highlights("BANR")


class TestFinancialHighlightsAnnual(_HighlightsSite):
    RADIO = "Annual"

    def test_row_formats_dollar_pct_per_share_and_missing_period(self):
        self._render(SEC_FACTS)
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "financial_highlights_BANR_Annual_FY2025.xlsx")
        self.assertEqual(kw["key"], "exp_fh_BANR_Annual")
        g = self._grid(ws)
        # Only the two December quarters are periods; Q3 '25 is not a column.
        self.assertEqual(g[0], ["Line item", "Section", "FY2024", "FY2025"])
        self.assertEqual(len(g), 1 + 26)
        rows = {r[0]: r for r in g[1:]}
        by_label = {r[0]: i + 1 for i, r in enumerate(g)}     # label → sheet row

        # ── FDIC dollar rows: $thousands UNSCALED, "($K)" in the label ──
        self.assertEqual(rows["Total assets ($K)"], ["Total assets ($K)", "Balance Sheet",
                                                     900_000, 1_000_000])
        self.assertEqual(rows["Net income (YTD) ($K)"][1:], ["Profitability", 9_900, 12_000])
        r = by_label["Total assets ($K)"]
        self.assertEqual(ws.cell(r, 3).number_format, "#,##0")
        self.assertEqual(ws.cell(r, 4).number_format, "#,##0")

        # ── FDIC reported ratios: percent units, "(%)" ──
        self.assertEqual(rows["ROAA (%)"][2:], [1.10, 1.20])
        self.assertEqual(rows["CET1 ratio (%)"][1:], ["Capital Adequacy (bank-level)", 10.5, 11.0])
        # ── computed ratios (hand-checked) ──
        # ROAE FY2025 = 12,000 / 100,000 × 100 = 12.00; FY2024 = 9,900 / 81,000 = 12.2222
        self.assertAlmostEqual(rows["ROAE (%)"][2], 12.222222222, places=8)
        self.assertEqual(rows["ROAE (%)"][3], 12.0)
        # ROATCE FY2025 = 12,000 / (100,000 − 10,000) = 13.3333; FY2024 = 9,900 / 72,000 = 13.75
        self.assertAlmostEqual(rows["ROATCE (%)"][2], 13.75, places=9)
        self.assertAlmostEqual(rows["ROATCE (%)"][3], 13.333333333, places=8)
        # Loans / deposits FY2024 = 504,000 / 720,000 = 70.0; FY2025 = 600/800 = 75.0
        self.assertEqual(rows["Loans / deposits (%)"][1:], ["Balance Sheet Ratios", 70.0, 75.0])
        self.assertEqual(rows["Securities / assets (%)"][2:], [18.0, 20.0])
        self.assertEqual(rows["Equity / assets (%)"][2:], [9.0, 10.0])
        # TCE/TA FY2025 = 90,000 / 990,000 = 9.0909; FY2024 = 72,000 / 891,000 = 8.0808
        tce = rows["Tang. common equity / tang. assets (%)"]
        self.assertAlmostEqual(tce[2], 8.080808081, places=8)
        self.assertAlmostEqual(tce[3], 9.090909091, places=8)
        r = by_label["ROAE (%)"]
        self.assertEqual(ws.cell(r, 3).number_format, '0.00"%"')
        self.assertEqual(ws.cell(r, 4).number_format, '0.00"%"')

        # ── SEC per-share rows: $/share for FY2025, n/a for FY2024 ──
        self.assertEqual(rows["Diluted EPS ($/share)"][1], "Per Share (HoldCo)")
        self.assertEqual(rows["Diluted EPS ($/share)"][3], 2.49)
        self.assertEqual(rows["Dividends / share ($/share)"][3], 0.65)
        self.assertEqual(rows["Book value / share ($/share)"][3], 15.0)
        self.assertEqual(rows["Tangible BV / share ($/share)"][3], 12.0)
        self.assertEqual(rows["Shares outstanding"][3], 100_000)
        for lab in ("Diluted EPS ($/share)", "Dividends / share ($/share)",
                    "Book value / share ($/share)", "Tangible BV / share ($/share)",
                    "Shares outstanding"):
            self._assert_na(ws, f"C{by_label[lab]}")
        r = by_label["Diluted EPS ($/share)"]
        self.assertEqual(ws.cell(r, 4).number_format, "$#,##0.00")
        r = by_label["Shares outstanding"]
        self.assertEqual(ws.cell(r, 4).number_format, "#,##0")
        # label columns carry no number format; two frozen text columns
        self.assertEqual(ws["A2"].number_format, "General")
        self.assertEqual(ws.freeze_panes, "C2")

        src = self._source(wb)
        self.assertEqual(src["Page"], "Company Analysis · Financial Highlights")
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["Company"], "Banner Corp")
        self.assertEqual(src["FDIC cert"], 28489)
        self.assertEqual(src["SEC CIK"], "946673")
        self.assertIn("FDIC Call Report (bank subsidiary)", src["Source"])
        self.assertIn("SEC companyfacts (holding company)", src["Source"])
        self.assertEqual(src["Period"], "Annual — FY2024 to FY2025")
        self.assertEqual(src["Report date"], "Dec-31-2025")
        self.assertIn("($K)", src["Units"])
        self.assertIn("percent units", src["Units"])


class TestFinancialHighlightsQuarterly(_HighlightsSite):
    RADIO = "Quarterly"

    def test_quarter_columns_and_annualized_roae(self):
        # No holding-company facts at all → every per-share cell is n/a.
        self._render({})
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "financial_highlights_BANR_Quarterly_Q4_25.xlsx")
        self.assertEqual(kw["key"], "exp_fh_BANR_Quarterly")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Line item", "Section", "Q4 '24", "Q3 '25", "Q4 '25"])
        rows = {r[0]: r for r in g[1:]}
        # Q3 '25 ROAE annualizes YTD NI: 9,000 × 12/9 / 95,000 × 100 = 12.6316
        self.assertAlmostEqual(rows["ROAE (%)"][3], 12.631578947, places=8)
        self.assertEqual(rows["ROAE (%)"][4], 12.0)
        # Q3 '25 ROATCE: 12,000 / (95,000 − 9,500) = 14.0351
        self.assertAlmostEqual(rows["ROATCE (%)"][3], 14.035087719, places=8)
        self.assertEqual(rows["Total assets ($K)"][2:], [900_000, 980_000, 1_000_000])
        for lab in ("Diluted EPS ($/share)", "Book value / share ($/share)", "Shares outstanding"):
            self.assertEqual(rows[lab][2:], ["n/a", "n/a", "n/a"])
        src = self._source(wb)
        self.assertEqual(src["Period"], "Quarterly — Q4 '24 to Q4 '25")
        self.assertEqual(src["Report date"], "Dec-31-2025")


# ═══════════════════════════════════════════════════════════════════════
# ui/capital_dynamics.py — RC-R Part I capital walk (bank subsidiary)
# ═══════════════════════════════════════════════════════════════════════

# Banner Bank's filed 12/31/2025 Schedule RC-R Part I ($000) — the same
# hand-checked fixture as tests/test_render_smoke.TestCapitalWalkRendersPopulated.
BANNER_Q4 = {
    "reporting_period": "12/31/2025", "rssd_id": 352772,
    "cet1_before_adjustments": 1_951_461.0,
    "goodwill_deduction": 370_753.0, "other_intangibles_deduction": 2_237.0,
    "dta_deduction": 6_912.0,
    "aoci_adj_unrealized_afs": -213_012.0, "aoci_adj_afs_preferred": None,
    "aoci_adj_cash_flow_hedges": None, "aoci_adj_pension": None, "aoci_adj_htm": None,
    "cet1": 1_784_571.0, "additional_tier1": 0.0, "tier1": 1_784_571.0,
    "t2_instruments": 0.0, "t2_nonqualifying_instruments": None,
    "t2_minority_interest": None, "t2_allowance": 173_048.0, "tier2": 173_048.0,
    "total_capital": 1_957_619.0, "rwa": 13_841_345.0,
    "intangibles_deduction": 372_990.0, "aoci_adjustment": 213_012.0,
    "other_cet1_adjustments": 0.0, "t2_other": 0.0,
}
# A prior quarter with the DTA line and RWA absent from the filing → n/a for
# the DTA row and for every ratio (never 0, never a division guess).
BANNER_Q3 = {**BANNER_Q4, "reporting_period": "09/30/2025",
             "cet1_before_adjustments": 1_900_000.0, "dta_deduction": None,
             "cet1": 1_740_000.0, "tier1": 1_740_000.0, "total_capital": 1_910_000.0,
             "rwa": None}


class TestCapitalWalkExport(_ExportSite):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # Leave the import state exactly as found: test_render_smoke's
        # TestCapitalWalkRendersPopulated re-installs its own streamlit stub
        # and patches `sys.modules["streamlit"].info`, which only reaches
        # ui.capital_dynamics if THAT class is the first to import it (the
        # stub-replace hazard). Un-importing in tearDownClass keeps this suite
        # composable under discovery without touching the smoke suite.
        cls._imported_here = "ui.capital_dynamics" not in sys.modules
        import ui.capital_dynamics
        cls.CD = ui.capital_dynamics
        cls.ST_MODULES = (ui.capital_dynamics,)

    @classmethod
    def tearDownClass(cls):
        if cls._imported_here:
            sys.modules.pop("ui.capital_dynamics", None)
        super().tearDownClass()

    def _render(self, stored):
        import data.bank_mapping as bm
        import data.call_report_store as crs
        self._patch(sys.modules["streamlit.components.v1"], "html", _noop)
        self._patch(crs, "get_stored_rcr_detail", lambda cert, quarters=8: list(stored))
        self._patch(bm, "get_name", lambda t: "Banner Corp")
        self._patch(self.CD, "get_fdic_cert", lambda t: 28489)
        self.CD._render_rcr_capital_walk("BANR")

    def test_walk_rows_k_dollars_ratios_and_absent_lines(self):
        self._render([dict(BANNER_Q4), dict(BANNER_Q3)])      # store: newest first
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "capital_walk_rcr_BANR_2025-12-31.xlsx")
        self.assertEqual(kw["key"], "exp_rcr_walk_BANR")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Line item", "Section", "Q3 '25", "Q4 '25"])   # oldest left
        self.assertEqual(len(g), 1 + 17)               # 6 CET1 + 2 T1 + 5 T2 + 4 RWA/ratios
        rows = {r[0]: r for r in g[1:]}
        by_label = {r[0]: i + 1 for i, r in enumerate(g)}
        self.assertEqual(rows["CET1 before adjustments & deductions ($K)"],
                         ["CET1 before adjustments & deductions ($K)", "Common Equity Tier 1",
                          1_900_000, 1_951_461])
        # derived: 370,753 + 2,237 = 372,990
        self.assertEqual(rows["Less: intangibles (goodwill + other) ($K)"][2:],
                         [372_990, 372_990])
        # AOCI add-back positive (losses removed from capital), never −213,012
        self.assertEqual(rows["AOCI adjustment (positive = added back) ($K)"][2:],
                         [213_012, 213_012])
        self.assertEqual(rows["Other tier 2 components ($K)"][2:], [0, 0])
        self.assertEqual(rows["Total capital ($K)"][1:],
                         ["Tier 2 & Total Capital", 1_910_000, 1_957_619])
        self.assertEqual(rows["Total risk-weighted assets ($K)"][3], 13_841_345)
        # Ratios = component ÷ RWA × 100: 1,784,571 / 13,841,345 = 12.8930462;
        # 1,957,619 / 13,841,345 = 14.1432715
        self.assertAlmostEqual(rows["CET1 ratio (%)"][3], 12.8930462, places=6)
        self.assertAlmostEqual(rows["Tier 1 ratio (%)"][3], 12.8930462, places=6)
        self.assertAlmostEqual(rows["Total capital ratio (%)"][3], 14.1432715, places=6)
        self.assertEqual(rows["CET1 ratio (%)"][1], "Risk-Weighted Assets & Ratios")
        # Q3: DTA line absent and RWA absent → n/a, not 0
        self._assert_na(ws, f"C{by_label['Less: DTAs from carryforwards ($K)']}")
        self.assertEqual(rows["Less: DTAs from carryforwards ($K)"][3], 6_912)
        for lab in ("Total risk-weighted assets ($K)", "CET1 ratio (%)",
                    "Tier 1 ratio (%)", "Total capital ratio (%)"):
            self._assert_na(ws, f"C{by_label[lab]}")
        r = by_label["Total capital ($K)"]
        self.assertEqual(ws.cell(r, 3).number_format, "#,##0")
        self.assertEqual(ws.cell(r, 4).number_format, "#,##0")
        r = by_label["CET1 ratio (%)"]
        self.assertEqual(ws.cell(r, 4).number_format, '0.00"%"')
        self.assertEqual(ws.freeze_panes, "C2")
        src = self._source(wb)
        self.assertIn("Regulatory Capital Walk", src["Page"])
        self.assertIn("bank subsidiary", src["Page"])
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["Company"], "Banner Corp")
        self.assertEqual(src["FDIC cert"], 28489)
        self.assertIn("Schedule RC-R Part I", src["Source"])
        self.assertIn("bank subsidiary", src["Source"])
        self.assertEqual(src["Periods"], "Q3 '25 to Q4 '25")
        self.assertEqual(src["Report date"], "12/31/2025")
        self.assertIn("($K)", src["Units"])


# ═══════════════════════════════════════════════════════════════════════
# ui/data_quality.py — findings, SEC sources, FDIC fields, FFIEC ladder
# ═══════════════════════════════════════════════════════════════════════

def _fdic_frame():
    return pd.DataFrame([{
        "REPDTE": pd.Timestamp("2025-12-31"), "REPNM": "Banner Bank",
        "ASSET": 16_000_000.0, "DEP": 13_500_000.0, "LNLSNET": 11_200_000.0,
        "EQTOT": 1_800_000.0, "NETINC": 190_000.0, "NIMY": 3.85, "ROA": 1.20,
        "IDT1CER": 12.89, "NCLNLSR": 0.45,
    }])


def _sec_prov():
    from data.provenance import Source

    def s(concept, unit, as_of="2025-12-31", form="10-K", notes=""):
        return Source(origin="SEC", identifier="946673", concept=concept, as_of=as_of,
                      filed="2026-02-20", form=form, unit=unit, notes=notes)
    return {
        "book_value_total": {"value": 1_500_000_000.0, "source": s("StockholdersEquity", "USD")},
        "eps": {"value": 2.49, "source": s("EarningsPerShareDiluted", "USD/shares")},
        "shares_outstanding": {"value": 100_000_000.0,
                               "source": s("EntityCommonStockSharesOutstanding", "shares",
                                           notes="dei cover-page count")},
        "revenue": {"value": None, "source": s("Revenues", "USD")},   # absent → not a row
    }


class _DataQualitySite(_ExportSite):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        import ui.data_quality
        cls.DQ = ui.data_quality
        cls.ST_MODULES = (ui.data_quality,)

    def _render(self, *, tab, sec_prov, bank_metrics):
        import data.cache as cache
        import data.ffiec_client as fc
        DQ = self.DQ
        self._patch(DQ, "get_fdic_cert", lambda t: 28489)
        self._patch(DQ, "get_cik", lambda t: 946673)
        self._patch(DQ, "get_name", lambda t: "Banner Corp")
        self._patch(DQ.sec_client, "get_fundamentals_with_provenance", lambda cik: sec_prov)
        self._patch(DQ, "fetch_financials", lambda cert, limit=1: _fdic_frame())
        self._patch(cache, "get", lambda key, *a, **k: bank_metrics)
        self._patch(DQ, "lazy_tabs", lambda labels, key=None, default=0: labels[tab])
        self._patch(fc, "is_configured", lambda: False)
        DQ.render_data_quality("BANR")


class TestDataQualityFindings(_DataQualitySite):

    def test_finding_value_is_a_number(self):
        # roaa 5.0 breaches the [-3.0, 3.5] % band → one range warning. No SEC
        # provenance, so no cross-source checks; the FDIC staleness check fires
        # too (a 2025-12-31 Call Report is >135 days old from today on) with
        # value=None — the honest n/a row.
        self._render(tab=0, sec_prov={}, bank_metrics=[{"ticker": "BANR", "roaa": 5.0}])
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "data_quality_findings_BANR_2025-12-31.xlsx")
        self.assertEqual(kw["key"], "exp_dq_findings_BANR")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Severity", "Field", "Issue", "Value", "Data source"])
        self.assertEqual(len(g), 3)
        self.assertEqual(g[1][:2], ["Warning", "roaa"])
        self.assertIn("outside expected range [-3.0, 3.5] %", g[1][2])
        self.assertEqual(g[1][3], 5.0)
        self.assertEqual(ws["D2"].number_format, "#,##0.00")
        self._assert_na(ws, "E2")                     # Finding.source is "" here
        self.assertEqual(g[2][:2], ["Warning", "fdic_call_report"])
        self.assertIn("days old", g[2][2])
        self._assert_na(ws, "D3")                     # staleness finding has no value
        src = self._source(wb)
        self.assertIn("Validation Findings", src["Page"])
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["FDIC cert"], 28489)
        self.assertEqual(src["SEC CIK"], 946673)
        self.assertIn("data/validation.py", src["Source"])
        self.assertEqual(src["Report date"], "2025-12-31")
        self.assertEqual(src["Checks"], "0 errors, 2 warnings")


class TestDataQualitySources(_DataQualitySite):

    def test_sec_and_fdic_source_tables(self):
        self._render(tab=1, sec_prov=_sec_prov(), bank_metrics=[])
        self.assertEqual(len(self.calls), 2, "SEC sources + FDIC fields exports")

        # ── SEC HoldCo Sources: raw values, format per XBRL unit ──
        wb, ws, kw = self._book(0, expect=2)
        self.assertEqual(kw["file_name"], "data_quality_sec_sources_BANR_2025-12-31.xlsx")
        self.assertEqual(kw["key"], "exp_dq_sec_BANR")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Metric", "Value", "XBRL Concept", "As Of", "Age (days)",
                                "Form", "Unit", "Notes"])
        self.assertEqual(len(g), 1 + 3)                # revenue (None) is not a row
        self.assertEqual(g[1][:4], ["book_value_total", 1_500_000_000, "StockholdersEquity",
                                    dt.datetime(2025, 12, 31)])
        self.assertEqual(g[2][:3], ["eps", 2.49, "EarningsPerShareDiluted"])
        self.assertEqual(g[3][:3], ["shares_outstanding", 100_000_000,
                                    "EntityCommonStockSharesOutstanding"])
        self.assertEqual(ws["B2"].number_format, "$#,##0")        # USD → whole dollars
        self.assertEqual(ws["B3"].number_format, "$#,##0.00")     # USD/shares
        self.assertEqual(ws["B4"].number_format, "#,##0")         # shares
        self.assertEqual(ws["D2"].number_format, "yyyy-mm-dd")
        self.assertIsInstance(ws["E2"].value, int)                # age in days, a number
        self.assertEqual(ws["E2"].number_format, "#,##0")         # column wins over row
        self.assertEqual(g[1][5:], ["10-K", "USD", "n/a"])
        self.assertEqual(g[3][7], "dei cover-page count")
        src = self._source(wb)
        self.assertIn("SEC HoldCo Sources", src["Page"])
        self.assertEqual(src["Source"], "SEC companyfacts (holding company)")
        self.assertEqual(src["SEC CIK"], 946673)
        self.assertEqual(src["Data as of"], "2025-12-31")
        self.assertNotIn("FDIC cert", src)

        # ── FDIC Call Report fields: $thousands UNSCALED under ($K), ratios (%) ──
        wb, ws, kw = self._book(1, expect=2)
        self.assertEqual(kw["file_name"], "data_quality_fdic_fields_BANR_2025-12-31.xlsx")
        self.assertEqual(kw["key"], "exp_dq_fdic_BANR")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Label", "FDIC Field", "Value", "Unit", "As Of"])
        self.assertEqual(len(g), 1 + 9)
        self.assertEqual(g[1], ["Total Assets ($K)", "ASSET", 16_000_000, "$thousands",
                                dt.datetime(2025, 12, 31)])
        self.assertEqual(g[5][:3], ["Net Income (YTD) ($K)", "NETINC", 190_000])
        self.assertEqual(g[6][:4], ["NIM (annualized) (%)", "NIMY", 3.85, "%"])
        self.assertEqual(g[8][:3], ["CET1 Ratio (%)", "IDT1CER", 12.89])
        self.assertEqual(ws["C2"].number_format, "#,##0")         # usd_k
        self.assertEqual(ws["C7"].number_format, '0.00"%"')       # pct
        self.assertEqual(ws["E2"].number_format, "yyyy-mm-dd")
        src = self._source(wb)
        self.assertIn("FDIC Call Report Source", src["Page"])
        self.assertEqual(src["Source"], "FDIC Call Report (bank subsidiary)")
        self.assertEqual(src["FDIC cert"], 28489)
        self.assertEqual(src["Institution"], "Banner Bank")
        self.assertEqual(src["Report date"], "2025-12-31")
        self.assertIn("($K)", src["Units"])
        self.assertIn("percent units", src["Units"])


class TestDataQualityFfiecLadder(_DataQualitySite):

    def _render_ladder(self, ladder):
        import data.call_report_store as crs
        import data.ffiec_client as fc
        self._patch(fc, "is_configured", lambda: True)
        self._patch(fc, "health_check", lambda: {"ok": True, "days_until_expiry": 120})
        self._patch(crs, "get_latest_ladder", lambda cert: ladder)
        self.DQ._render_ffiec_status(28489, "BANR")

    def test_ladder_duration_years_and_share_percent(self):
        # floating_loan_share is a FRACTION at source (0.42) → 42.0 (%) exported.
        self._render_ladder({"reporting_period": "12/31/2025",
                             "weighted_avg_duration_years": 4.37,
                             "floating_loan_share": 0.42, "source": "ffiec"})
        wb, ws, kw = self._book()
        self.assertEqual(kw["file_name"], "ffiec_ladder_BANR_12_31_2025.xlsx")
        self.assertEqual(kw["key"], "exp_dq_ffiec_BANR")
        g = self._grid(ws)
        self.assertEqual(g[0], ["Field", "Value"])
        self.assertEqual(g[1], ["Reporting period", "12/31/2025"])
        self.assertEqual(g[2], ["Securities duration (wtd-avg, years)", 4.37])
        self.assertAlmostEqual(g[3][1], 42.0, places=9)
        self.assertEqual(g[3][0], "Floating-loan share (RC-C Memo 2, %)")
        self.assertEqual(g[4], ["Source", "ffiec"])
        self.assertEqual(ws["B3"].number_format, "#,##0.00")
        self.assertEqual(ws["B4"].number_format, '0.00"%"')
        src = self._source(wb)
        self.assertIn("FFIEC Call Report Ladder", src["Page"])
        self.assertEqual(src["Ticker"], "BANR")
        self.assertEqual(src["FDIC cert"], 28489)
        self.assertIn("FFIEC Call Report (bank subsidiary)", src["Source"])
        self.assertEqual(src["Report date"], "12/31/2025")

    def test_unreported_share_and_duration_are_na_not_zero(self):
        # The display falls back to "0.00 yrs" for a missing duration; the
        # export must say n/a (never a fabricated 0).
        self._render_ladder({"reporting_period": "12/31/2025",
                             "weighted_avg_duration_years": None,
                             "floating_loan_share": None, "source": "ffiec"})
        wb, ws, kw = self._book()
        self._assert_na(ws, "B3")
        self._assert_na(ws, "B4")


if __name__ == "__main__":
    unittest.main()
