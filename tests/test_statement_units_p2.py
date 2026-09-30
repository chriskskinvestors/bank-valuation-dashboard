"""Statement units / dead-cell reasons (docs/REVIEW-2026-09-24-numbers.md P2).

  * P2-2  Templated section headers said "ASSETS ($000)" over $-compact cells
          ("$4,091.39B"). The screen header drops the "($000)"; the Excel
          export keeps the spec's section name and its "($K)" row labels
          (raw FDIC thousands, unscaled).
  * P2-6  ONB's stated-value common-stock row ("… $1.00 per share stated
          value …", 389,662,000 raw dollars) rendered as "$389,662,000.00":
          the "per share" label read as EPS. Its XBRL kind is monetary, so it
          renders $-compact; a true EPS row stays $x.xx. (Render-path pin:
          tests/test_export_sites_cr.py TestBareLabelEpsRows.)
  * P2-10 FTE adjustment / NII (FTE) rendered a bare "—" for a quarter the
          FFIEC Schedule RI store has not ingested yet. Now "n/a" with a
          click-through reason — never an imputed value.
"""
import io
import re
import types
import unittest
from unittest import mock

import pandas as pd
from openpyxl import load_workbook

import ui.export as ex
import ui.financials_statements as FS


class _Ctx:
    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def _stub_st(period):
    st = types.SimpleNamespace()
    st.radio = lambda *a, **k: period
    st.markdown = st.caption = st.warning = st.info = lambda *a, **k: None
    st.spinner = lambda *a, **k: _Ctx()
    st.columns = lambda spec, **k: [_Ctx() for _ in (spec if isinstance(spec, list) else range(spec))]
    st.container = lambda *a, **k: _Ctx()
    return st


def _render(spec, hist, period="Annual", ri_rows=(), flow_fields=None):
    """Render one Templated statement; returns (iframe html, export workbook)."""
    calls, html = [], []
    fake_ex = types.SimpleNamespace(
        download_button=lambda label, data, **k: calls.append(data),
        container=lambda *a, **k: _Ctx())
    fake_components = types.SimpleNamespace(html=lambda h, **k: html.append(h))
    import data.loaders as loaders
    import data.call_report_store as crs
    with mock.patch.object(FS, "st", _stub_st(period)), \
         mock.patch.object(FS, "components", fake_components), \
         mock.patch.object(FS, "get_bank_info",
                           lambda t: {"name": "Test Bancorp", "fdic_cert": 4242, "cik": None}), \
         mock.patch.object(FS, "_render_statement_trends", lambda *a, **k: None), \
         mock.patch.object(loaders, "load_fdic_hist_df", lambda t, q: hist.copy()), \
         mock.patch.object(crs, "get_stored_ri_detail", lambda cert, quarters=40: list(ri_rows)), \
         mock.patch.object(crs, "get_stored_rie_detail", lambda cert, quarters=40: []), \
         mock.patch.object(ex, "st", fake_ex):
        FS.render_statement("TBNK", "stmt", "Statement", spec, trends=[],
                            with_ri=True, flow_fields=flow_fields)
    assert len(html) == 1 and len(calls) == 1, (html, calls)
    return html[0], load_workbook(io.BytesIO(calls[0]())).worksheets[0]


def _sec_headers(html):
    return re.findall(r'<td class="sec"[^>]*>(.*?)</td>', html, flags=re.S)


def _row_cells(html, label):
    """Raw <td class="val…"> tags + text for one row label."""
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", html, flags=re.S):
        m = re.search(r'<td class="lbl">(.*?)</td>', tr, flags=re.S)
        if m and re.sub(r"<.*?>", "", m.group(1)).strip() == label:
            return re.findall(r'(<td class="val[^"]*"[^>]*>)(.*?)</td>', tr, flags=re.S)
    raise AssertionError(f"row {label!r} not rendered")


class TestSectionHeaderUnits(unittest.TestCase):
    """P2-2: screen header never claims $000 over $-compact cells."""

    HIST = pd.DataFrame([{"REPDTE": "2025-12-31", "ASSET": 4_091_390_000,
                          "DEP": 2_500_000_000, "EQTOT": 350_000_000}])
    SPEC = [("Assets ($000)", [("Total assets", "dollar", "ASSET")]),
            ("Liabilities ($000)", [("Total deposits", "dollar", "DEP")]),
            ("Balance Sheet Analysis (%)", [("Equity", "dollar", "EQTOT")])]

    @classmethod
    def setUpClass(cls):
        cls.html, cls.ws = _render(cls.SPEC, cls.HIST)

    def test_screen_header_drops_000(self):
        self.assertEqual(_sec_headers(self.html),
                         ["Assets", "Liabilities", "Balance Sheet Analysis (%)"])
        self.assertNotIn("($000)", "".join(_sec_headers(self.html)))
        # the cell beneath is $-compact: 4,091,390,000 $K → "$4,091.39B"
        self.assertEqual([t for _, t in _row_cells(self.html, "Total assets")],
                         ["$4,091.39B"])

    def test_export_keeps_k_labels_and_raw_thousands(self):
        rows = [[c.value for c in r] for r in self.ws.iter_rows()]
        self.assertEqual(rows[1], ["Assets ($000)", "n/a"])      # export header unchanged
        self.assertEqual(rows[2], ["Total assets ($K)", 4_091_390_000])
        self.assertEqual(rows[3], ["Liabilities ($000)", "n/a"])

    def test_every_real_spec_header_is_unit_clean_on_screen(self):
        specs = [v for k, v in vars(FS).items()
                 if k.startswith("_") and isinstance(v, list)
                 and v and all(isinstance(s, tuple) and len(s) == 2
                               and isinstance(s[0], str) and isinstance(s[1], list)
                               for s in v)]
        names = [s[0] for spec in specs for s in spec]
        self.assertIn("Regulatory Capital ($000)", names)       # the scan found the specs
        self.assertIn("Assets ($000)", names)
        for n in names:
            self.assertNotIn("$000", FS._screen_section(n), n)


class TestOnbStatedValueCommonStock(unittest.TestCase):
    """P2-6: a monetary row whose label mentions "per share" is $-compact."""

    ONB_LABEL = ("Common stock, no par value, $1.00 per share stated value, 600,000 "
                 "shares authorized, 389,662 and 318,980 shares issued and outstanding")

    def test_monetary_kind_beats_per_share_label(self):
        row = {"label": self.ONB_LABEL, "header": False, "kind": "monetary",
               "values": [389_662_000.0, 318_980_000.0]}
        kind = FS._cr_line_kind(row)
        self.assertEqual(kind, "usd")
        self.assertEqual(FS._cr_fmt(kind, 389_662_000.0), "$389.7M")
        self.assertNotEqual(FS._cr_fmt(kind, 389_662_000.0), "$389,662,000.00")

    def test_true_eps_row_stays_dollars_and_cents(self):
        row = {"label": "Diluted earnings per share", "kind": "pershare", "values": [0.55]}
        kind = FS._cr_line_kind(row)
        self.assertEqual(kind, "eps")
        self.assertEqual(FS._cr_fmt(kind, 0.55), "$0.55")


class TestFteRowsNotIngestedReason(unittest.TestCase):
    """P2-10: a column with no Schedule RI detail is n/a WITH a reason."""

    HIST = pd.DataFrame([
        {"REPDTE": "2025-12-31", "INTINC": 1_300_000, "EINTEXP": 460_000},
        {"REPDTE": "2026-03-31", "INTINC": 360_000, "EINTEXP": 130_000},
        {"REPDTE": "2026-06-30", "INTINC": 740_000, "EINTEXP": 270_000},
    ])
    # Store is a quarter behind: Q1 '26 ingested, Q2 '26 not yet.
    RI = [{"reporting_period": "2026-03-31", "tax_exempt_loan_income": 4_000,
           "tax_exempt_sec_income": 1_000}]
    SPEC = [("Interest Income & Expense", [
        ("FTE adjustment", "fte_adj"),
        ("Net interest income (FTE)", "nii_fte"),
    ])]
    FLOWS = frozenset({"INTINC", "EINTEXP"})

    @classmethod
    def setUpClass(cls):
        cls.q_html, _ = _render(cls.SPEC, cls.HIST, "Quarterly", cls.RI, cls.FLOWS)
        cls.a_html, cls.a_ws = _render(cls.SPEC, cls.HIST, "Annual", cls.RI, cls.FLOWS)

    def test_latest_quarter_is_na_with_clickable_reason(self):
        for label in ("FTE adjustment", "Net interest income (FTE)"):
            cells = _row_cells(self.q_html, label)
            # columns: Q4 '25 (no RI), Q1 '26 (ingested), Q2 '26 (not yet)
            tag, text = cells[-1]
            self.assertEqual(text, "n/a", label)
            self.assertNotEqual(text, "—", label)
            self.assertIn("data-cid", tag, label)             # click-through, not a dead cell
            self.assertNotIn("dead", tag, label)
        self.assertIn("Schedule RI detail not yet ingested for this period", self.q_html)
        self.assertIn("needed to de-cumulate YTD", self.q_html)

    def test_ingested_quarter_still_computes(self):
        # Q1 as filed: (4,000 + 1,000) × 0.21 ÷ 0.79 = 1,329.1 ($000) → $1.3M;
        # NII (FTE) = (360,000 − 130,000) + 1,329.1 = 231,329.1 → $231.3M.
        self.assertEqual(_row_cells(self.q_html, "FTE adjustment")[1][1], "$1.3M")
        self.assertEqual(_row_cells(self.q_html, "Net interest income (FTE)")[1][1], "$231.3M")

    def test_annual_column_reason_has_no_decumulation_clause(self):
        tag, text = _row_cells(self.a_html, "FTE adjustment")[0]      # FY2025, no RI
        self.assertEqual(text, "n/a")
        self.assertIn("data-cid", tag)
        self.assertIn("Schedule RI detail not yet ingested for this period", self.a_html)
        self.assertNotIn("needed to de-cumulate YTD", self.a_html)
        # export: absent stays n/a, never 0
        rows = [[c.value for c in r] for r in self.a_ws.iter_rows()]
        fte = next(r for r in rows if str(r[0]).startswith("FTE adjustment"))
        self.assertEqual(fte[1], "n/a")


if __name__ == "__main__":
    unittest.main()
